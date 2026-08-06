"""
Otimizador de Portfólio - Versão Desktop Completa (PyQt5)
=========================================================

Port da interface Tkinter (desktop_app_FULL.py) para PyQt5.

Toda a LÓGICA DE NEGÓCIO (cálculo de ranking, base zero, otimização,
walk-forward, métricas mensais) é idêntica ao arquivo original.
Apenas a camada de interface (widgets) foi reescrita usando PyQt5.

Dependências: PyQt5, pandas, numpy, matplotlib, scipy, openpyxl.
Também requer o módulo `optimizer.py` (classe PortfolioOptimizer),
o mesmo usado pela versão Tkinter.
"""

import os
import sys

# Garantir que o backend Qt do matplotlib use PyQt5
os.environ.setdefault("QT_API", "pyqt5")

import pandas as pd
import numpy as np
from datetime import datetime, timedelta, date

import matplotlib
matplotlib.use("QtAgg")
from matplotlib.figure import Figure
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.backends.backend_qtagg import NavigationToolbar2QT as NavigationToolbar
import matplotlib.dates as mdates

from PyQt5.QtCore import Qt, QDate, QObject, pyqtSignal
from PyQt5.QtGui import QFont, QColor
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QTabWidget, QVBoxLayout, QHBoxLayout,
    QGridLayout, QLabel, QPushButton, QCheckBox, QRadioButton, QButtonGroup,
    QLineEdit, QSlider, QGroupBox, QScrollArea, QListWidget, QAbstractItemView,
    QTableWidget, QTableWidgetItem, QHeaderView, QDialog, QProgressBar,
    QMessageBox, QFileDialog, QDateEdit, QFrame, QSizePolicy, QComboBox,
    QPlainTextEdit, QDialogButtonBox, QSpinBox
)

from optimizer import PortfolioOptimizer

# yfinance é OPCIONAL: sem ele o aplicativo funciona normalmente, apenas a
# importação pelo Yahoo Finance fica indisponível (com aviso ao usuário).
try:
    import yfinance as yf
    YFINANCE_OK = True
except Exception:
    yf = None
    YFINANCE_OK = False

# Fontes nacionais — também OPCIONAIS (cada uma tem suas dependências):
#  - B3 (COTAHIST): baixa os arquivos anuais da B3; requer 'requests'.
#  - Excel/STOCKHISTORY: fonte complementar p/ renda fixa; requer Windows +
#    Excel 365 + 'xlwings'. O módulo importa, mas a busca só roda no Windows.
# Guardamos o MOTIVO real da falha de importação (arquivo ausente, dependência
# faltando, etc.) para exibir ao usuário em vez de um palpite genérico.
try:
    import b3_series_wide_com_limpeza as b3src
    B3_OK, B3_IMPORT_ERROR = True, None
except Exception as _e:
    b3src, B3_OK = None, False
    B3_IMPORT_ERROR = f"{type(_e).__name__}: {_e}"

try:
    import b3_excel_rendafixa as rfsrc
    EXCELRF_OK, EXCELRF_IMPORT_ERROR = True, None
except Exception as _e:
    rfsrc, EXCELRF_OK = None, False
    EXCELRF_IMPORT_ERROR = f"{type(_e).__name__}: {_e}"

# Sentinela equivalente ao tk.END, usado pelo adaptador de listbox
END = "end"


# =============================================================================
# FUNÇÕES PARA RANKING DE ATIVOS  (IDÊNTICO À VERSÃO TKINTER)
# =============================================================================

def calculate_asset_ranking(df_base_zero, risk_free_column=None, peso_inc=0.33, peso_desv=0.33, peso_cor=0.33):
    """
    Calcula ranking de ativos - VERSÃO CORRIGIDA
    ATUALIZADO: Correlação agora é entre integrais (ativo vs referência)
    COM: Normalização final (0 a 1)
    """
    try:
        from scipy import stats
        import pandas as pd
        import numpy as np

        # Verificar se dados estão em base zero
        if df_base_zero is None or df_base_zero.empty:
            return None

        # Identificar colunas
        if 'Data' in df_base_zero.columns:
            df_work = df_base_zero.copy()

            # CORREÇÃO: Forçar formato brasileiro de data
            try:
                dates_col = pd.to_datetime(df_work['Data'], format='%d/%m/%Y')
            except Exception:
                dates_col = pd.to_datetime(df_work['Data'])

            # Identificar coluna de referência (taxa livre de risco)
            if risk_free_column and risk_free_column in df_work.columns:
                ref_col = risk_free_column
                asset_columns = [col for col in df_work.columns if col not in ['Data', risk_free_column]]
            elif len(df_work.columns) > 2:
                # Assumir segunda coluna como referência se contém palavras-chave
                second_col = df_work.columns[1]
                if any(term in second_col.lower() for term in ['taxa', 'livre', 'risco', 'ibov', 'ref', 'cdi', 'selic']):
                    ref_col = second_col
                    asset_columns = [col for col in df_work.columns if col not in ['Data', second_col]]
                else:
                    ref_col = df_work.columns[1]
                    asset_columns = df_work.columns[2:].tolist()
            else:
                return None
        else:
            return None

        # ===============================
        # PASSO 1: CRIAR ABA "DIFERENÇA"
        # ===============================
        diferenca_data = {}
        diferenca_data['Data'] = dates_col

        for asset in asset_columns:
            diferenca_data[f"{asset}_diff"] = df_work[asset] - df_work[ref_col]

        df_diferenca = pd.DataFrame(diferenca_data)

        # ===============================
        # PASSO 2: CRIAR ABA "INTEGRAL"
        # ===============================
        integral_data = {}
        integral_data['Data'] = dates_col

        for asset in asset_columns:
            integral_data[f"{asset}_integral"] = df_diferenca[f"{asset}_diff"].cumsum()

        df_integral = pd.DataFrame(integral_data)

        # ===============================
        # PASSO 3: CALCULAR PARÂMETROS (ÍNDICE BRUTO)
        # ===============================
        rankings = []
        all_slopes = []
        all_deviations = []

        # Primeira passada: coletar TODOS os valores
        for asset in asset_columns:
            try:
                x_data = np.arange(len(df_integral))
                y_data = df_integral[f"{asset}_integral"].values

                slope, intercept, r_value, p_value, std_err = stats.linregress(x_data, y_data)

                std_dev = df_diferenca[f"{asset}_diff"].std()

                all_slopes.append(slope)
                all_deviations.append(std_dev)

            except Exception:
                continue

        # Encontrar máximos para normalização de componentes
        max_slope = max(all_slopes) if all_slopes else 1
        max_deviation = max(all_deviations) if all_deviations else 1

        # Segunda passada: calcular índices BRUTOS
        for asset in asset_columns:
            try:
                x_data = np.arange(len(df_integral))
                y_data = df_integral[f"{asset}_integral"].values

                slope, intercept, r_value, p_value, std_err = stats.linregress(x_data, y_data)
                r_squared = r_value ** 2

                # NOVA CORRELAÇÃO: Entre integrais (evoluções acumuladas)
                asset_integral = df_work[asset].cumsum().values
                ref_integral = df_work[ref_col].cumsum().values
                correlation = np.corrcoef(asset_integral, ref_integral)[0, 1]

                std_dev = df_diferenca[f"{asset}_diff"].std()

                slope_norm = slope / max_slope if max_slope > 0 else 0
                std_dev_norm = std_dev / max_deviation if max_deviation > 0 else 0

                correlation_norm = correlation

                numerador = (peso_inc * slope_norm +
                             peso_desv * (1 - std_dev_norm) +
                             peso_cor * correlation_norm)
                denominador = peso_inc + peso_desv + peso_cor

                indice_bruto = numerador / denominador if denominador > 0 else 0

                rankings.append({
                    'Ativo': asset,
                    'Inclinação': slope,
                    'Inclinação_Norm': slope_norm,
                    'R²': r_squared,
                    'Correlação': correlation,
                    'Desvio_Padrão': std_dev,
                    'Desvio_Norm': std_dev_norm,
                    'Índice_Bruto': indice_bruto
                })

            except Exception:
                rankings.append({
                    'Ativo': asset,
                    'Inclinação': 0,
                    'Inclinação_Norm': 0,
                    'R²': 0,
                    'Correlação': 0,
                    'Desvio_Padrão': 0,
                    'Desvio_Norm': 0,
                    'Índice_Bruto': 0
                })

        # ===============================
        # PASSO 4: NORMALIZAÇÃO FINAL (0 a 1)
        # ===============================
        df_ranking = pd.DataFrame(rankings)

        if len(df_ranking) > 0:
            max_idx = df_ranking['Índice_Bruto'].max()
            min_idx = df_ranking['Índice_Bruto'].min()

            if max_idx > min_idx:
                df_ranking['Índice'] = (df_ranking['Índice_Bruto'] - min_idx) / (max_idx - min_idx)
            else:
                df_ranking['Índice'] = 0.5

            df_ranking = df_ranking.drop(columns=['Índice_Bruto'])
        else:
            df_ranking['Índice'] = 0

        df_ranking = df_ranking.sort_values('Índice', ascending=False).reset_index(drop=True)
        df_ranking['Posição'] = range(1, len(df_ranking) + 1)

        cols = ['Posição', 'Ativo', 'Índice', 'Inclinação', 'Inclinação_Norm',
                'R²', 'Correlação', 'Desvio_Padrão', 'Desvio_Norm']
        df_ranking = df_ranking[cols]

        return {
            'ranking': df_ranking,
            'referencia': ref_col,
            'total_ativos': len(asset_columns)
        }

    except Exception as e:
        print(f"❌ Erro no cálculo de ranking: {str(e)}")
        return None


# =============================================================================
# IMPORTAÇÃO DE RESTRIÇÕES (MÍN/MÁX OU PESOS) DE ARQUIVO EXCEL/CSV
# =============================================================================

def load_constraints_from_file(file_path):
    """
    Lê um arquivo Excel/CSV com restrições por ativo.

    Formatos aceitos (nomes de coluna flexíveis, sem diferenciar maiúsculas):
      • 'Ativo', 'Min', 'Max'  -> usa mínimo e máximo por ativo (faixas)
      • 'Ativo', 'Peso'        -> fixa min = max = peso
                                  (útil para ANALISAR um portfólio pronto:
                                   quando os pesos somam 100%, o otimizador
                                   é forçado exatamente àquela composição)

    Valores em PERCENTUAL (ex.: 30 = 30%). Se TODOS os valores forem <= 1,
    assume-se que já estão em fração (ex.: 0.30 = 30%).

    Retorna: (constraints, warnings)
      constraints = {ativo: {'min': fração, 'max': fração}}
      warnings    = lista de avisos (strings)
    """
    warnings = []
    ext = os.path.splitext(file_path)[1].lower()

    if ext in ('.xlsx', '.xls'):
        df = pd.read_excel(file_path)
    else:
        # CSV: detecta separador automaticamente (vírgula, ponto-e-vírgula, tab)
        df = pd.read_csv(file_path, sep=None, engine='python')

    if df is None or df.empty or len(df.columns) < 2:
        raise ValueError("Arquivo vazio ou sem colunas suficientes "
                         "(mínimo: coluna de ativo + 1 coluna de valor).")

    low = {c: str(c).strip().lower() for c in df.columns}

    def find(terms, exclude=()):
        for col, name in low.items():
            if any(t in name for t in terms) and not any(e in name for e in exclude):
                return col
        return None

    asset_col = find(['ativo', 'asset', 'ticker', 'papel', 'codigo', 'código', 'symbol', 'nome'])
    min_col = find(['min', 'mín'])
    max_col = find(['max', 'máx'])
    weight_col = find(['peso', 'weight', 'quant', 'aloca', 'aloc', 'percent', 'participa'])

    if asset_col is None:
        asset_col = df.columns[0]

    def to_float(v):
        if pd.isna(v):
            return None
        s = str(v).strip().replace('%', '').replace(' ', '')
        if s == '':
            return None
        if ',' in s and '.' in s:
            s = s.replace('.', '').replace(',', '.')   # 1.234,56 -> 1234.56
        elif ',' in s:
            s = s.replace(',', '.')                     # 30,5 -> 30.5
        try:
            return float(s)
        except ValueError:
            return None

    raw = {}
    if min_col is not None and max_col is not None:
        for _, row in df.iterrows():
            asset = str(row[asset_col]).strip()
            if not asset or asset.lower() == 'nan':
                continue
            mn = to_float(row[min_col])
            mx = to_float(row[max_col])
            if mn is None or mx is None:
                warnings.append(f"'{asset}': valor mín/máx inválido, ignorado.")
                continue
            raw[asset] = [mn, mx]
    else:
        # Coluna única de peso -> min = max = peso
        if weight_col is None:
            weight_col = df.columns[1]   # assume 2ª coluna como peso
        for _, row in df.iterrows():
            asset = str(row[asset_col]).strip()
            if not asset or asset.lower() == 'nan':
                continue
            w = to_float(row[weight_col])
            if w is None:
                warnings.append(f"'{asset}': peso inválido, ignorado.")
                continue
            raw[asset] = [w, w]

    if not raw:
        raise ValueError("Nenhum ativo válido encontrado no arquivo.")

    # Percentual (0-100) vs fração (0-1): decide pela maior magnitude presente
    all_vals = [abs(v) for pair in raw.values() for v in pair]
    max_val = max(all_vals) if all_vals else 0
    scale = 1.0 if max_val <= 1.0 else 0.01

    constraints = {}
    for asset, (mn, mx) in raw.items():
        mn *= scale
        mx *= scale
        if mn > mx:
            mn, mx = mx, mn
            warnings.append(f"'{asset}': mín > máx, valores invertidos.")
        constraints[asset] = {'min': max(0.0, mn), 'max': max(0.0, mx)}

    return constraints, warnings


# =============================================================================
# ADAPTADORES / HELPERS PARA APROXIMAR A API DO TKINTER
# =============================================================================

class BoolVar:
    """Adaptador de tk.BooleanVar sobre um QCheckBox."""
    def __init__(self, checkbox):
        self.w = checkbox

    def get(self):
        return self.w.isChecked()

    def set(self, value):
        self.w.setChecked(bool(value))


class NumVar:
    """Adaptador de tk.DoubleVar/IntVar sobre um QLineEdit."""
    def __init__(self, lineedit, value=0.0):
        self.w = lineedit
        if value is not None:
            self.w.setText(self._fmt(value))

    @staticmethod
    def _fmt(v):
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)

    def get(self):
        text = self.w.text().strip().replace(',', '.')
        try:
            return float(text)
        except Exception:
            return 0.0

    def set(self, value):
        self.w.setText(self._fmt(value))


class StrVar:
    """Adaptador de tk.StringVar sobre um QLineEdit."""
    def __init__(self, lineedit):
        self.w = lineedit

    def get(self):
        return self.w.text()

    def set(self, value):
        self.w.setText(str(value))


class SliderVar:
    """Adaptador de tk.DoubleVar sobre um QSlider (valores fracionários)."""
    def __init__(self, slider, factor):
        self.s = slider
        self.f = factor

    def get(self):
        return self.s.value() / self.f

    def set(self, value):
        self.s.setValue(int(round(value * self.f)))


class RadioVar:
    """Adaptador de tk.StringVar para grupos de QRadioButton."""
    def __init__(self, default=None):
        self.buttons = {}
        self._default = default

    def add(self, text, button):
        self.buttons[text] = button

    def get(self):
        for text, btn in self.buttons.items():
            if btn.isChecked():
                return text
        return self._default

    def set(self, text):
        if text in self.buttons:
            self.buttons[text].setChecked(True)


class DateEntry(QDateEdit):
    """Substituto de tkcalendar.DateEntry usando QDateEdit."""
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setCalendarPopup(True)
        self.setDisplayFormat("dd/MM/yyyy")

    def get(self):
        return self.date().toString("dd/MM/yyyy")

    def get_date(self):
        return self.date().toPyDate()

    def set_date(self, d):
        if isinstance(d, (datetime, pd.Timestamp)):
            d = d.date() if hasattr(d, "date") else d
        self.setDate(QDate(d.year, d.month, d.day))


class AssetListBox(QListWidget):
    """Adaptador de tk.Listbox (seleção múltipla) sobre QListWidget."""
    def __init__(self):
        super().__init__()
        self.setSelectionMode(QAbstractItemView.ExtendedSelection)

    def delete(self, *args):
        self.clear()

    def insert(self, index, text):
        self.addItem(str(text))

    def get(self, a=None, b=None):
        if b is not None:
            return [self.item(i).text() for i in range(self.count())]
        return self.item(a).text()

    def size(self):
        return self.count()

    def curselection(self):
        return sorted(self.row(it) for it in self.selectedItems())

    def selection_set(self, i):
        item = self.item(i)
        if item is not None:
            item.setSelected(True)

    def select_set(self, *args):
        self.selectAll()

    def selection_clear(self, *args):
        self.clearSelection()


# ---- Shims para messagebox / filedialog (mesma assinatura do tkinter) --------

def _active():
    return QApplication.activeWindow()


class _MessageBox:
    @staticmethod
    def showerror(title, msg):
        QMessageBox.critical(_active(), str(title), str(msg))

    @staticmethod
    def showinfo(title, msg):
        QMessageBox.information(_active(), str(title), str(msg))

    @staticmethod
    def showwarning(title, msg):
        QMessageBox.warning(_active(), str(title), str(msg))


messagebox = _MessageBox()


def _convert_filters(filetypes):
    """Converte filetypes do tkinter [(label, '*.ext ...')] para o formato Qt."""
    if not filetypes:
        return ""
    parts = []
    for label, patterns in filetypes:
        parts.append(f"{label} ({patterns})")
    return ";;".join(parts)


class _FileDialog:
    @staticmethod
    def askopenfilename(title="", filetypes=None):
        fn, _ = QFileDialog.getOpenFileName(_active(), title, "", _convert_filters(filetypes))
        return fn

    @staticmethod
    def asksaveasfilename(title="", defaultextension="", filetypes=None):
        fn, _ = QFileDialog.getSaveFileName(_active(), title, "", _convert_filters(filetypes))
        if fn and defaultextension and not os.path.splitext(fn)[1]:
            fn += defaultextension
        return fn


filedialog = _FileDialog()


def clear_widget(widget):
    """Remove todos os widgets/layouts filhos do layout de `widget`."""
    layout = widget.layout()
    if layout is not None:
        _clear_layout(layout)


def _clear_layout(layout):
    while layout.count():
        item = layout.takeAt(0)
        child = item.widget()
        if child is not None:
            child.setParent(None)
            child.deleteLater()
        else:
            sub = item.layout()
            if sub is not None:
                _clear_layout(sub)


def bold(size=None):
    f = QFont()
    f.setBold(True)
    if size:
        f.setPointSize(size)
    return f


# Sinais para comunicação segura thread -> UI (auto-otimização)
class _AutoSignals(QObject):
    status = pyqtSignal(str, str)
    results_ready = pyqtSignal(object)
    finished = pyqtSignal()


# =============================================================================
# IMPORTAÇÃO DE DADOS PELO YAHOO FINANCE
# =============================================================================

# Tipos de ativo e o sufixo aplicado ao código digitado.
# "LIVRE" é um marcador: nesse modo o código vai exatamente como digitado.
YAHOO_ASSET_TYPES = (
    ("Ações Brasileiras (.SA)", ".SA"),
    ("Ações Americanas", ""),
    ("ETFs Americanos", ""),
    ("Criptomoedas", ""),
    ("Códigos Livres do Yahoo", "LIVRE"),
)


def yahoo_symbol(simbolo, sufixo):
    """Monta o código final enviado ao Yahoo a partir do que foi digitado."""
    simbolo = simbolo.strip()
    if sufixo in ("", None, "LIVRE"):
        return simbolo          # modo livre / mercados sem sufixo
    if "." in simbolo:
        return simbolo          # já veio com sufixo: respeita o que o usuário digitou
    return simbolo + sufixo


def fetch_yahoo_prices(simbolos, data_inicio, data_fim, sufixo=".SA", progress=None):
    """
    Baixa o histórico diário de cada símbolo no Yahoo Finance.

    progress: callable(indice, total, simbolo) chamado antes de cada busca,
    para a interface poder atualizar a barra de progresso.

    Retorna (dados_por_simbolo, erros) — as chaves usam o código ORIGINAL
    digitado, para o restante do fluxo não depender do sufixo.
    """
    if not YFINANCE_OK:
        raise RuntimeError(
            "A biblioteca 'yfinance' não está instalada.\n\n"
            "Instale com:  pip install yfinance"
        )

    dados, erros = {}, []
    inicio = data_inicio.strftime('%Y-%m-%d')
    fim = data_fim.strftime('%Y-%m-%d')
    total = len(simbolos)

    for i, simbolo in enumerate(simbolos):
        if progress is not None:
            progress(i, total, simbolo)
        try:
            hist = yf.Ticker(yahoo_symbol(simbolo, sufixo)).history(
                start=inicio, end=fim, interval="1d")
            # Exige um mínimo de pontos: séries muito curtas quebram a base zero
            if hist is not None and not hist.empty and len(hist) > 5:
                dados[simbolo] = hist
            else:
                erros.append(simbolo)
        except Exception:
            erros.append(simbolo)

    if progress is not None:
        progress(total, total, "")
    return dados, erros


def consolidate_yahoo_prices(dados_historicos):
    """Junta os fechamentos ('Close') num único DataFrame indexado por data."""
    if not dados_historicos:
        return None

    series = []
    for simbolo, hist in dados_historicos.items():
        if 'Close' in hist.columns and not hist['Close'].empty:
            df_temp = pd.DataFrame({simbolo: hist['Close']})
            # Yahoo devolve índice com fuso; remover evita conflito ao mesclar
            if getattr(df_temp.index, 'tz', None) is not None:
                df_temp.index = df_temp.index.tz_localize(None)
            series.append(df_temp)

    if not series:
        return None

    consolidado = pd.concat(series, axis=1, sort=True)
    consolidado.index.name = "Data"
    return consolidado


def promote_reference_column(df, ativo_referencia):
    """
    Renomeia a coluna do ativo de referência para 'Taxa_Ref_<CÓDIGO>' e a move
    para a segunda posição (logo após 'Data'). É esse prefixo + posição que faz
    a detecção automática reconhecer a taxa de referência e habilitar os
    objetivos de excesso. Retorna (df, nome_da_coluna) — nome None se o ativo
    não estiver presente. Usada por todas as fontes online (Yahoo, B3, Excel).
    """
    if not ativo_referencia:
        return df, None
    ref = ativo_referencia.strip().upper()
    if ref not in df.columns:
        return df, None
    nome_ref = f"Taxa_Ref_{ref}"
    df = df.rename(columns={ref: nome_ref})
    outras = [c for c in df.columns if c not in ('Data', nome_ref)]
    return df[['Data', nome_ref] + outras], nome_ref


def build_yahoo_dataframe(dados_historicos, ativo_referencia=None):
    """
    Monta o DataFrame no layout que o aplicativo espera:
    Data | Taxa_Ref_<ATIVO> (opcional) | demais ativos...
    """
    consolidado = consolidate_yahoo_prices(dados_historicos)
    if consolidado is None:
        return None, None
    df = consolidado.reset_index()
    return promote_reference_column(df, ativo_referencia)


class _SortableItem(QTableWidgetItem):
    """Item de tabela que ordena numericamente quando há uma chave numérica
    armazenada em Qt.UserRole; caso contrário, ordena como texto."""
    def __lt__(self, other):
        a = self.data(Qt.UserRole)
        b = other.data(Qt.UserRole)
        if a is not None and b is not None:
            try:
                return float(a) < float(b)
            except (TypeError, ValueError):
                pass
        return self.text() < other.text()


class YahooImportDialog(QDialog):
    """
    Diálogo de importação de cotações do Yahoo Finance.

    Ao ser aceito, expõe:
      result_df  -> DataFrame pronto (Data | Taxa_Ref_X | ativos...)
      nome_ref   -> nome da coluna de referência criada (ou None)
      erros      -> símbolos sem dados suficientes
    """

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("🌐 Importar do Yahoo Finance")
        self.setMinimumSize(600, 660)

        self.result_df = None
        self.nome_ref = None
        self.erros = []
        self.origem_desc = "Yahoo Finance"

        layout = QVBoxLayout(self)

        # ----- Símbolos -----
        sym_box = QGroupBox("📝 Símbolos dos Ativos (um por linha)")
        sym_l = QVBoxLayout(sym_box)
        self.symbols_edit = QPlainTextEdit()
        self.symbols_edit.setPlainText("PETR4\nVALE3\nITUB4\nBBDC4\nABEV3")
        self.symbols_edit.setFixedHeight(120)
        sym_l.addWidget(self.symbols_edit)
        layout.addWidget(sym_box)

        # ----- Tipo de ativo -----
        tipo_row = QHBoxLayout()
        tipo_row.addWidget(QLabel("🏷️ Tipo de ativo:"))
        self.tipo_combo = QComboBox()
        for nome, _ in YAHOO_ASSET_TYPES:
            self.tipo_combo.addItem(nome)
        self.tipo_combo.currentIndexChanged.connect(self._update_tipo_hint)
        tipo_row.addWidget(self.tipo_combo, 1)
        layout.addLayout(tipo_row)

        self.tipo_hint = QLabel("")
        self.tipo_hint.setStyleSheet("color: gray;")
        self.tipo_hint.setWordWrap(True)
        layout.addWidget(self.tipo_hint)

        # ----- Ativo de referência -----
        ref_box = QGroupBox("🏛️ Ativo de Referência (benchmark / taxa livre)")
        ref_l = QVBoxLayout(ref_box)
        ref_row = QHBoxLayout()
        chk_ref = QCheckBox("Incluir")
        chk_ref.setChecked(True)
        self.use_ref = BoolVar(chk_ref)
        ref_row.addWidget(chk_ref)
        self.ref_entry = QLineEdit("BOVA11")
        self.ref_entry.setFixedWidth(140)
        ref_row.addWidget(self.ref_entry)
        ref_row.addStretch()
        ref_l.addLayout(ref_row)
        hint_ref = QLabel(
            "Sugestões: BOVA11 (Ibovespa), LFTS11 (CDI), SMAL11 (Small Caps), IVV (S&P 500).\n"
            "O ativo escolhido vira a coluna de referência e é detectado automaticamente.")
        hint_ref.setStyleSheet("color: gray;")
        hint_ref.setWordWrap(True)
        ref_l.addWidget(hint_ref)
        layout.addWidget(ref_box)

        # ----- Período -----
        per_box = QGroupBox("📅 Período")
        per_l = QHBoxLayout(per_box)
        per_l.addWidget(QLabel("Início:"))
        self.date_ini = DateEntry()
        self.date_ini.set_date(date.today() - timedelta(days=365 * 3))
        self.date_ini.setMaximumDate(QDate.currentDate())
        per_l.addWidget(self.date_ini)
        per_l.addSpacing(20)
        per_l.addWidget(QLabel("Fim:"))
        self.date_fim = DateEntry()
        self.date_fim.set_date(date.today())
        self.date_fim.setMaximumDate(QDate.currentDate())
        per_l.addWidget(self.date_fim)
        per_l.addStretch()
        layout.addWidget(per_box)

        # ----- Progresso -----
        self.progress = QProgressBar()
        self.progress.setVisible(False)
        layout.addWidget(self.progress)
        self.status = QLabel("")
        self.status.setWordWrap(True)
        layout.addWidget(self.status)

        # ----- Botões -----
        btns = QHBoxLayout()
        self.btn_fetch = QPushButton("🚀 Buscar e Importar")
        self.btn_fetch.setStyleSheet(
            "QPushButton { background-color: #0078d4; color: white; font-weight: bold; padding: 8px; }")
        self.btn_fetch.clicked.connect(self._fetch)
        btns.addWidget(self.btn_fetch)
        btn_cancel = QPushButton("Cancelar")
        btn_cancel.clicked.connect(self.reject)
        btns.addWidget(btn_cancel)
        layout.addLayout(btns)

        self._update_tipo_hint()

    def _sufixo(self):
        return YAHOO_ASSET_TYPES[self.tipo_combo.currentIndex()][1]

    def _update_tipo_hint(self):
        sufixo = self._sufixo()
        if sufixo == "LIVRE":
            self.tipo_hint.setText(
                "🔥 Modo códigos livres: digite exatamente como aparece no Yahoo "
                "(ex.: PETR4.SA, MSFT, BTC-USD). Nenhum sufixo é acrescentado.")
        elif sufixo:
            self.tipo_hint.setText(
                f"O sufixo '{sufixo}' é acrescentado automaticamente (ex.: PETR4 → PETR4{sufixo}). "
                "Códigos que já contenham ponto são usados como digitados.")
        else:
            self.tipo_hint.setText(
                "Os códigos são usados como digitados (ex.: MSFT, AAPL, BTC-USD).")

    def _fetch(self):
        simbolos = [s.strip().upper() for s in self.symbols_edit.toPlainText().split('\n') if s.strip()]
        if len(simbolos) < 2:
            messagebox.showerror("Erro", "Digite pelo menos 2 símbolos.")
            return

        d_ini, d_fim = self.date_ini.get_date(), self.date_fim.get_date()
        if d_ini >= d_fim:
            messagebox.showerror("Erro", "A data de início deve ser anterior à data de fim.")
            return

        ativo_ref = self.ref_entry.text().strip().upper() if self.use_ref.get() else None

        # O ativo de referência entra na mesma busca dos demais
        busca = list(simbolos)
        if ativo_ref and ativo_ref not in busca:
            busca.append(ativo_ref)

        self.btn_fetch.setEnabled(False)
        self.progress.setVisible(True)
        self.progress.setMaximum(len(busca))

        def _progress(i, total, simbolo):
            self.progress.setValue(i)
            self.status.setText(f"Buscando {simbolo}... ({min(i + 1, total)}/{total})" if simbolo else "Consolidando...")
            QApplication.processEvents()

        try:
            dados, erros = fetch_yahoo_prices(busca, d_ini, d_fim,
                                              sufixo=self._sufixo(), progress=_progress)
        except Exception as e:
            self.progress.setVisible(False)
            self.btn_fetch.setEnabled(True)
            self.status.setText("")
            messagebox.showerror("Erro", f"Falha ao consultar o Yahoo Finance:\n{str(e)}")
            return

        self.progress.setVisible(False)
        self.btn_fetch.setEnabled(True)
        self.status.setText("")

        if not dados:
            messagebox.showerror(
                "Erro",
                "Nenhum dado encontrado.\n\nVerifique os símbolos e o tipo de ativo "
                "(ações brasileiras normalmente exigem o sufixo .SA).")
            return

        # Se a referência falhou, avisa: sem ela os objetivos de excesso somem
        if ativo_ref and ativo_ref not in dados:
            messagebox.showwarning(
                "Referência não encontrada",
                f"Não foi possível obter dados de '{ativo_ref}'.\n\n"
                "A importação segue sem taxa de referência — os objetivos que "
                "dependem dela ficarão indisponíveis.")
            ativo_ref = None

        df, nome_ref = build_yahoo_dataframe(dados, ativo_ref)
        if df is None or df.empty:
            messagebox.showerror("Erro", "Não foi possível consolidar os preços obtidos.")
            return

        n_ativos = len(df.columns) - 1
        self.result_df = df
        self.nome_ref = nome_ref
        self.erros = erros
        self.origem_desc = f"Yahoo Finance ({n_ativos} séries)"
        self.accept()


class _PriceImportDialog(QDialog):
    """
    Base dos diálogos de importação por PREÇO (B3 e Excel/renda fixa), que
    compartilham quase toda a interface: símbolos, tipo de preço, limpeza (k),
    ativo de referência, período e barra de progresso.

    A subclasse só precisa implementar _run_fetch(...), devolvendo
    (wide_df, faltantes, relatorio) — o resto do fluxo (validação, promoção da
    referência, empacotamento do resultado) fica aqui.
    """
    progress_label = "Buscando"

    def __init__(self, parent=None, titulo="Importar", banner=None):
        super().__init__(parent)
        self.setWindowTitle(titulo)
        self.setMinimumSize(600, 680)

        self.result_df = None
        self.nome_ref = None
        self.erros = []
        self.relatorio = None
        self.origem_desc = titulo

        layout = QVBoxLayout(self)

        if banner:
            lbl_banner = QLabel(banner)
            lbl_banner.setWordWrap(True)
            lbl_banner.setStyleSheet(
                "background:#fff3cd; color:#664d03; padding:8px; border-radius:4px;")
            layout.addWidget(lbl_banner)

        # ----- Símbolos -----
        sym_box = QGroupBox("📝 Símbolos dos Ativos (um por linha)")
        sym_l = QVBoxLayout(sym_box)
        self.symbols_edit = QPlainTextEdit()
        self.symbols_edit.setPlainText(self.default_symbols())
        self.symbols_edit.setFixedHeight(110)
        sym_l.addWidget(self.symbols_edit)
        layout.addWidget(sym_box)

        # ----- Tipo de preço + limpeza -----
        opts = QHBoxLayout()
        opts.addWidget(QLabel("💰 Preço:"))
        self.preco_combo = QComboBox()
        for nome in ("Abertura", "Máximo", "Mínimo", "Fechamento"):
            self.preco_combo.addItem(nome)
        self.preco_combo.setCurrentText("Fechamento")
        opts.addWidget(self.preco_combo)
        opts.addSpacing(20)
        opts.addWidget(QLabel("🧹 Elimina após N dias sem dado:"))
        self.k_spin = QSpinBox()
        self.k_spin.setRange(1, 999)
        self.k_spin.setValue(10)
        self.k_spin.setToolTip(
            "Se um ativo ficar mais de N pregões seguidos sem cotação no período, "
            "ele é descartado. Lacunas menores são preenchidas com o dia anterior.")
        opts.addWidget(self.k_spin)
        opts.addStretch()
        layout.addLayout(opts)

        # ----- Ativo de referência -----
        ref_box = QGroupBox("🏛️ Ativo de Referência (benchmark / taxa livre)")
        ref_l = QVBoxLayout(ref_box)
        ref_row = QHBoxLayout()
        chk_ref = QCheckBox("Incluir")
        chk_ref.setChecked(True)
        self.use_ref = BoolVar(chk_ref)
        ref_row.addWidget(chk_ref)
        self.ref_entry = QLineEdit(self.default_reference())
        self.ref_entry.setFixedWidth(140)
        ref_row.addWidget(self.ref_entry)
        ref_row.addStretch()
        ref_l.addLayout(ref_row)
        hint = QLabel(
            "O ativo escolhido é buscado junto, vira a coluna de referência "
            "(Taxa_Ref_<código>) e habilita os objetivos de excesso.")
        hint.setStyleSheet("color: gray;")
        hint.setWordWrap(True)
        ref_l.addWidget(hint)
        layout.addWidget(ref_box)

        # ----- Período -----
        per_box = QGroupBox("📅 Período")
        per_l = QHBoxLayout(per_box)
        per_l.addWidget(QLabel("Início:"))
        self.date_ini = DateEntry()
        self.date_ini.set_date(date.today() - timedelta(days=365 * 3))
        self.date_ini.setMaximumDate(QDate.currentDate())
        per_l.addWidget(self.date_ini)
        per_l.addSpacing(20)
        per_l.addWidget(QLabel("Fim:"))
        self.date_fim = DateEntry()
        self.date_fim.set_date(date.today())
        self.date_fim.setMaximumDate(QDate.currentDate())
        per_l.addWidget(self.date_fim)
        per_l.addStretch()
        layout.addWidget(per_box)

        # ----- Progresso -----
        self.progress = QProgressBar()
        self.progress.setVisible(False)
        layout.addWidget(self.progress)
        self.status = QLabel("")
        self.status.setWordWrap(True)
        layout.addWidget(self.status)

        # ----- Botões -----
        btns = QHBoxLayout()
        self.btn_fetch = QPushButton("🚀 Buscar e Importar")
        self.btn_fetch.setStyleSheet(
            "QPushButton { background-color: #0078d4; color: white; font-weight: bold; padding: 8px; }")
        self.btn_fetch.clicked.connect(self._fetch)
        btns.addWidget(self.btn_fetch)
        btn_cancel = QPushButton("Cancelar")
        btn_cancel.clicked.connect(self.reject)
        btns.addWidget(btn_cancel)
        layout.addLayout(btns)

    # ---- pontos de personalização das subclasses ----
    def default_symbols(self):
        return "PETR4\nVALE3\nITUB4\nBBDC4\nABEV3"

    def default_reference(self):
        return "BOVA11"

    def _run_fetch(self, simbolos, d_ini, d_fim, tipo_preco, k, progress):
        """Deve devolver (wide_df, faltantes, relatorio). Implementada na subclasse."""
        raise NotImplementedError

    # ---- fluxo comum ----
    def _fetch(self):
        simbolos = [s.strip().upper() for s in self.symbols_edit.toPlainText().split('\n') if s.strip()]
        if len(simbolos) < 2:
            messagebox.showerror("Erro", "Digite pelo menos 2 símbolos.")
            return

        d_ini = datetime.combine(self.date_ini.get_date(), datetime.min.time())
        d_fim = datetime.combine(self.date_fim.get_date(), datetime.min.time())
        if d_ini >= d_fim:
            messagebox.showerror("Erro", "A data de início deve ser anterior à data de fim.")
            return

        ativo_ref = self.ref_entry.text().strip().upper() if self.use_ref.get() else None
        busca = list(simbolos)
        if ativo_ref and ativo_ref not in busca:
            busca.append(ativo_ref)

        tipo_preco = self.preco_combo.currentText()
        k = self.k_spin.value()

        self.btn_fetch.setEnabled(False)
        self.progress.setVisible(True)
        self.progress.setMaximum(0)  # indeterminado até o 1º progresso

        def _progress(i, total, item):
            self.progress.setMaximum(max(int(total), 1))
            self.progress.setValue(int(i))
            if item is not None:
                self.status.setText(f"{self.progress_label}: {item} ({min(i + 1, total)}/{total})")
            else:
                self.status.setText("Processando...")
            QApplication.processEvents()

        try:
            wide, faltantes, relatorio = self._run_fetch(busca, d_ini, d_fim, tipo_preco, k, _progress)
        except Exception as e:
            self.progress.setVisible(False)
            self.btn_fetch.setEnabled(True)
            self.status.setText("")
            messagebox.showerror("Erro", f"Falha na importação:\n{str(e)}")
            return

        self.progress.setVisible(False)
        self.btn_fetch.setEnabled(True)
        self.status.setText("")

        if wide is None or wide.empty or (len(wide.columns) <= 1):
            messagebox.showerror(
                "Erro", "Nenhum dado utilizável foi obtido.\nVerifique os símbolos e o período.")
            return

        # Se a referência não sobreviveu (não veio, ou saiu na limpeza), avisa
        if ativo_ref and ativo_ref not in wide.columns:
            messagebox.showwarning(
                "Referência indisponível",
                f"'{ativo_ref}' não retornou dados suficientes.\n\n"
                "A importação segue sem taxa de referência — os objetivos que "
                "dependem dela ficarão indisponíveis.")
            ativo_ref = None

        wide, nome_ref = promote_reference_column(wide, ativo_ref)

        n_ativos = len(wide.columns) - 1
        self.result_df = wide
        self.nome_ref = nome_ref
        self.erros = list(faltantes or [])
        self.relatorio = relatorio
        self.origem_desc = f"{self.origem_desc} ({n_ativos} séries)"
        self.accept()


class B3ImportDialog(_PriceImportDialog):
    """Importa cotações da B3 (arquivos COTAHIST anuais)."""

    def __init__(self, parent=None):
        super().__init__(
            parent,
            titulo="🇧🇷 Importar da B3 (COTAHIST)",
            banner="Baixa os arquivos COTAHIST anuais da B3. O primeiro download de "
                   "cada ano é grande (dezenas de MB) e pode levar de alguns segundos "
                   "a minutos — a janela pode parecer parada durante cada ano.")
        self.progress_label = "Baixando COTAHIST"

    def _run_fetch(self, simbolos, d_ini, d_fim, tipo_preco, k, progress):
        df = b3src.baixar_periodo(d_ini, d_fim, set(simbolos),
                                  somente_vista=b3src.SOMENTE_VISTA, progress=progress)
        wide, faltantes = b3src.montar_planilha(
            df, simbolos, tipo_preco, d_ini, d_fim, somente_vista=b3src.SOMENTE_VISTA)
        wide, relatorio = b3src.limpar_dados(wide, k=k)
        return wide, faltantes, relatorio


class ExcelRFImportDialog(_PriceImportDialog):
    """Importa cotações via Excel/STOCKHISTORY — complementar (renda fixa)."""

    def __init__(self, parent=None):
        super().__init__(
            parent,
            titulo="📊 Importar via Excel (renda fixa)",
            banner="Fonte COMPLEMENTAR (LSEG/Refinitiv via STOCKHISTORY do Excel), para "
                   "ETFs que o COTAHIST não cobre bem (FIXA11, IMAB11, B5P211...). "
                   "Requer Windows com Excel 365 instalado e logado, e a biblioteca "
                   "'xlwings'. Não misture esta fonte com a B3 na mesma carteira.")
        self.progress_label = "Excel: buscando"

    def default_symbols(self):
        return "FIXA11\nIMAB11\nB5P211\nIRFM11"

    def default_reference(self):
        return "LFTS11"

    def _run_fetch(self, simbolos, d_ini, d_fim, tipo_preco, k, progress):
        wide, faltantes = rfsrc.fetch_excel_wide(
            simbolos, d_ini, d_fim, tipo_preco, visivel=False, progress=progress)
        wide, relatorio = rfsrc.limpar_dados(wide, k=k)
        return wide, faltantes, relatorio


# =============================================================================
# JANELA PRINCIPAL
# =============================================================================

class PortfolioOptimizerGUI(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("📊 Otimizador de Portfólio - Versão Desktop Completa")
        self.resize(1400, 900)

        # Variáveis de dados
        self.df = None
        self.optimizer = None
        self.result = None
        self.has_risk_free = False
        self.risk_free_column_name = None
        self.detected_risk_free_rate = 0.0

        # Variáveis para restrições individuais
        self.individual_constraints = {}
        self.constraint_widgets = {}

        # Variáveis para short selling
        self.short_weights = {}
        self.short_widgets = {}
        self.auto_short_weights = {}   # shorts fixos da Auto-Otimização

        # Tabelas mensais
        self.monthly_table = None
        self.excess_table = None

        # Janelas temporais
        self.dados_brutos = None
        self.periodo_disponivel = None
        self.df_otimizacao = None
        self.df_analise = None
        self.periodo_otimizacao = None
        self.periodo_analise = None

        # Sinais para thread de auto-otimização
        self.auto_signals = _AutoSignals()
        self.auto_signals.status.connect(self._set_auto_status)
        self.auto_signals.results_ready.connect(self.process_and_display_results)
        self.auto_signals.finished.connect(self._auto_finished)

        # Grupo exclusivo para radios de objetivo
        self.objective_group = QButtonGroup(self)
        self.objective_group.setExclusive(True)

        self.setup_ui()

    # -------------------------------------------------------------------------
    # Construção da interface
    # -------------------------------------------------------------------------
    def setup_ui(self):
        self.notebook = QTabWidget()
        self.setCentralWidget(self.notebook)

        self.tab_data = QWidget()
        self.notebook.addTab(self.tab_data, "📁 Dados")
        self.setup_data_tab()

        self.tab_config = QWidget()
        self.notebook.addTab(self.tab_config, "⚙️ Configuração")
        self.setup_config_tab()

        self.tab_advanced = QWidget()
        self.notebook.addTab(self.tab_advanced, "🔧 Avançado")
        self.setup_advanced_tab()

        self.tab_short = QWidget()
        self.notebook.addTab(self.tab_short, "🔄 Short/Hedge")
        self.setup_short_tab()

        self.tab_ranking = QWidget()
        self.notebook.addTab(self.tab_ranking, "🏆 Ranking")
        self.setup_ranking_tab()

        self.tab_results = QWidget()
        self.notebook.addTab(self.tab_results, "📈 Resultados")
        self.setup_results_tab()

        self.tab_monthly = QWidget()
        self.notebook.addTab(self.tab_monthly, "📅 Retornos Mensais")
        self.setup_monthly_tab()

        self.tab_auto = QWidget()
        self.notebook.addTab(self.tab_auto, "🤖 Auto-Otimização")
        self.setup_auto_optimization_tab()

    # ---- Helpers de layout ----

    @staticmethod
    def _scroll_area(parent_widget):
        """Cria (scroll, container, layout) e insere o scroll no parent_widget."""
        outer = QVBoxLayout(parent_widget)
        outer.setContentsMargins(0, 0, 0, 0)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        container = QWidget()
        layout = QVBoxLayout(container)
        scroll.setWidget(container)
        outer.addWidget(scroll)
        return scroll, container, layout

    def _add_slider(self, layout, caption, lo, hi, initial, decimals, suffix="%"):
        row = QHBoxLayout()
        row.addWidget(QLabel(caption))
        row.addStretch()
        fmt = f"{{:.{decimals}f}}{suffix}"
        val_label = QLabel(fmt.format(initial))
        row.addWidget(val_label)
        layout.addLayout(row)

        slider = QSlider(Qt.Horizontal)
        factor = 10 ** decimals
        slider.setRange(int(lo * factor), int(hi * factor))
        slider.setValue(int(round(initial * factor)))
        slider.valueChanged.connect(lambda v: val_label.setText(fmt.format(v / factor)))
        layout.addWidget(slider)
        return SliderVar(slider, factor)

    # -------------------------------------------------------------------------
    # ABA 1: DADOS
    # -------------------------------------------------------------------------
    def setup_data_tab(self):
        root = QHBoxLayout(self.tab_data)

        left = QVBoxLayout()
        right = QVBoxLayout()
        root.addLayout(left, 1)
        root.addLayout(right, 1)

        # ----- Carregar Dados -----
        load_box = QGroupBox("Carregar Dados")
        load_l = QVBoxLayout(load_box)
        btn_load = QPushButton("📂 Carregar Planilha Excel")
        btn_load.clicked.connect(self.load_excel_file)
        load_l.addWidget(btn_load)

        btn_yahoo = QPushButton("🌐 Importar do Yahoo Finance")
        btn_yahoo.clicked.connect(self.open_yahoo_import)
        if not YFINANCE_OK:
            btn_yahoo.setToolTip("Requer a biblioteca 'yfinance' (pip install yfinance)")
        load_l.addWidget(btn_yahoo)

        btn_b3 = QPushButton("🇧🇷 Importar da B3 (COTAHIST)")
        btn_b3.clicked.connect(self.open_b3_import)
        if not B3_OK:
            btn_b3.setToolTip("Requer a biblioteca 'requests' (pip install requests)")
        load_l.addWidget(btn_b3)

        btn_rf = QPushButton("📊 Importar via Excel (renda fixa)")
        btn_rf.clicked.connect(self.open_excel_rf_import)
        btn_rf.setToolTip("Fonte complementar: requer Windows + Excel 365 + xlwings")
        load_l.addWidget(btn_rf)

        left.addWidget(load_box)

        # ----- Informações do Arquivo -----
        info_box = QGroupBox("Informações do Arquivo")
        info_l = QVBoxLayout(info_box)
        self.status_label = QLabel("Nenhum arquivo carregado")
        self.status_label.setWordWrap(True)
        info_l.addWidget(self.status_label)

        rf_box = QGroupBox("Taxa de Referência Detectada")
        rf_l = QVBoxLayout(rf_box)
        self.risk_free_info = QLabel("Nenhuma taxa detectada")
        rf_l.addWidget(self.risk_free_info)
        info_l.addWidget(rf_box)
        left.addWidget(info_box)

        # ----- Janelas Temporais -----
        temp_box = QGroupBox("📅 Configurar Janelas Temporais")
        temp_l = QVBoxLayout(temp_box)

        instr = QLabel("🎯 Configure as 3 datas críticas para análise:")
        instr.setFont(bold())
        temp_l.addWidget(instr)

        # Data 1
        r1 = QHBoxLayout()
        r1.addWidget(QLabel("📊 Início da Otimização:"))
        r1.addStretch()
        self.data_inicio_otim = DateEntry()
        r1.addWidget(self.data_inicio_otim)
        temp_l.addLayout(r1)

        # Data 2
        r2 = QHBoxLayout()
        r2.addWidget(QLabel("🎯 Fim da Otimização:"))
        r2.addStretch()
        self.data_fim_otim = DateEntry()
        r2.addWidget(self.data_fim_otim)
        temp_l.addLayout(r2)

        # Data 3
        r3 = QHBoxLayout()
        r3.addWidget(QLabel("📈 Fim da Análise:"))
        r3.addStretch()
        self.data_fim_analise = DateEntry()
        r3.addWidget(self.data_fim_analise)
        temp_l.addLayout(r3)

        hoje = datetime.now().date()
        self.data_inicio_otim.set_date(hoje - timedelta(days=365 * 2))
        self.data_fim_otim.set_date(hoje - timedelta(days=365))
        self.data_fim_analise.set_date(hoje)

        # Checkbox validação
        chk_valid = QCheckBox("✅ Usar validação (forward test)")
        chk_valid.setChecked(True)
        self.usar_validacao = BoolVar(chk_valid)
        chk_valid.toggled.connect(self.toggle_validacao)
        temp_l.addWidget(chk_valid)

        # Conectar mudança de datas
        self.data_inicio_otim.dateChanged.connect(self.atualizar_metricas_janelas)
        self.data_fim_otim.dateChanged.connect(self.atualizar_metricas_janelas)
        self.data_fim_analise.dateChanged.connect(self.atualizar_metricas_janelas)

        self.info_otimizacao = QLabel("📊 Configure as datas acima")
        self.info_otimizacao.setWordWrap(True)
        temp_l.addWidget(self.info_otimizacao)

        self.processar_btn = QPushButton("⚡ Processar Período Selecionado")
        self.processar_btn.setEnabled(False)
        self.processar_btn.clicked.connect(self.processar_periodo)
        temp_l.addWidget(self.processar_btn)

        self.status_processamento = QLabel("")
        self.status_processamento.setWordWrap(True)
        temp_l.addWidget(self.status_processamento)
        temp_l.addStretch()

        left.addWidget(temp_box, 1)

        # ----- Coluna direita: Seleção de Ativos -----
        assets_box = QGroupBox("📋 Seleção de Ativos")
        assets_l = QVBoxLayout(assets_box)

        lbl1 = QLabel("Selecione os ativos:")
        lbl1.setFont(bold())
        assets_l.addWidget(lbl1)
        assets_l.addWidget(QLabel("(Ctrl+clique para múltiplos)"))

        self.assets_listbox = AssetListBox()
        self.assets_listbox.itemSelectionChanged.connect(self._update_selection_info)
        assets_l.addWidget(self.assets_listbox, 1)

        btns = QHBoxLayout()
        b_all = QPushButton("✅ Todos")
        b_all.clicked.connect(self.select_all_assets)
        b_clear = QPushButton("❌ Limpar")
        b_clear.clicked.connect(self.clear_selection)
        btns.addWidget(b_all)
        btns.addWidget(b_clear)
        btns.addStretch()
        assets_l.addLayout(btns)

        self.selection_info = QLabel("")
        self.selection_info.setStyleSheet("color: blue;")
        assets_l.addWidget(self.selection_info)

        right.addWidget(assets_box, 1)

        self._update_selection_info()

    def _update_selection_info(self):
        try:
            selected_count = len(self.assets_listbox.curselection())
            total_count = self.assets_listbox.size()
            if total_count > 0:
                self.selection_info.setText(f"Selecionados: {selected_count}/{total_count}")
            else:
                self.selection_info.setText("Carregue dados primeiro")
        except Exception:
            pass

    def toggle_validacao(self):
        self.data_fim_analise.setEnabled(self.usar_validacao.get())

    def processar_periodo(self):
        """Processar período selecionado - VERSÃO SIMPLIFICADA igual ao Streamlit"""
        if self.dados_brutos is None:
            messagebox.showerror("Erro", "Carregue dados primeiro!")
            return

        try:
            data_inicio_otim = pd.to_datetime(self.data_inicio_otim.get(), format='%d/%m/%Y')
            data_fim_otim = pd.to_datetime(self.data_fim_otim.get(), format='%d/%m/%Y')

            if self.usar_validacao.get():
                data_fim_analise = pd.to_datetime(self.data_fim_analise.get(), format='%d/%m/%Y')
            else:
                data_fim_analise = None

            if data_inicio_otim >= data_fim_otim:
                messagebox.showerror("Erro", "Data de início deve ser anterior à data fim!")
                return

            if data_fim_analise is not None and data_fim_analise <= data_fim_otim:
                messagebox.showerror("Erro", "Data fim da análise deve ser posterior ao fim da otimização!")
                return

            self.status_processamento.setText("🔄 Processando dados...")
            self.status_processamento.setStyleSheet("color: blue;")
            self.processar_btn.setEnabled(False)
            QApplication.processEvents()

            # 1. Configurar dados
            if 'Data' in self.dados_brutos.columns:
                df_trabalho = self.dados_brutos.copy()
                df_trabalho['Data'] = pd.to_datetime(df_trabalho['Data'])
                df_trabalho = df_trabalho.set_index('Data')
            else:
                df_trabalho = self.dados_brutos.copy()
                if not isinstance(df_trabalho.index, pd.DatetimeIndex):
                    df_trabalho.index = pd.to_datetime(df_trabalho.index)

            # 2. Filtrar período para otimização
            df_otimizacao = df_trabalho[(df_trabalho.index >= data_inicio_otim) &
                                        (df_trabalho.index <= data_fim_otim)].copy()

            # 3. Período estendido (se houver validação)
            df_analise_estendida = None
            if data_fim_analise is not None and data_fim_analise > data_fim_otim:
                df_analise_estendida = df_trabalho[(df_trabalho.index >= data_inicio_otim) &
                                                   (df_trabalho.index <= data_fim_analise)].copy()

            # 4. Base zero - otimização
            df_base0_otimizacao, cols_removidas_otim = self.transformar_base_zero(df_otimizacao)

            # 5. Base zero - estendido
            df_base0_analise = None
            if df_analise_estendida is not None:
                df_base0_analise, _ = self.transformar_base_zero(df_analise_estendida)

            # 6. Adicionar coluna de data de volta
            if df_base0_otimizacao is not None:
                df_base0_otimizacao = df_base0_otimizacao.reset_index()
                df_base0_otimizacao.rename(columns={'index': 'Data'}, inplace=True)

            if df_base0_analise is not None:
                df_base0_analise = df_base0_analise.reset_index()
                df_base0_analise.rename(columns={'index': 'Data'}, inplace=True)

            if df_base0_otimizacao is not None:
                self.df = df_base0_otimizacao
                self.df_otimizacao = df_base0_otimizacao
                self.df_analise = df_base0_analise

                self.periodo_otimizacao = {'inicio': data_inicio_otim, 'fim': data_fim_otim}
                self.periodo_analise = {
                    'inicio': data_inicio_otim,
                    'fim': data_fim_analise if data_fim_analise is not None else data_fim_otim
                }

                if self.has_risk_free:
                    try:
                        temp_optimizer = PortfolioOptimizer(df_base0_otimizacao, [])
                        if hasattr(temp_optimizer, 'risk_free_rate_total'):
                            self.detected_risk_free_rate = temp_optimizer.risk_free_rate_total
                    except Exception:
                        pass

                # Atualizar listbox preservando seleção
                selected_indices = self.assets_listbox.curselection()
                selected_assets = [self.assets_listbox.get(i) for i in selected_indices]

                self.assets_listbox.delete(0, END)

                ativos_validos = []
                for col in df_base0_otimizacao.columns:
                    if col != 'Data' and col != self.risk_free_column_name:
                        ativos_validos.append(col)

                for asset in ativos_validos:
                    self.assets_listbox.insert(END, asset)

                for i in range(self.assets_listbox.size()):
                    asset = self.assets_listbox.get(i)
                    if asset in selected_assets:
                        self.assets_listbox.selection_set(i)

                self.update_advanced_widgets()

                dias_otim = (data_fim_otim - data_inicio_otim).days
                if data_fim_analise is not None:
                    dias_total = (data_fim_analise - data_inicio_otim).days
                    dias_valid = (data_fim_analise - data_fim_otim).days
                    info_text = (
                        f"✅ Período processado com sucesso!\n"
                        f"📊 Otimização: {dias_otim} dias "
                        f"({data_inicio_otim.strftime('%d/%m/%Y')} a {data_fim_otim.strftime('%d/%m/%Y')})\n"
                        f"🔍 Validação: {dias_valid} dias (até {data_fim_analise.strftime('%d/%m/%Y')})\n"
                        f"📈 Total: {dias_total} dias | {len(ativos_validos)} ativos disponíveis"
                    )
                else:
                    info_text = (
                        f"✅ Período processado com sucesso!\n"
                        f"📊 Otimização: {dias_otim} dias "
                        f"({data_inicio_otim.strftime('%d/%m/%Y')} a {data_fim_otim.strftime('%d/%m/%Y')})\n"
                        f"📈 {len(ativos_validos)} ativos disponíveis"
                    )

                if cols_removidas_otim:
                    info_text += f"\n⚠️ Removidos: {', '.join(cols_removidas_otim)}"

                self.status_processamento.setText(info_text)
                self.status_processamento.setStyleSheet("color: green;")
                self.processar_btn.setEnabled(True)

                messagebox.showinfo("Sucesso", "🎉 Dados processados!\n🎯 Agora você pode otimizar o portfólio.")

                self.notebook.setCurrentWidget(self.tab_config)
            else:
                self.status_processamento.setText("❌ Erro no processamento")
                self.status_processamento.setStyleSheet("color: red;")
                self.processar_btn.setEnabled(True)

        except Exception as e:
            self.status_processamento.setText("❌ Erro no processamento")
            self.status_processamento.setStyleSheet("color: red;")
            self.processar_btn.setEnabled(True)
            messagebox.showerror("Erro", f"Erro ao processar período:\n{str(e)}")

    def atualizar_datas_automaticas(self):
        """Atualizar datas automaticamente quando carregar arquivo"""
        if self.periodo_disponivel:
            total_dias = (self.periodo_disponivel['fim'] - self.periodo_disponivel['inicio']).days
            dias_otimizacao = int(total_dias * 0.7)

            data_inicio = self.periodo_disponivel['inicio']
            data_fim_otim = data_inicio + timedelta(days=dias_otimizacao)
            data_fim_analise = self.periodo_disponivel['fim']

            self.data_inicio_otim.set_date(data_inicio.date())
            self.data_fim_otim.set_date(data_fim_otim.date())
            self.data_fim_analise.set_date(data_fim_analise.date())

            self.processar_btn.setEnabled(True)

            dias_otim = (data_fim_otim - data_inicio).days
            dias_valid = (data_fim_analise - data_fim_otim).days

            info_text = (
                f"📊 Otimização: {dias_otim} dias ({(dias_otim/total_dias*100):.0f}% do total)\n"
                f"🔍 Validação: {dias_valid} dias ({(dias_valid/total_dias*100):.0f}% do total)\n"
                f"📈 Total: {total_dias} dias"
            )
            self.info_otimizacao.setText(info_text)

    def atualizar_metricas_janelas(self, *args):
        """Atualizar informações das janelas quando datas mudarem"""
        try:
            data_inicio = self.data_inicio_otim.get_date()
            data_fim_otim = self.data_fim_otim.get_date()
            data_fim_analise = self.data_fim_analise.get_date()

            dias_otim = (data_fim_otim - data_inicio).days
            dias_valid = (data_fim_analise - data_fim_otim).days
            dias_total = dias_otim + dias_valid

            if dias_total > 0:
                info_text = (
                    f"📊 Otimização: {dias_otim} dias ({(dias_otim/dias_total*100):.0f}% do total)\n"
                    f"🔍 Validação: {dias_valid} dias ({(dias_valid/dias_total*100):.0f}% do total)\n"
                    f"📈 Total: {dias_total} dias"
                )
            else:
                info_text = "⚠️ Configure datas válidas"

            self.info_otimizacao.setText(info_text)

            if self.dados_brutos is not None and dias_otim > 0:
                self.processar_btn.setEnabled(True)
            else:
                self.processar_btn.setEnabled(False)
        except Exception:
            self.info_otimizacao.setText("⚠️ Configure datas válidas")
            self.processar_btn.setEnabled(False)

    def transformar_base_zero(self, df_precos):
        """Transforma dados de preços para base 0 (IDÊNTICO À VERSÃO TKINTER)"""
        if df_precos is None or df_precos.empty:
            return None, []

        df_limpo = df_precos.copy()

        colunas_removidas = []
        for coluna in df_limpo.columns:
            if len(df_limpo[coluna]) == 0:
                colunas_removidas.append(coluna)
                continue

            primeiro_valor = df_limpo[coluna].iloc[0]
            if pd.isna(primeiro_valor) or primeiro_valor == 0:
                colunas_removidas.append(coluna)

        if colunas_removidas:
            df_limpo = df_limpo.drop(columns=colunas_removidas)

        if df_limpo.empty:
            return None, colunas_removidas

        for coluna in df_limpo.columns:
            df_limpo[coluna] = df_limpo[coluna].replace(0, np.nan)
            df_limpo[coluna] = df_limpo[coluna].ffill()

        df_limpo = df_limpo.fillna(0)

        base_zero_data = {}
        for coluna in df_limpo.columns:
            valores = df_limpo[coluna].values

            if len(valores) == 0:
                continue

            cota_1 = valores[0]
            if cota_1 == 0:
                continue

            novos_valores = np.zeros(len(valores))
            novos_valores[0] = 0.0

            for i in range(1, len(valores)):
                cota_n = valores[i]
                cota_anterior = valores[i - 1]
                novo_valor = (cota_n - cota_anterior) / cota_1
                novos_valores[i] = novo_valor

            base_zero_data[coluna] = novos_valores

        if base_zero_data:
            df_base_zero = pd.DataFrame(base_zero_data, index=df_limpo.index)
            return df_base_zero, colunas_removidas
        else:
            return None, colunas_removidas

    # -------------------------------------------------------------------------
    # ABA 2: CONFIGURAÇÃO
    # -------------------------------------------------------------------------
    def setup_config_tab(self):
        scroll, container, layout = self._scroll_area(self.tab_config)

        # 1. Objetivo
        obj_box = QGroupBox("🎯 Objetivo de Otimização")
        obj_l = QVBoxLayout(obj_box)

        self.objective_var = RadioVar(default="Maximizar Sharpe Ratio")

        self.base_objectives = [
            "Maximizar Sharpe Ratio",
            "Maximizar Sortino Ratio",
            "Minimizar Risco",
            "Minimizar Under Water",
            "Maximizar Inclinação/[(1-R²)×Vol]",
            "Maximizar Qualidade da Linearidade"
        ]
        self.risk_free_objectives = [
            "Maximizar Linearidade do Excesso",
            "Maximizar Sharpe do Excesso"
        ]

        self.objective_buttons = {}
        for obj in self.base_objectives:
            rb = QRadioButton(obj)
            self.objective_group.addButton(rb)
            obj_l.addWidget(rb)
            self.objective_buttons[obj] = rb
            self.objective_var.add(obj, rb)
        self.objective_buttons["Maximizar Sharpe Ratio"].setChecked(True)

        # Placeholder para objetivos de taxa livre
        self.risk_free_obj_frame = QWidget()
        self.risk_free_obj_layout = QVBoxLayout(self.risk_free_obj_frame)
        self.risk_free_obj_layout.setContentsMargins(0, 0, 0, 0)
        obj_l.addWidget(self.risk_free_obj_frame)

        layout.addWidget(obj_box)

        # 2. Limites de Peso Globais
        limits_box = QGroupBox("📊 Limites de Peso Globais")
        limits_l = QVBoxLayout(limits_box)
        self.min_weight_var = self._add_slider(limits_l, "Peso mínimo por ativo (%):", 0, 20, 0.0, 1)
        self.max_weight_var = self._add_slider(limits_l, "Peso máximo por ativo (%):", 5, 100, 30.0, 1)
        layout.addWidget(limits_box)

        # 3. Taxa de Referência
        risk_box = QGroupBox("🏛️ Taxa de Referência")
        risk_l = QVBoxLayout(risk_box)
        risk_l.addWidget(QLabel("Taxa de referência manual (% acumulada):"))
        self.manual_risk_entry = QLineEdit()
        self.manual_risk_entry.setFixedWidth(120)
        self.risk_free_var = NumVar(self.manual_risk_entry, 0.0)
        risk_l.addWidget(self.manual_risk_entry)
        layout.addWidget(risk_box)

        # 3b. Meta de Retorno (opcional)
        meta_box = QGroupBox("🎯 Meta de Retorno (opcional)")
        meta_l = QVBoxLayout(meta_box)
        chk_meta = QCheckBox("Exigir meta de retorno mínima")
        self.use_meta = BoolVar(chk_meta)
        meta_l.addWidget(chk_meta)

        # Tipo de meta: relativa à referência ou absoluta (% ao ano)
        self.meta_mode_group = QButtonGroup(self)
        self.meta_mode_var = RadioVar(default="relativa")

        rel_row = QHBoxLayout()
        rb_rel = QRadioButton("Relativa:")
        self.meta_mode_group.addButton(rb_rel)
        self.meta_mode_var.add("relativa", rb_rel)
        rel_row.addWidget(rb_rel)
        self.meta_entry = QLineEdit()
        self.meta_entry.setFixedWidth(70)
        self.meta_var = NumVar(self.meta_entry, 5.0)
        rel_row.addWidget(self.meta_entry)
        rel_row.addWidget(QLabel("% acima da referência"))
        rel_row.addStretch()
        meta_l.addLayout(rel_row)

        abs_row = QHBoxLayout()
        rb_abs = QRadioButton("Absoluta:")
        self.meta_mode_group.addButton(rb_abs)
        self.meta_mode_var.add("absoluta", rb_abs)
        abs_row.addWidget(rb_abs)
        self.meta_abs_entry = QLineEdit()
        self.meta_abs_entry.setFixedWidth(70)
        self.meta_abs_var = NumVar(self.meta_abs_entry, 15.0)
        abs_row.addWidget(self.meta_abs_entry)
        abs_row.addWidget(QLabel("% ao ano (não depende da referência)"))
        abs_row.addStretch()
        meta_l.addLayout(abs_row)

        meta_hint = QLabel(
            "Combina com o objetivo escolhido acima: maximiza o objetivo garantindo um retorno "
            "mínimo. Na prática é o menor risco que alcança a meta.\n"
            "• RELATIVA: alvo = referência × (1 + meta/100) no período "
            "(ex.: referência 12% e meta 5% → alvo 12,6%). Exige taxa de referência.\n"
            "• ABSOLUTA: meta em % ao ano, convertida para o período "
            "(ex.: 15% ao ano em 126 pregões → alvo 7,2%). Funciona sem taxa de referência.\n"
            "Se a meta for inatingível, retorna a carteira de MAIOR retorno possível e avisa.")
        meta_hint.setStyleSheet("color: gray;")
        meta_hint.setWordWrap(True)
        meta_l.addWidget(meta_hint)
        layout.addWidget(meta_box)

        # 4. Botão Otimizar
        btn_opt = QPushButton("🚀 OTIMIZAR PORTFÓLIO")
        btn_opt.setStyleSheet(
            "QPushButton { background-color: #0078d4; color: white; font-weight: bold; padding: 10px; }"
        )
        btn_opt.clicked.connect(self.optimize_portfolio)
        layout.addWidget(btn_opt)

        layout.addStretch()

    # -------------------------------------------------------------------------
    # ABA 3: RESTRIÇÕES INDIVIDUAIS
    # -------------------------------------------------------------------------
    def setup_advanced_tab(self):
        box = QGroupBox("🚫 Restrições Individuais por Ativo")
        outer = QVBoxLayout(self.tab_advanced)
        outer.addWidget(box)
        box_l = QVBoxLayout(box)

        chk = QCheckBox("Habilitar limites específicos para ativos selecionados")
        self.use_individual_constraints = BoolVar(chk)
        chk.toggled.connect(self.toggle_individual_constraints)
        box_l.addWidget(chk)

        # Importação de restrições a partir de arquivo
        import_row = QHBoxLayout()
        btn_import = QPushButton("📂 Importar Restrições (Excel/CSV)")
        btn_import.clicked.connect(self.import_constraints_file)
        import_row.addWidget(btn_import)
        import_row.addStretch()
        box_l.addLayout(import_row)

        hint = QLabel("Colunas aceitas: 'Ativo, Min, Max' (faixas) ou 'Ativo, Peso' "
                      "(pesos fixos → analisa o portfólio). Valores em % (ex.: 30 = 30%).")
        hint.setStyleSheet("color: gray;")
        hint.setWordWrap(True)
        box_l.addWidget(hint)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        self.constraints_frame = QWidget()
        self.constraints_layout = QVBoxLayout(self.constraints_frame)
        scroll.setWidget(self.constraints_frame)
        box_l.addWidget(scroll, 1)

        self.constraints_info = QLabel("Carregue dados e selecione ativos primeiro")
        self.constraints_layout.addWidget(self.constraints_info)

    def import_constraints_file(self):
        """Importar restrições mín/máx (ou pesos) de um arquivo Excel/CSV."""
        if self.assets_listbox.size() == 0:
            messagebox.showerror("Erro", "Carregue os dados primeiro (aba Dados)!")
            return

        file_path = filedialog.askopenfilename(
            title="Importar restrições (Excel/CSV)",
            filetypes=[("Planilhas", "*.xlsx *.xls *.csv"),
                       ("Excel", "*.xlsx *.xls"),
                       ("CSV", "*.csv"),
                       ("Todos", "*.*")]
        )
        if not file_path:
            return

        try:
            constraints, warnings = load_constraints_from_file(file_path)
        except Exception as e:
            messagebox.showerror("Erro", f"Falha ao ler o arquivo:\n{str(e)}")
            return

        self._apply_imported_constraints(constraints, warnings)

    def _apply_imported_constraints(self, constraints, warnings):
        """Casa os ativos do arquivo com os dados, seleciona e ativa as restrições."""
        available = list(self.assets_listbox.get(0, END))
        avail_lower = {a.lower(): a for a in available}

        matched = {}
        not_found = []
        for asset, lim in constraints.items():
            if asset in available:
                matched[asset] = lim
            elif asset.lower() in avail_lower:
                matched[avail_lower[asset.lower()]] = lim
            else:
                not_found.append(asset)

        if not matched:
            exemplos = ', '.join(list(constraints.keys())[:10])
            messagebox.showwarning(
                "Atenção",
                "Nenhum ativo do arquivo corresponde aos ativos carregados.\n"
                f"Ativos no arquivo: {exemplos}")
            return

        # Selecionar exatamente os ativos importados e ativar restrições
        self.use_individual_constraints.set(True)
        self.individual_constraints = matched

        self.assets_listbox.selection_clear(0, END)
        for i in range(self.assets_listbox.size()):
            if self.assets_listbox.get(i) in matched:
                self.assets_listbox.selection_set(i)

        self.update_advanced_widgets()
        self._update_selection_info()

        # Diagnóstico: portfólio (min == max) e soma
        is_portfolio = all(abs(v['min'] - v['max']) < 1e-9 for v in matched.values())
        soma = sum((v['min'] + v['max']) / 2 for v in matched.values()) * 100

        msg = f"✅ {len(matched)} ativos importados e selecionados."
        if is_portfolio:
            msg += f"\n📊 Modo portfólio (pesos fixos). Soma = {soma:.1f}%."
            if abs(soma - 100) < 0.5:
                msg += "\n🎯 Soma ≈ 100%: o software atuará como analisador deste portfólio."
        if not_found:
            extra = ', '.join(not_found[:8]) + ('...' if len(not_found) > 8 else '')
            msg += f"\n⚠️ {len(not_found)} não encontrados nos dados: {extra}"
        if warnings:
            msg += "\n\n" + "\n".join(warnings[:6])
            if len(warnings) > 6:
                msg += "\n..."

        messagebox.showinfo("Importação concluída", msg)
        self.notebook.setCurrentWidget(self.tab_advanced)

    # -------------------------------------------------------------------------
    # ABA 4: SHORT SELLING
    # -------------------------------------------------------------------------
    def setup_short_tab(self):
        box = QGroupBox("🔄 Posições Short / Hedge")
        outer = QVBoxLayout(self.tab_short)
        outer.addWidget(box)
        box_l = QVBoxLayout(box)

        chk = QCheckBox("Habilitar posições short/hedge (venda a descoberto)")
        self.use_short = BoolVar(chk)
        chk.toggled.connect(self.toggle_short_selling)
        box_l.addWidget(chk)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        self.short_frame = QWidget()
        self.short_layout = QVBoxLayout(self.short_frame)
        scroll.setWidget(self.short_frame)
        box_l.addWidget(scroll, 1)

        self.short_info = QLabel("Carregue dados e selecione ativos principais primeiro")
        self.short_layout.addWidget(self.short_info)

    # -------------------------------------------------------------------------
    # ABA 5: RANKING
    # -------------------------------------------------------------------------
    def setup_ranking_tab(self):
        scroll, container, layout = self._scroll_area(self.tab_ranking)

        title = QLabel("🏆 Sistema de Ranking de Ativos")
        title.setFont(bold(14))
        layout.addWidget(title)

        chk = QCheckBox("🤖 Ativar ranking automático de ativos")
        self.use_ranking = BoolVar(chk)
        chk.toggled.connect(self.toggle_ranking)
        layout.addWidget(chk)

        # Config de pesos (oculto inicialmente)
        self.ranking_config_frame = QGroupBox("⚙️ Configurar Pesos dos Parâmetros")
        cfg_l = QVBoxLayout(self.ranking_config_frame)
        self.peso_inc_var = self._add_slider(cfg_l, "📈 Peso Inclinação:", 0, 1, 0.33, 2, suffix="")
        self.peso_desv_var = self._add_slider(cfg_l, "📊 Peso Estabilidade:", 0, 1, 0.33, 2, suffix="")
        self.peso_cor_var = self._add_slider(cfg_l, "🎯 Peso Correlação:", 0, 1, 0.33, 2, suffix="")

        btn_calc = QPushButton("🔄 Calcular Ranking")
        btn_calc.clicked.connect(self.calculate_ranking)
        cfg_l.addWidget(btn_calc)

        self.ranking_config_frame.setVisible(False)
        layout.addWidget(self.ranking_config_frame)

        # Resultados (oculto inicialmente)
        self.ranking_results_frame = QGroupBox("📊 Resultados do Ranking")
        self.ranking_results_layout = QVBoxLayout(self.ranking_results_frame)
        self.ranking_status = QLabel("Ative o ranking e clique em 'Calcular Ranking'")
        self.ranking_results_layout.addWidget(self.ranking_status)
        self.ranking_results_frame.setVisible(False)
        layout.addWidget(self.ranking_results_frame, 1)

        layout.addStretch()

    def toggle_ranking(self):
        show = self.use_ranking.get()
        self.ranking_config_frame.setVisible(show)
        self.ranking_results_frame.setVisible(show)

    def calculate_ranking(self):
        if self.df is None:
            messagebox.showerror("Erro", "Carregue dados primeiro!")
            return

        try:
            peso_inc = self.peso_inc_var.get()
            peso_desv = self.peso_desv_var.get()
            peso_cor = self.peso_cor_var.get()

            ranking_result = calculate_asset_ranking(
                self.df, self.risk_free_column_name, peso_inc, peso_desv, peso_cor
            )

            if ranking_result is not None:
                self.display_ranking_results(ranking_result)
            else:
                self.ranking_status.setText("❌ Erro ao calcular ranking")
        except Exception as e:
            messagebox.showerror("Erro", f"Erro no cálculo: {str(e)}")

    def display_ranking_results(self, ranking_result):
        clear_widget(self.ranking_results_frame)

        df_ranking = ranking_result['ranking']

        info_text = (f"✅ Ranking calculado: {ranking_result['total_ativos']} ativos\n"
                     f"📊 Referência: {ranking_result['referencia']}")
        self.ranking_results_layout.addWidget(QLabel(info_text))

        columns = ('Posição', 'Ativo', 'Índice', 'Inclinação', 'R²', 'Correlação', 'Desvio')
        table = QTableWidget(0, len(columns))
        table.setHorizontalHeaderLabels(columns)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)

        for _, row in df_ranking.head(20).iterrows():
            values = (
                str(int(row['Posição'])),
                str(row['Ativo']),
                f"{row['Índice']:.4f}",
                f"{row['Inclinação_Norm']:.3f}",
                f"{row['R²']:.3f}",
                f"{row['Correlação']:.3f}",
                f"{row['Desvio_Norm']:.3f}",
            )
            r = table.rowCount()
            table.insertRow(r)
            for c, val in enumerate(values):
                item = QTableWidgetItem(val)
                item.setTextAlignment(Qt.AlignCenter)
                table.setItem(r, c, item)

        table.setMinimumHeight(360)
        self.ranking_results_layout.addWidget(table)

        # Seleção Automática por Score
        selection_box = QGroupBox("🎯 Seleção Automática de Ativos")
        sel_l = QVBoxLayout(selection_box)

        range_l = QHBoxLayout()
        min_col = QVBoxLayout()
        min_col.addWidget(QLabel("📉 Score mínimo:"))
        e_min = QLineEdit()
        self.score_min_var = NumVar(e_min, 0.7)
        min_col.addWidget(e_min)
        range_l.addLayout(min_col)

        max_col = QVBoxLayout()
        max_col.addWidget(QLabel("📈 Score máximo:"))
        e_max = QLineEdit()
        self.score_max_var = NumVar(e_max, 1.0)
        max_col.addWidget(e_max)
        range_l.addLayout(max_col)
        sel_l.addLayout(range_l)

        btn_row = QHBoxLayout()
        btn_sel = QPushButton("✅ Selecionar Ativos por Score")
        btn_sel.clicked.connect(self.select_assets_by_score)
        btn_row.addWidget(btn_sel)
        self.selection_result_label = QLabel("")
        btn_row.addWidget(self.selection_result_label)
        btn_row.addStretch()
        sel_l.addLayout(btn_row)

        self.ranking_results_layout.addWidget(selection_box)

        self.ranking_result = ranking_result

    def select_assets_by_score(self):
        if not hasattr(self, 'ranking_result') or self.ranking_result is None:
            messagebox.showerror("Erro", "Calcule o ranking primeiro!")
            return

        try:
            score_min = self.score_min_var.get()
            score_max = self.score_max_var.get()

            if score_min > score_max:
                messagebox.showerror("Erro", "Score mínimo deve ser menor que o máximo!")
                return

            df_ranking = self.ranking_result['ranking']
            filtered_ranking = df_ranking[
                (df_ranking['Índice'] >= score_min) &
                (df_ranking['Índice'] <= score_max)
            ]

            if len(filtered_ranking) == 0:
                self.selection_result_label.setText(
                    f"⚠️ Nenhum ativo no range {score_min:.2f} - {score_max:.2f}")
                self.selection_result_label.setStyleSheet("color: orange;")
                return

            selected_assets = filtered_ranking['Ativo'].tolist()

            short_backup = self.short_weights.copy() if hasattr(self, 'short_weights') else {}
            individual_backup = self.individual_constraints.copy() if hasattr(self, 'individual_constraints') else {}

            self.assets_listbox.selection_clear(0, END)

            for i in range(self.assets_listbox.size()):
                asset = self.assets_listbox.get(i)
                if asset in selected_assets:
                    self.assets_listbox.selection_set(i)

            self.update_advanced_widgets()

            self.short_weights = short_backup
            if hasattr(self, 'use_short') and self.use_short.get():
                try:
                    self.update_short_summary()
                except Exception:
                    pass

            self.individual_constraints = individual_backup

            avg_score = filtered_ranking['Índice'].mean()
            self.selection_result_label.setText(
                f"✅ {len(selected_assets)} ativos selecionados (score médio: {avg_score:.3f})")
            self.selection_result_label.setStyleSheet("color: green;")

            self.notebook.setCurrentWidget(self.tab_config)
        except Exception as e:
            messagebox.showerror("Erro", f"Erro na seleção: {str(e)}")

    # -------------------------------------------------------------------------
    # ABA 6: RESULTADOS
    # -------------------------------------------------------------------------
    def setup_results_tab(self):
        outer = QVBoxLayout(self.tab_results)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        self.chart_frame = QWidget()
        self.chart_layout = QVBoxLayout(self.chart_frame)
        scroll.setWidget(self.chart_frame)
        outer.addWidget(scroll)

        initial = QLabel("Execute uma otimização para ver os resultados")
        initial.setStyleSheet("color: gray;")
        initial.setAlignment(Qt.AlignCenter)
        self.chart_layout.addWidget(initial)

        self.export_csv_btn = None
        self.export_excel_btn = None
        self.portfolio_tree = None

    # -------------------------------------------------------------------------
    # ABA 7: RETORNOS MENSAIS
    # -------------------------------------------------------------------------
    def setup_monthly_tab(self):
        outer = QVBoxLayout(self.tab_monthly)
        self.monthly_notebook = QTabWidget()
        outer.addWidget(self.monthly_notebook)

        # Aba portfólio
        tab_pf = QWidget()
        self.monthly_notebook.addTab(tab_pf, "📊 Retornos do Portfólio")
        scroll1 = QScrollArea()
        scroll1.setWidgetResizable(True)
        self.portfolio_monthly_frame = QWidget()
        self.portfolio_monthly_layout = QVBoxLayout(self.portfolio_monthly_frame)
        scroll1.setWidget(self.portfolio_monthly_frame)
        QVBoxLayout(tab_pf).addWidget(scroll1)

        # Aba excesso
        tab_ex = QWidget()
        self.monthly_notebook.addTab(tab_ex, "📈 Excesso de Retorno")
        scroll2 = QScrollArea()
        scroll2.setWidgetResizable(True)
        self.excess_monthly_frame = QWidget()
        self.excess_monthly_layout = QVBoxLayout(self.excess_monthly_frame)
        scroll2.setWidget(self.excess_monthly_frame)
        QVBoxLayout(tab_ex).addWidget(scroll2)

        self.portfolio_monthly_layout.addWidget(
            QLabel("Execute uma otimização para ver as tabelas mensais"))
        self.excess_monthly_layout.addWidget(
            QLabel("Execute uma otimização para ver as tabelas mensais"))

    def create_monthly_returns_table_desktop(self, returns_data, weights, dates=None, risk_free_returns=None):
        """Cria tabela de retornos mensais (IDÊNTICO À VERSÃO TKINTER)"""
        portfolio_returns_daily = np.dot(returns_data.values, weights)
        portfolio_cumulative = np.cumsum(portfolio_returns_daily)

        if dates is not None:
            portfolio_df = pd.DataFrame({'cumulative': portfolio_cumulative}, index=dates)
        else:
            start_date = pd.Timestamp('2020-01-01')
            dates = pd.date_range(start=start_date, periods=len(portfolio_cumulative), freq='D')
            portfolio_df = pd.DataFrame({'cumulative': portfolio_cumulative}, index=dates)

        monthly_cumulative = portfolio_df['cumulative'].resample('ME').last()

        monthly_returns = []
        previous_cumulative = 0

        for month_date, current_cumulative in monthly_cumulative.items():
            if previous_cumulative != 0:
                monthly_return = (current_cumulative - previous_cumulative) / (1 + previous_cumulative)
            else:
                monthly_return = current_cumulative
            monthly_returns.append(monthly_return)
            previous_cumulative = current_cumulative

        monthly_returns_series = pd.Series(monthly_returns, index=monthly_cumulative.index)

        monthly_risk_free = None
        if risk_free_returns is not None:
            risk_free_cumulative = np.cumsum(risk_free_returns.values)

            if dates is not None:
                risk_free_df = pd.DataFrame({'cumulative': risk_free_cumulative}, index=dates)
            else:
                risk_free_df = pd.DataFrame({'cumulative': risk_free_cumulative}, index=dates)

            monthly_rf_cumulative = risk_free_df['cumulative'].resample('ME').last()

            monthly_rf_returns = []
            previous_rf_cumulative = 0

            for month_date, current_rf_cumulative in monthly_rf_cumulative.items():
                monthly_rf_return = current_rf_cumulative - previous_rf_cumulative
                monthly_rf_returns.append(monthly_rf_return)
                previous_rf_cumulative = current_rf_cumulative

            monthly_risk_free = pd.Series(monthly_rf_returns, index=monthly_rf_cumulative.index)

        monthly_df = pd.DataFrame({
            'Year': monthly_returns_series.index.year,
            'Month': monthly_returns_series.index.month,
            'Return': monthly_returns_series.values
        })

        pivot_table = monthly_df.pivot(index='Year', columns='Month', values='Return')

        month_names = {
            1: 'Jan', 2: 'Fev', 3: 'Mar', 4: 'Abr', 5: 'Mai', 6: 'Jun',
            7: 'Jul', 8: 'Ago', 9: 'Set', 10: 'Out', 11: 'Nov', 12: 'Dez'
        }
        pivot_table.columns = [month_names.get(col, f'M{col}') for col in pivot_table.columns]

        yearly_returns = []
        for year in pivot_table.index:
            year_data = pivot_table.loc[year].dropna()
            if len(year_data) > 0:
                annual_return = 1.0
                for monthly_return in year_data:
                    annual_return *= (1 + monthly_return)
                annual_return -= 1
                yearly_returns.append(annual_return)
            else:
                yearly_returns.append(np.nan)

        pivot_table['Total Anual'] = yearly_returns

        comparison_table = None
        if monthly_risk_free is not None:
            rf_monthly_df = pd.DataFrame({
                'Year': monthly_risk_free.index.year,
                'Month': monthly_risk_free.index.month,
                'Return': monthly_risk_free.values
            })

            rf_pivot = rf_monthly_df.pivot(index='Year', columns='Month', values='Return')
            rf_pivot.columns = [month_names.get(col, f'M{col}') for col in rf_pivot.columns]

            rf_yearly = []
            for year in rf_pivot.index:
                year_data = rf_pivot.loc[year].dropna()
                if len(year_data) > 0:
                    annual_return = year_data.sum()
                    rf_yearly.append(annual_return)
                else:
                    rf_yearly.append(np.nan)

            rf_pivot['Total Anual'] = rf_yearly
            comparison_table = pivot_table - rf_pivot

        return pivot_table, comparison_table

    def display_monthly_tables(self):
        if not self.result or not self.result['success']:
            return

        try:
            if hasattr(self, 'df_analise') and self.df_analise is not None:
                selected_assets = self.get_selected_assets()

                if hasattr(self, 'use_short') and self.use_short.get() and hasattr(self, 'short_weights'):
                    short_assets = list(self.short_weights.keys())
                    assets_used_in_optimization = selected_assets + short_assets
                else:
                    assets_used_in_optimization = selected_assets

                optimizer_to_use = PortfolioOptimizer(self.df_analise, assets_used_in_optimization)
            else:
                optimizer_to_use = self.optimizer

            dates = getattr(optimizer_to_use, 'dates', None)
            risk_free_returns = getattr(optimizer_to_use, 'risk_free_returns', None)

            self.monthly_table, self.excess_table = self.create_monthly_returns_table_desktop(
                optimizer_to_use.returns_data, self.result['weights'], dates, risk_free_returns
            )

            clear_widget(self.portfolio_monthly_frame)
            clear_widget(self.excess_monthly_frame)

            if hasattr(self, 'df_analise') and self.df_analise is not None:
                periodo_otim = self.periodo_otimizacao
                periodo_analise = self.periodo_analise

                info_l = QHBoxLayout()
                lbl_p = QLabel(f"📊 Período: {periodo_otim['inicio'].strftime('%d/%m/%Y')} a "
                               f"{periodo_analise['fim'].strftime('%d/%m/%Y')}")
                lbl_p.setFont(bold())
                info_l.addWidget(lbl_p)
                info_l.addStretch()
                lbl_i = QLabel("🔍 Incluindo: Otimização + Validação (período completo)")
                lbl_i.setStyleSheet("color: blue;")
                info_l.addWidget(lbl_i)
                self.portfolio_monthly_layout.addLayout(info_l)
            else:
                periodo_otim = self.periodo_otimizacao
                info_l = QHBoxLayout()
                lbl_p = QLabel(f"📊 Período: {periodo_otim['inicio'].strftime('%d/%m/%Y')} a "
                               f"{periodo_otim['fim'].strftime('%d/%m/%Y')}")
                lbl_p.setFont(bold())
                info_l.addWidget(lbl_p)
                info_l.addStretch()
                lbl_i = QLabel("⚠️ Apenas: Período de otimização")
                lbl_i.setStyleSheet("color: orange;")
                info_l.addWidget(lbl_i)
                self.portfolio_monthly_layout.addLayout(info_l)

            self.create_table_widget(self.portfolio_monthly_layout, self.monthly_table,
                                     "Retornos Mensais do Portfólio (%)")

            if self.excess_table is not None:
                self.create_table_widget(self.excess_monthly_layout, self.excess_table,
                                         "Excesso de Retorno Mensal (%)")
            else:
                self.excess_monthly_layout.addWidget(
                    QLabel("Não disponível (sem taxa de referência detectada)"))

        except Exception as e:
            self.portfolio_monthly_layout.addWidget(QLabel(f"Erro ao gerar tabelas: {str(e)}"))
            self.excess_monthly_layout.addWidget(QLabel(f"Erro ao gerar tabelas: {str(e)}"))

    def create_table_widget(self, parent_layout, data_df, title):
        """Criar tabela mensal com cores para uma DataFrame"""
        lbl = QLabel(title)
        lbl.setFont(bold(12))
        parent_layout.addWidget(lbl)

        columns = ['Ano'] + list(data_df.columns)
        table = QTableWidget(0, len(columns))
        table.setHorizontalHeaderLabels([str(c) for c in columns])
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)

        for year, row in data_df.iterrows():
            # Determinar prefixo (📈/📉) com base na média dos valores da linha
            numeric_values = [v for v in row.values if pd.notna(v)]
            year_str = str(year)
            if numeric_values:
                avg_return = sum(numeric_values) / len(numeric_values)
                if avg_return > 0:
                    year_str = f"📈 {year}"
                elif avg_return < 0:
                    year_str = f"📉 {year}"

            r = table.rowCount()
            table.insertRow(r)

            item_year = QTableWidgetItem(year_str)
            item_year.setTextAlignment(Qt.AlignCenter)
            table.setItem(r, 0, item_year)

            for c, val in enumerate(row.values, start=1):
                text = f"{val:.2%}" if pd.notna(val) else "-"
                item = QTableWidgetItem(text)
                item.setTextAlignment(Qt.AlignCenter)
                if pd.notna(val):
                    if val > 0:
                        item.setForeground(QColor("#127a12"))
                    elif val < 0:
                        item.setForeground(QColor("#b00020"))
                table.setItem(r, c, item)

        table.setMinimumHeight(300)
        parent_layout.addWidget(table)

    # -------------------------------------------------------------------------
    # EXPORTAÇÕES
    # -------------------------------------------------------------------------
    def export_csv(self):
        if not self.result or not self.result['success']:
            messagebox.showerror("Erro", "Execute uma otimização primeiro!")
            return
        try:
            filename = filedialog.asksaveasfilename(
                title="Salvar Composição do Portfólio",
                defaultextension=".csv",
                filetypes=[("CSV files", "*.csv")]
            )
            if filename:
                portfolio_summary = self.get_portfolio_summary_desktop(
                    self.result['weights'], self.result['assets'], self.optimizer.returns_data)
                portfolio_summary.to_csv(filename, index=False)
                messagebox.showinfo("Sucesso", f"CSV exportado com todas as colunas!\n{filename}")
        except Exception as e:
            messagebox.showerror("Erro", f"Erro ao exportar CSV:\n{str(e)}")

    def export_excel(self):
        if not self.result or not self.result['success']:
            messagebox.showerror("Erro", "Execute uma otimização primeiro!")
            return
        try:
            filename = filedialog.asksaveasfilename(
                title="Salvar Composição do Portfólio",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")]
            )
            if filename:
                portfolio_summary = self.get_portfolio_summary_desktop(
                    self.result['weights'], self.result['assets'], self.optimizer.returns_data)
                portfolio_summary.to_excel(filename, index=False)
                messagebox.showinfo("Sucesso", f"Excel exportado com todas as colunas!\n{filename}")
        except Exception as e:
            messagebox.showerror("Erro", f"Erro ao exportar Excel:\n{str(e)}")

    def prepare_export_data(self):
        all_weights = self.result['weights']
        export_list = []
        for i, asset in enumerate(self.result['assets']):
            weight = all_weights[i]
            if abs(weight) > 0.001:
                export_list.append({
                    'Ativo': asset,
                    'Peso_Decimal': weight,
                    'Peso_Percentual': weight * 100,
                    'Tipo': 'LONG' if weight > 0 else 'SHORT'
                })
        return pd.DataFrame(export_list)

    def prepare_metrics_export(self):
        metrics = self.result['metrics']
        metrics_list = [
            {'Métrica': 'Retorno Total', 'Valor': metrics['gv_final'], 'Formato': f"{metrics['gv_final']:.4f}"},
            {'Métrica': 'Retorno Anualizado', 'Valor': metrics['annual_return'], 'Formato': f"{metrics['annual_return']:.4f}"},
            {'Métrica': 'Volatilidade', 'Valor': metrics['volatility'], 'Formato': f"{metrics['volatility']:.4f}"},
            {'Métrica': 'Sharpe Ratio', 'Valor': metrics['sharpe_ratio'], 'Formato': f"{metrics['sharpe_ratio']:.4f}"},
            {'Métrica': 'Downside Deviation', 'Valor': metrics['downside_deviation'], 'Formato': f"{metrics['downside_deviation']:.4f}"},
            {'Métrica': 'Excesso de Retorno', 'Valor': metrics['excess_return'], 'Formato': f"{metrics['excess_return']:.4f}"},
            {'Métrica': 'R²', 'Valor': metrics['r_squared'], 'Formato': f"{metrics['r_squared']:.4f}"},
            {'Métrica': 'VaR 95% Diário', 'Valor': metrics['var_95_daily'], 'Formato': f"{metrics['var_95_daily']:.4f}"},
            {'Métrica': 'CVaR 95% Diário', 'Valor': metrics['cvar_95_daily'], 'Formato': f"{metrics['cvar_95_daily']:.4f}"},
            {'Métrica': 'Taxa de Referência', 'Valor': metrics['risk_free_rate'], 'Formato': f"{metrics['risk_free_rate']:.4f}"},
        ]

        if self.objective_var.get() in ("Maximizar Linearidade do Excesso", "Maximizar Sharpe do Excesso") and metrics.get('excess_r_squared') is not None:
            if hasattr(self.optimizer, 'risk_free_returns') and self.optimizer.risk_free_returns is not None:
                excess_returns_daily = metrics['portfolio_returns_daily'] - self.optimizer.risk_free_returns.values
                excess_vol = np.std(excess_returns_daily, ddof=0) * np.sqrt(252)
                metrics_list.extend([
                    {'Métrica': 'R² do Excesso', 'Valor': metrics['excess_r_squared'], 'Formato': f"{metrics['excess_r_squared']:.4f}"},
                    {'Métrica': 'Volatilidade do Excesso', 'Valor': excess_vol, 'Formato': f"{excess_vol:.4f}"},
                ])

        return pd.DataFrame(metrics_list)

    def load_excel_file(self):
        file_path = filedialog.askopenfilename(
            title="Selecionar planilha Excel",
            filetypes=[("Excel files", "*.xlsx *.xls")]
        )
        if file_path:
            try:
                df = pd.read_excel(file_path)
                if len(df.columns) > 0:
                    df.columns.values[0] = "Data"
                self._apply_raw_data(df, f"Arquivo: {os.path.basename(file_path)}")
                messagebox.showinfo("Sucesso", "📥 Dados brutos carregados!\n🎯 Agora configure as janelas temporais.")
            except Exception as e:
                messagebox.showerror("Erro", f"Erro ao carregar arquivo:\n{str(e)}")

    def _apply_raw_data(self, df, origem):
        """
        Recebe um DataFrame bruto (Data na 1ª coluna) e prepara todo o estado da
        aplicação: detecta a taxa de referência, popula a lista de ativos e
        ajusta as janelas temporais.

        Compartilhado pelo carregamento de planilha e pela importação do Yahoo
        Finance, para que as duas origens se comportem exatamente igual.
        """
        self.dados_brutos = df

        self.df = None
        self.df_otimizacao = None
        self.df_analise = None
        self.has_risk_free = False
        self.risk_free_column_name = None
        self.detected_risk_free_rate = 0.0
        self.periodo_disponivel = None

        if 'Data' in self.dados_brutos.columns:
            try:
                datas = pd.to_datetime(self.dados_brutos['Data'])
                self.periodo_disponivel = {
                    'inicio': datas.min(),
                    'fim': datas.max(),
                    'total_dias': len(datas)
                }
            except Exception:
                self.periodo_disponivel = None

        # Taxa de referência: 2ª coluna cujo nome contenha um dos termos-chave
        if len(self.dados_brutos.columns) > 2 and isinstance(self.dados_brutos.columns[1], str):
            col_name = self.dados_brutos.columns[1].lower()
            if any(term in col_name for term in ['taxa', 'livre', 'risco', 'ibov', 'ref', 'cdi', 'selic']):
                self.has_risk_free = True
                self.risk_free_column_name = self.dados_brutos.columns[1]
                asset_columns = self.dados_brutos.columns[2:].tolist()
            else:
                asset_columns = self.dados_brutos.columns[1:].tolist()
        else:
            asset_columns = self.dados_brutos.columns[1:].tolist()

        if self.periodo_disponivel:
            status_text = (
                f"✅ {origem}\n"
                f"📊 Dimensões: {self.dados_brutos.shape[0]} linhas x {self.dados_brutos.shape[1]} colunas\n"
                f"📅 Período: {self.periodo_disponivel['inicio'].strftime('%d/%m/%Y')} a "
                f"{self.periodo_disponivel['fim'].strftime('%d/%m/%Y')}\n"
                f"🗓️ Total: {self.periodo_disponivel['total_dias']} dias"
            )
        else:
            status_text = (f"✅ {origem}\n"
                           f"📊 Dimensões: {self.dados_brutos.shape[0]} linhas x "
                           f"{self.dados_brutos.shape[1]} colunas")

        self.status_label.setText(status_text)

        if self.has_risk_free:
            self.risk_free_info.setText(f"✅ Detectada: '{self.risk_free_column_name}'")
            self.risk_free_info.setStyleSheet("color: green;")
            self.manual_risk_entry.setEnabled(False)
            self.update_objective_options()
        else:
            self.risk_free_info.setText("❌ Nenhuma taxa detectada")
            self.risk_free_info.setStyleSheet("color: red;")
            self.manual_risk_entry.setEnabled(True)
            self.update_objective_options()

        self.assets_listbox.delete(0, END)
        for asset in asset_columns:
            self.assets_listbox.insert(END, asset)

        self.update_advanced_widgets()
        self.atualizar_datas_automaticas()
        self._update_selection_info()

    def open_yahoo_import(self):
        """Abre o diálogo de importação de cotações do Yahoo Finance."""
        if not YFINANCE_OK:
            messagebox.showerror(
                "yfinance não instalado",
                "A importação pelo Yahoo Finance exige a biblioteca 'yfinance'.\n\n"
                "Instale com:\n    pip install yfinance\n\n"
                "Depois reinicie o aplicativo."
            )
            return

        self._run_import_dialog(YahooImportDialog(self), "Yahoo Finance")

    def open_b3_import(self):
        """Abre o diálogo de importação de cotações da B3 (COTAHIST)."""
        if not B3_OK:
            messagebox.showerror(
                "Importação da B3 indisponível",
                "Não foi possível carregar o módulo de importação da B3.\n\n"
                f"Motivo: {B3_IMPORT_ERROR}\n\n"
                "Verifique:\n"
                "• se o arquivo 'b3_series_wide_com_limpeza.py' está na MESMA "
                "pasta que 'desktop_app_qt.py';\n"
                "• se a biblioteca 'requests' está instalada no mesmo Python "
                "(pip install requests).\n\n"
                "Depois reinicie o aplicativo.")
            return
        self._run_import_dialog(B3ImportDialog(self), "B3")

    def open_excel_rf_import(self):
        """Abre o diálogo de importação via Excel/STOCKHISTORY (renda fixa)."""
        if not EXCELRF_OK:
            messagebox.showerror(
                "Importação via Excel indisponível",
                "Não foi possível carregar o módulo de importação via Excel.\n\n"
                f"Motivo: {EXCELRF_IMPORT_ERROR}\n\n"
                "Verifique se o arquivo 'b3_excel_rendafixa.py' está na MESMA "
                "pasta que 'desktop_app_qt.py'.\n\n"
                "Depois reinicie o aplicativo.")
            return
        self._run_import_dialog(ExcelRFImportDialog(self), "Excel (renda fixa)")

    def _run_import_dialog(self, dlg, fonte):
        """Executa um diálogo de importação e aplica o resultado, se aceito."""
        if dlg.exec_() != QDialog.Accepted or dlg.result_df is None:
            return
        try:
            self._apply_raw_data(dlg.result_df, dlg.origem_desc)
            msg = f"📥 {len(dlg.result_df.columns) - 1} séries importadas ({fonte})!"
            if dlg.nome_ref:
                msg += f"\n🏛️ Taxa de referência: '{dlg.nome_ref}'"

            # Relatório de limpeza (só as fontes B3/Excel devolvem)
            rel = getattr(dlg, 'relatorio', None)
            if rel:
                if rel.get('regra1'):
                    msg += f"\n🧹 Excluídos (sem dado na 1ª data): {', '.join(rel['regra1'])}"
                if rel.get('regra2'):
                    msg += (f"\n🧹 Excluídos (>{rel.get('k')} dias seguidos sem dado): "
                            f"{', '.join(rel['regra2'])}")

            if dlg.erros:
                msg += f"\n\n⚠️ Sem dados para: {', '.join(dlg.erros)}"
            msg += "\n\n🎯 Agora configure as janelas temporais."
            messagebox.showinfo("Sucesso", msg)
        except Exception as e:
            messagebox.showerror("Erro", f"Erro ao aplicar os dados importados:\n{str(e)}")

    def update_objective_options(self):
        """Atualizar opções de objetivo baseado na taxa livre"""
        _clear_layout(self.risk_free_obj_layout)

        if self.has_risk_free:
            sep = QFrame()
            sep.setFrameShape(QFrame.HLine)
            self.risk_free_obj_layout.addWidget(sep)

            lbl = QLabel("Objetivos com Taxa de Referência:")
            lbl.setFont(bold())
            self.risk_free_obj_layout.addWidget(lbl)

            for obj in self.risk_free_objectives:
                rb = QRadioButton(obj)
                self.objective_group.addButton(rb)
                self.risk_free_obj_layout.addWidget(rb)
                self.objective_buttons[obj] = rb
                self.objective_var.add(obj, rb)

    def update_advanced_widgets(self):
        self.update_constraints_widgets()
        self.update_short_widgets()

    # -------------------------------------------------------------------------
    # RESTRIÇÕES INDIVIDUAIS
    # -------------------------------------------------------------------------
    def update_constraints_widgets(self):
        existing_constraints = getattr(self, 'individual_constraints', {}).copy()

        _clear_layout(self.constraints_layout)
        self.constraint_widgets.clear()

        if not self.use_individual_constraints.get():
            self.constraints_layout.addWidget(QLabel("Habilite as restrições individuais acima"))
            return

        selected_assets = self.get_selected_assets()
        if len(selected_assets) < 2:
            self.constraints_layout.addWidget(QLabel("Selecione pelo menos 2 ativos na aba Dados"))
            return

        lbl = QLabel("Configure limites específicos por ativo:")
        lbl.setFont(bold())
        self.constraints_layout.addWidget(lbl)

        btn = QPushButton("📋 Configurar Restrições Individuais")
        btn.clicked.connect(lambda: self.open_constraints_window(selected_assets))
        self.constraints_layout.addWidget(btn)

        self.constraints_summary_frame = QGroupBox("Restrições Configuradas")
        self.constraints_summary_layout = QVBoxLayout(self.constraints_summary_frame)
        self.constraints_layout.addWidget(self.constraints_summary_frame, 1)

        if existing_constraints and self.use_individual_constraints.get():
            self.individual_constraints = existing_constraints
            try:
                self.update_constraints_summary()
            except Exception:
                pass
        else:
            self.constraints_summary_layout.addWidget(QLabel("Nenhuma restrição individual configurada"))

    def open_constraints_window(self, selected_assets):
        popup = QDialog(self)
        popup.setWindowTitle("Configuração de Restrições Individuais")
        popup.resize(700, 600)
        popup.setModal(True)
        main_l = QVBoxLayout(popup)

        lbl = QLabel("Configure limites específicos por ativo:")
        lbl.setFont(bold(12))
        main_l.addWidget(lbl)

        global_min = self.min_weight_var.get()
        global_max = self.max_weight_var.get()

        info_box = QGroupBox("Limites Globais Atuais")
        info_l = QVBoxLayout(info_box)
        info_l.addWidget(QLabel(f"📊 Globais: Mín={global_min:.1f}% | Máx={global_max:.1f}%"))
        main_l.addWidget(info_box)

        search_l = QHBoxLayout()
        search_l.addWidget(QLabel("🔍 Buscar:"))
        search_entry = QLineEdit()
        search_l.addWidget(search_entry)
        main_l.addLayout(search_l)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll_content = QWidget()
        scroll_layout = QVBoxLayout(scroll_content)
        scroll.setWidget(scroll_content)
        main_l.addWidget(scroll, 1)

        constraint_widgets = {}

        def toggle_constraint(asset):
            if asset in constraint_widgets:
                w = constraint_widgets[asset]
                enabled = w['use_var'].get()
                w['min_entry'].setEnabled(enabled)
                w['max_entry'].setEnabled(enabled)

        def create_constraint_widgets():
            _clear_layout(scroll_layout)
            constraint_widgets.clear()

            search_term = search_entry.text().lower()
            filtered_assets = [a for a in selected_assets if search_term in a.lower()]

            for asset in filtered_assets:
                asset_box = QGroupBox(asset)
                asset_l = QVBoxLayout(asset_box)

                chk = QCheckBox("Usar limites específicos")
                chk.setChecked(asset in self.individual_constraints)
                use_var = BoolVar(chk)
                asset_l.addWidget(chk)

                limits_l = QHBoxLayout()

                if asset in self.individual_constraints:
                    initial_min = self.individual_constraints[asset]['min'] * 100
                    initial_max = self.individual_constraints[asset]['max'] * 100
                else:
                    initial_min = global_min
                    initial_max = global_max

                min_col = QVBoxLayout()
                min_col.addWidget(QLabel("Mín (%):"))
                e_min = QLineEdit()
                min_var = NumVar(e_min, initial_min)
                min_col.addWidget(e_min)
                limits_l.addLayout(min_col)

                max_col = QVBoxLayout()
                max_col.addWidget(QLabel("Máx (%):"))
                e_max = QLineEdit()
                max_var = NumVar(e_max, initial_max)
                max_col.addWidget(e_max)
                limits_l.addLayout(max_col)

                asset_l.addLayout(limits_l)
                scroll_layout.addWidget(asset_box)

                e_min.setEnabled(chk.isChecked())
                e_max.setEnabled(chk.isChecked())

                chk.toggled.connect(lambda _c, a=asset: toggle_constraint(a))

                constraint_widgets[asset] = {
                    'use_var': use_var, 'min_var': min_var, 'max_var': max_var,
                    'min_entry': e_min, 'max_entry': e_max
                }

            scroll_layout.addStretch()

        search_entry.textChanged.connect(create_constraint_widgets)
        create_constraint_widgets()

        # Ações rápidas
        quick_box = QGroupBox("Ações Rápidas")
        quick_l = QHBoxLayout(quick_box)

        def apply_global_limits():
            for asset, w in constraint_widgets.items():
                if w['use_var'].get():
                    w['min_var'].set(global_min)
                    w['max_var'].set(global_max)

        def clear_all_constraints():
            for asset, w in constraint_widgets.items():
                w['use_var'].set(False)
                toggle_constraint(asset)

        def apply_batch_limits():
            batch_min = batch_min_var.get()
            batch_max = batch_max_var.get()
            if batch_min > batch_max:
                messagebox.showerror("Erro", "Mínimo deve ser menor que máximo!")
                return
            for asset, w in constraint_widgets.items():
                if w['use_var'].get():
                    w['min_var'].set(batch_min)
                    w['max_var'].set(batch_max)

        b1 = QPushButton("Aplicar Globais")
        b1.clicked.connect(apply_global_limits)
        quick_l.addWidget(b1)
        b2 = QPushButton("Limpar Todos")
        b2.clicked.connect(clear_all_constraints)
        quick_l.addWidget(b2)

        quick_l.addWidget(QLabel("Lote - Mín:"))
        e_bmin = QLineEdit()
        e_bmin.setFixedWidth(60)
        batch_min_var = NumVar(e_bmin, global_min)
        quick_l.addWidget(e_bmin)
        quick_l.addWidget(QLabel("Máx:"))
        e_bmax = QLineEdit()
        e_bmax.setFixedWidth(60)
        batch_max_var = NumVar(e_bmax, global_max)
        quick_l.addWidget(e_bmax)
        b3 = QPushButton("Aplicar aos Selecionados")
        b3.clicked.connect(apply_batch_limits)
        quick_l.addWidget(b3)

        main_l.addWidget(quick_box)

        def apply_constraints():
            new_constraints = {}
            for asset, w in constraint_widgets.items():
                if w['use_var'].get():
                    min_val = w['min_var'].get() / 100
                    max_val = w['max_var'].get() / 100
                    if min_val > max_val:
                        messagebox.showerror("Erro", f"Mínimo > máximo para {asset}!")
                        return
                    new_constraints[asset] = {'min': min_val, 'max': max_val}

            self.individual_constraints = new_constraints
            self.update_constraints_summary()
            popup.accept()

            if new_constraints:
                messagebox.showinfo("Sucesso", f"Configuradas restrições para {len(new_constraints)} ativos!")
            else:
                messagebox.showinfo("Info", "Todas as restrições foram removidas!")

        btn_l = QHBoxLayout()
        btn_l.addStretch()
        b_cancel = QPushButton("❌ Cancelar")
        b_cancel.clicked.connect(popup.reject)
        btn_l.addWidget(b_cancel)
        b_apply = QPushButton("✅ Aplicar Configuração")
        b_apply.clicked.connect(apply_constraints)
        btn_l.addWidget(b_apply)
        main_l.addLayout(btn_l)

        popup.exec_()

    def update_constraints_summary(self):
        _clear_layout(self.constraints_summary_layout)

        if not self.individual_constraints:
            self.constraints_summary_layout.addWidget(QLabel("Nenhuma restrição individual configurada"))
            self.constraints_summary_layout.addStretch()
            return

        for asset, limits in self.individual_constraints.items():
            self.constraints_summary_layout.addWidget(
                QLabel(f"📊 {asset}: {limits['min']*100:.1f}% - {limits['max']*100:.1f}%"))

        total = QLabel(f"Total: {len(self.individual_constraints)} ativos com limites específicos")
        total.setFont(bold())
        self.constraints_summary_layout.addWidget(total)
        self.constraints_summary_layout.addStretch()

    # -------------------------------------------------------------------------
    # SHORT SELLING
    # -------------------------------------------------------------------------
    def update_short_widgets(self):
        existing_short_weights = getattr(self, 'short_weights', {}).copy()

        _clear_layout(self.short_layout)
        self.short_widgets.clear()

        if not self.use_short.get():
            self.short_layout.addWidget(QLabel("Habilite posições short acima"))
            return

        if self.df is None:
            self.short_layout.addWidget(QLabel("Carregue dados primeiro"))
            return

        selected_assets = self.get_selected_assets()
        if len(selected_assets) < 1:
            self.short_layout.addWidget(QLabel("Selecione ativos principais na aba Dados"))
            return

        all_assets = list(self.assets_listbox.get(0, END))
        available_for_short = [a for a in all_assets if a not in selected_assets]

        if len(available_for_short) == 0:
            self.short_layout.addWidget(
                QLabel("Selecione menos ativos principais para liberar opções de short"))
            return

        lbl = QLabel("Configure posições short (venda a descoberto):")
        lbl.setFont(bold())
        self.short_layout.addWidget(lbl)

        btn = QPushButton("📋 Selecionar Ativos para Short")
        btn.clicked.connect(lambda: self.open_short_selection_window(available_for_short))
        self.short_layout.addWidget(btn)

        self.short_summary_frame = QGroupBox("Ativos Short Configurados")
        self.short_summary_layout = QVBoxLayout(self.short_summary_frame)
        self.short_layout.addWidget(self.short_summary_frame, 1)
        self.short_summary_layout.addWidget(QLabel("Nenhum ativo short configurado"))

        if existing_short_weights and self.use_short.get():
            self.short_weights = existing_short_weights
            try:
                self.update_short_summary()
            except Exception:
                pass

    def open_short_selection_window(self, available_assets):
        popup = QDialog(self)
        popup.setWindowTitle("Configuração de Ativos Short")
        popup.resize(600, 700)
        popup.setModal(True)
        main_l = QVBoxLayout(popup)

        lbl = QLabel("Configure posições short individuais:")
        lbl.setFont(bold(12))
        main_l.addWidget(lbl)

        search_l = QHBoxLayout()
        search_l.addWidget(QLabel("🔍 Buscar:"))
        search_entry = QLineEdit()
        search_l.addWidget(search_entry)
        main_l.addLayout(search_l)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll_content = QWidget()
        scroll_layout = QVBoxLayout(scroll_content)
        scroll.setWidget(scroll_content)
        main_l.addWidget(scroll, 1)

        asset_widgets = {}

        def toggle_asset_weight(asset):
            if asset in asset_widgets:
                w = asset_widgets[asset]
                w['weight_entry'].setEnabled(w['use_var'].get())

        def create_asset_widgets():
            _clear_layout(scroll_layout)
            asset_widgets.clear()

            search_term = search_entry.text().lower()
            filtered_assets = [a for a in available_assets if search_term in a.lower()]

            for asset in filtered_assets:
                row = QHBoxLayout()

                chk = QCheckBox(asset)
                chk.setChecked(asset in self.short_weights)
                use_var = BoolVar(chk)
                row.addWidget(chk)
                row.addStretch()

                row.addWidget(QLabel("Peso (%):"))
                e_w = QLineEdit()
                e_w.setFixedWidth(80)
                initial_weight = self.short_weights.get(asset, -1.0) * 100
                weight_var = NumVar(e_w, initial_weight)
                row.addWidget(e_w)

                e_w.setEnabled(chk.isChecked())
                chk.toggled.connect(lambda _c, a=asset: toggle_asset_weight(a))

                scroll_layout.addLayout(row)

                asset_widgets[asset] = {
                    'use_var': use_var, 'weight_var': weight_var, 'weight_entry': e_w
                }

            scroll_layout.addStretch()

        search_entry.textChanged.connect(create_asset_widgets)
        create_asset_widgets()

        # Ações rápidas
        quick_box = QGroupBox("Ações Rápidas")
        quick_l = QHBoxLayout(quick_box)

        def select_all(select):
            for asset, w in asset_widgets.items():
                w['use_var'].set(select)
                toggle_asset_weight(asset)

        def apply_default_weight():
            dw = default_weight_var.get()
            for asset, w in asset_widgets.items():
                if w['use_var'].get():
                    w['weight_var'].set(dw)

        b_all = QPushButton("Selecionar Todos")
        b_all.clicked.connect(lambda: select_all(True))
        quick_l.addWidget(b_all)
        b_clear = QPushButton("Limpar Todos")
        b_clear.clicked.connect(lambda: select_all(False))
        quick_l.addWidget(b_clear)

        quick_l.addWidget(QLabel("Peso padrão:"))
        e_dw = QLineEdit()
        e_dw.setFixedWidth(80)
        default_weight_var = NumVar(e_dw, -10.0)
        quick_l.addWidget(e_dw)
        b_apply_dw = QPushButton("Aplicar aos Selecionados")
        b_apply_dw.clicked.connect(apply_default_weight)
        quick_l.addWidget(b_apply_dw)

        main_l.addWidget(quick_box)

        def apply_configuration():
            new_short_weights = {}
            for asset, w in asset_widgets.items():
                if w['use_var'].get():
                    weight = w['weight_var'].get() / 100
                    if weight >= 0:
                        messagebox.showerror("Erro", f"Peso short deve ser negativo para {asset}!")
                        return
                    new_short_weights[asset] = weight

            self.short_weights = new_short_weights
            self.update_short_summary()
            popup.accept()
            messagebox.showinfo("Sucesso", f"Configurados {len(new_short_weights)} ativos short!")

        btn_l = QHBoxLayout()
        btn_l.addStretch()
        b_cancel = QPushButton("❌ Cancelar")
        b_cancel.clicked.connect(popup.reject)
        btn_l.addWidget(b_cancel)
        b_ok = QPushButton("✅ Aplicar Configuração")
        b_ok.clicked.connect(apply_configuration)
        btn_l.addWidget(b_ok)
        main_l.addLayout(btn_l)

        popup.exec_()

    def update_short_summary(self):
        _clear_layout(self.short_summary_layout)

        if not self.short_weights:
            self.short_summary_layout.addWidget(QLabel("Nenhum ativo short configurado"))
            self.short_summary_layout.addStretch()
            return

        for asset, weight in self.short_weights.items():
            self.short_summary_layout.addWidget(QLabel(f"📉 {asset}: {weight*100:.1f}%"))

        total_short = sum(self.short_weights.values()) * 100
        total = QLabel(f"Total Short: {total_short:.1f}%")
        total.setFont(bold())
        self.short_summary_layout.addWidget(total)
        self.short_summary_layout.addStretch()

    def toggle_individual_constraints(self):
        self.update_constraints_widgets()

    def toggle_short_selling(self):
        self.update_short_widgets()

    def get_individual_constraints(self):
        if not self.use_individual_constraints.get():
            return None

        # Habilitado mas sem nenhuma restrição configurada (ex.: após "Limpar Todos"):
        # devolver dict vazio -> otimizador usa os limites globais e NÃO aborta.
        if not hasattr(self, 'individual_constraints') or not self.individual_constraints:
            return {}

        for asset, limits in self.individual_constraints.items():
            min_val = limits['min']
            max_val = limits['max']
            if min_val > max_val:
                messagebox.showerror("Erro", f"Peso mínimo > máximo para {asset}")
                return None

        return self.individual_constraints.copy()

    def get_short_configuration(self):
        if not self.use_short.get():
            return [], {}

        if not hasattr(self, 'short_weights') or not self.short_weights:
            return [], {}

        short_assets = list(self.short_weights.keys())
        short_weights = self.short_weights.copy()

        for asset, weight in short_weights.items():
            if weight >= 0:
                messagebox.showerror("Erro", f"Peso short deve ser negativo para {asset}")
                return [], {}

        return short_assets, short_weights

    def select_all_assets(self):
        self.assets_listbox.select_set(0, END)
        self.update_advanced_widgets()

    def clear_selection(self):
        self.assets_listbox.selection_clear(0, END)
        self.update_advanced_widgets()

    def get_selected_assets(self):
        selected_indices = self.assets_listbox.curselection()
        return [self.assets_listbox.get(i) for i in selected_indices]

    # -------------------------------------------------------------------------
    # OTIMIZAÇÃO
    # -------------------------------------------------------------------------
    def optimize_portfolio(self):
        if self.df is None:
            messagebox.showerror("Erro", "Carregue um arquivo Excel primeiro!")
            return

        selected_assets = self.get_selected_assets()
        if len(selected_assets) < 2:
            messagebox.showerror("Erro", "Selecione pelo menos 2 ativos!")
            return

        individual_constraints = self.get_individual_constraints()
        if individual_constraints is None and self.use_individual_constraints.get():
            return

        short_assets, short_weights = self.get_short_configuration()
        if short_assets is None:
            return

        progress_window = None
        try:
            progress_window = QDialog(self)
            progress_window.setWindowTitle("Otimizando...")
            progress_window.resize(350, 120)
            progress_window.setModal(True)
            pv = QVBoxLayout(progress_window)
            pv.addWidget(QLabel("🔄 Otimizando portfólio..."))
            bar = QProgressBar()
            bar.setRange(0, 0)  # indeterminado
            pv.addWidget(bar)
            status_label = QLabel("Inicializando...")
            pv.addWidget(status_label)
            progress_window.show()
            QApplication.processEvents()

            all_assets = selected_assets.copy()
            if len(short_assets) > 0:
                all_assets.extend(short_assets)

            status_label.setText("Inicializando otimizador...")
            QApplication.processEvents()

            self.optimizer = PortfolioOptimizer(self.df, all_assets)

            status_label.setText("Configurando parâmetros...")
            QApplication.processEvents()

            objective_map = {
                "Maximizar Sharpe Ratio": 'sharpe',
                "Maximizar Sortino Ratio": 'sortino',
                "Minimizar Risco": 'volatility',
                "Minimizar Under Water": 'under_water',
                "Maximizar Inclinação/[(1-R²)×Vol]": 'hc10',
                "Maximizar Qualidade da Linearidade": 'quality_linear',
                "Maximizar Linearidade do Excesso": 'excess_hc10',
                "Maximizar Sharpe do Excesso": 'excess_sharpe'
            }

            objective_type = objective_map[self.objective_var.get()]
            min_weight = self.min_weight_var.get() / 100
            max_weight = self.max_weight_var.get() / 100

            if self.has_risk_free and hasattr(self.optimizer, 'risk_free_rate_total'):
                risk_free_rate = self.optimizer.risk_free_rate_total
            else:
                risk_free_rate = self.risk_free_var.get() / 100

            # Meta de retorno (opcional), em um de dois modos:
            #  - relativa: excesso mínimo sobre a referência no período
            #  - absoluta: % ao ano, convertida no otimizador para o período
            target_return = None
            target_annual = None
            if self.use_meta.get():
                if self.meta_mode_var.get() == 'absoluta':
                    target_annual = self.meta_abs_var.get() / 100
                else:
                    target_return = self.meta_var.get() / 100

            status_label.setText("Executando otimização...")
            QApplication.processEvents()

            if len(short_assets) > 0:
                self.result = self.optimizer.optimize_portfolio_with_shorts(
                    selected_assets=selected_assets,
                    short_assets=short_assets,
                    short_weights=short_weights,
                    objective_type=objective_type,
                    target_return=target_return,
                    target_annual=target_annual,
                    max_weight=max_weight,
                    min_weight=min_weight,
                    risk_free_rate=risk_free_rate,
                    individual_constraints=individual_constraints
                )
            else:
                self.result = self.optimizer.optimize_portfolio(
                    objective_type=objective_type,
                    target_return=target_return,
                    target_annual=target_annual,
                    max_weight=max_weight,
                    min_weight=min_weight,
                    risk_free_rate=risk_free_rate,
                    individual_constraints=individual_constraints
                )

            progress_window.close()
            progress_window = None

            if self.result['success']:
                self.display_results()
                self.display_monthly_tables()
                if self.export_csv_btn is not None:
                    self.export_csv_btn.setEnabled(True)
                if self.export_excel_btn is not None:
                    self.export_excel_btn.setEnabled(True)
                if self.result.get('degraded'):
                    # Degradação graciosa: houve resultado, mas com ressalvas
                    messagebox.showwarning(
                        "Otimização concluída com ressalvas",
                        "⚠️ Otimização concluída, mas SEM convergência plena.\n\n"
                        f"{self.result['degraded_message']}"
                        + self._meta_message()
                    )
                else:
                    messagebox.showinfo("Sucesso", "🎉 Otimização concluída com sucesso!" + self._meta_message())
                self.notebook.setCurrentWidget(self.tab_results)
            else:
                messagebox.showerror("Erro", f"❌ {self.result['message']}")

        except Exception as e:
            if progress_window is not None:
                progress_window.close()
            messagebox.showerror("Erro", f"Erro durante otimização:\n{str(e)}")

    def _meta_message(self):
        """Texto sobre o resultado da meta (vazio se meta não foi usada)."""
        if not self.result or not self.result.get('meta_used'):
            return ""
        req = self.result['meta_required'] * 100
        ach = self.result['meta_achieved'] * 100

        # Como o alvo do período foi obtido, conforme o modo escolhido
        if self.result.get('meta_mode') == 'absoluta':
            anual = (self.result.get('meta_annual') or 0) * 100
            ach_anual = (self.result.get('meta_achieved_annual') or 0) * 100
            origem = f"alvo = {anual:.1f}% ao ano → {req:.2f}% no período"
            extra = f"\nRetorno anualizado obtido: {ach_anual:.2f}%."
        else:
            meta = (self.result.get('meta_target') or 0) * 100
            ref = (self.result.get('meta_ref') or 0) * 100
            origem = f"alvo = referência {ref:.2f}% × (1+{meta:.1f}%) = {req:.2f}%"
            extra = ""

        if self.result.get('meta_atingida'):
            return (f"\n\n🎯 Meta atingida: retorno do período = {ach:.2f}% "
                    f"({origem}).{extra}")
        return (f"\n\n⚠️ Meta NÃO atingida com os limites atuais.\n"
                f"Melhor possível: retorno = {ach:.2f}% ({origem}).{extra}")

    def display_results(self):
        if not self.result or not self.result['success']:
            return

        clear_widget(self.chart_frame)

        main_horizontal = QHBoxLayout()
        self.chart_layout.addLayout(main_horizontal)

        # ---- IN-SAMPLE ----
        in_sample_box = QGroupBox("📊 IN-SAMPLE (Otimização)")
        in_l = QVBoxLayout(in_sample_box)

        metrics = self.result['metrics']

        if hasattr(self, 'periodo_otimizacao') and self.periodo_otimizacao:
            data_inicio_otim = self.periodo_otimizacao['inicio']
            data_fim_otim = self.periodo_otimizacao['fim']
            n_dias_otim = (data_fim_otim - data_inicio_otim).days
        else:
            n_dias_otim = len(self.optimizer.returns_data)

        risk_free_annual = (1 + metrics['risk_free_rate']) ** (365 / n_dias_otim) - 1
        sharpe_corrected = (metrics['annual_return'] - risk_free_annual) / metrics['volatility']

        excess_annual_in = metrics['annual_return'] - risk_free_annual
        ratio_in = f"{excess_annual_in / risk_free_annual * 100:.1f}%" if risk_free_annual != 0 else "n/d"

        in_sample_text = f"""🎯 RETORNOS:
  • Total: {metrics['gv_final']:.2%}
  • Anualizado: {metrics['annual_return']:.2%}

⚡ PERFORMANCE:
  • Sharpe: {sharpe_corrected:.3f}

📊 RISCO:
  • Volatilidade: {metrics['volatility']:.2%}
  • VaR 95%: {metrics['var_95_daily']:.2%}
  • CVaR 95%: {metrics['cvar_95_daily']:.2%}

📈 QUALIDADE:
  • R²: {metrics['r_squared']:.3f}

🛡️ REFERÊNCIA:
  • Taxa Ref Período: {metrics['risk_free_rate']:.2%}
  • Taxa Ref Anualizada: {risk_free_annual*100:.2f}
  • Excesso Período: {metrics['excess_return']:.2%}
  • Excesso Anualizado: {excess_annual_in:.2%}
  • Excesso/Ref: {ratio_in}

📅 PERÍODO:
  • Dias: {n_dias_otim}"""

        in_l.addWidget(self._metrics_label(in_sample_text))
        main_horizontal.addWidget(in_sample_box, 1)

        # ---- OUT-OF-SAMPLE ----
        out_sample_box = QGroupBox("🔍 OUT-OF-SAMPLE (Validação)")
        out_l = QVBoxLayout(out_sample_box)

        annual_return_valid = sharpe_valid = vol_valid = 0
        var_95_daily_valid = excess_return_valid = 0

        if hasattr(self, 'df_analise') and self.df_analise is not None:
            try:
                selected_assets = self.get_selected_assets()

                if hasattr(self, 'use_short') and self.use_short.get() and hasattr(self, 'short_weights'):
                    short_assets = list(self.short_weights.keys())
                    assets_used_in_optimization = selected_assets + short_assets
                else:
                    assets_used_in_optimization = selected_assets

                optimizer_valid = PortfolioOptimizer(self.df_analise, assets_used_in_optimization)

                n_assets_optimization = len(self.result['weights'])
                n_assets_validation = optimizer_valid.returns_data.shape[1] if len(optimizer_valid.returns_data.shape) > 1 else 1

                if n_assets_optimization == n_assets_validation:
                    portfolio_returns_valid = np.dot(optimizer_valid.returns_data.values, self.result['weights'])
                    cumulative_valid = np.cumsum(portfolio_returns_valid)

                    n_dias_otim = len(self.optimizer.returns_data)

                    if len(portfolio_returns_valid) > n_dias_otim:
                        returns_valid_only = portfolio_returns_valid[n_dias_otim:]

                        if hasattr(self, 'periodo_otimizacao') and hasattr(self, 'periodo_analise'):
                            data_fim_otim = self.periodo_otimizacao['fim']
                            data_fim_analise = self.periodo_analise['fim']
                            n_dias_valid = (data_fim_analise - data_fim_otim).days
                        else:
                            n_dias_valid = len(returns_valid_only)

                        portfolio_total_ate_validacao = cumulative_valid[-1]
                        portfolio_total_ate_otimizacao = cumulative_valid[n_dias_otim - 1]
                        retorno_total_valid = (1 + portfolio_total_ate_validacao) / (1 + portfolio_total_ate_otimizacao) - 1

                        if n_dias_valid > 0:
                            annual_return_valid = (1 + retorno_total_valid) ** (365 / n_dias_valid) - 1
                        else:
                            annual_return_valid = 0

                        portfolio_cumulative_validacao = cumulative_valid[n_dias_otim - 1:]

                        if len(portfolio_cumulative_validacao) > 1:
                            portfolio_cumulative_with_zero = np.concatenate(
                                [[portfolio_cumulative_validacao[0]], portfolio_cumulative_validacao])
                            variac_result_pu = (1 + portfolio_cumulative_with_zero[1:]) / (1 + portfolio_cumulative_with_zero[:-1])
                            portfolio_returns_pct_valid = variac_result_pu - 1
                            vol_valid = np.std(portfolio_returns_pct_valid, ddof=0) * np.sqrt(252)
                        else:
                            vol_valid = 0

                        if hasattr(optimizer_valid, 'risk_free_cumulative') and optimizer_valid.risk_free_cumulative is not None:
                            try:
                                risk_free_total_ate_validacao = optimizer_valid.risk_free_cumulative.iloc[-1]
                                risk_free_total_ate_otimizacao = optimizer_valid.risk_free_cumulative.iloc[n_dias_otim - 1]
                                risk_free_total_valid = (1 + risk_free_total_ate_validacao) / (1 + risk_free_total_ate_otimizacao) - 1

                                if n_dias_valid > 0:
                                    risk_free_annual_valid = (1 + risk_free_total_valid) ** (365 / n_dias_valid) - 1
                                else:
                                    risk_free_annual_valid = 0
                            except Exception:
                                risk_free_annual_valid = 0
                        else:
                            risk_free_annual_valid = 0

                        if vol_valid > 0:
                            sharpe_valid = (annual_return_valid - risk_free_annual_valid) / vol_valid
                        else:
                            sharpe_valid = 0

                        mean_daily_return = np.mean(portfolio_returns_pct_valid)
                        std_daily_return = np.std(portfolio_returns_pct_valid, ddof=0)
                        var_95_daily_valid = mean_daily_return - 1.65 * std_daily_return

                        excess_return_valid = annual_return_valid - risk_free_annual_valid
                        ratio_valid = f"{excess_return_valid / risk_free_annual_valid * 100:.1f}%" if risk_free_annual_valid != 0 else "n/d"

                        out_sample_text = f"""🎯 RETORNOS:
  • Total: {retorno_total_valid:.2%}
  • Anualizado: {annual_return_valid:.2%}

⚡ PERFORMANCE:
  • Sharpe: {sharpe_valid:.3f}

📊 RISCO:
  • Volatilidade: {vol_valid:.2%}
  • VaR 95%: {var_95_daily_valid:.2%}

🛡️ REFERÊNCIA:
  • Taxa Ref Anualizada: {risk_free_annual_valid:.2%}
  • Excesso Anualizado: {excess_return_valid:.2%}
  • Excesso/Ref: {ratio_valid}

📅 PERÍODO:
  • Dias: {n_dias_valid}"""
                    else:
                        out_sample_text = "⚠️ Período de validação muito curto\nconfigurado."
                        annual_return_valid = sharpe_valid = vol_valid = 0
                        var_95_daily_valid = excess_return_valid = 0
                else:
                    out_sample_text = f"❌ Incompatibilidade:\n{n_assets_optimization} pesos vs {n_assets_validation} ativos"
                    annual_return_valid = sharpe_valid = vol_valid = 0
                    var_95_daily_valid = excess_return_valid = 0
            except Exception as e:
                out_sample_text = f"❌ Erro na validação:\n{str(e)}"
                annual_return_valid = sharpe_valid = vol_valid = 0
                var_95_daily_valid = excess_return_valid = 0
        else:
            out_sample_text = "🔍 Configure um período de\nvalidação para ver resultados\nout-of-sample"
            annual_return_valid = sharpe_valid = vol_valid = 0
            var_95_daily_valid = excess_return_valid = 0

        out_l.addWidget(self._metrics_label(out_sample_text))
        main_horizontal.addWidget(out_sample_box, 1)

        # ---- COMPARAÇÃO ----
        comparison_box = QGroupBox("⚖️ COMPARAÇÃO")
        comp_l = QVBoxLayout(comparison_box)

        if annual_return_valid != 0:
            diff_return = annual_return_valid - metrics['annual_return']
            diff_sharpe = sharpe_valid - metrics['sharpe_ratio']
            diff_vol = vol_valid - metrics['volatility']

            better_return = "📈 Out" if diff_return > 0 else "📉 In"
            better_sharpe = "📈 Out" if diff_sharpe > 0 else "📉 In"
            better_vol = "📈 In" if diff_vol > 0 else "📉 Out"

            comparison_text = f"""📊 DIFERENÇAS (Out - In):

🎯 RETORNO:
  • Anual: {diff_return:+.2%}
  • Melhor: {better_return}

⚡ SHARPE:
  • Diferença: {diff_sharpe:+.3f}
  • Melhor: {better_sharpe}

📊 VOLATILIDADE:
  • Diferença: {diff_vol:+.2%}
  • Menor risco: {better_vol}

💡 RESUMO:
{"✅ Out-of-sample superior" if diff_sharpe > 0 else "⚠️ Degradação out-of-sample"}

🔍 ANÁLISE:
{"Boa generalização!" if abs(diff_sharpe) < 0.1 else "Cuidado: overfitting?"}"""
        else:
            comparison_text = """⚖️ COMPARAÇÃO

Configure um período de validação
para ver a comparação entre
in-sample e out-of-sample.

Isso ajuda a detectar:
  • Overfitting
  • Generalização
  • Robustez da estratégia"""

        comp_l.addWidget(self._metrics_label(comparison_text))
        main_horizontal.addWidget(comparison_box, 1)

        # ---- COMPOSIÇÃO ----
        self.display_portfolio_composition()

    @staticmethod
    def _metrics_label(text):
        lbl = QLabel(text)
        lbl.setAlignment(Qt.AlignTop | Qt.AlignLeft)
        lbl.setTextInteractionFlags(Qt.TextSelectableByMouse)
        f = QFont("Monospace")
        f.setStyleHint(QFont.TypeWriter)
        lbl.setFont(f)
        return lbl

    def display_portfolio_composition(self):
        composition_box = QGroupBox("📋 Composição do Portfólio")
        comp_main = QVBoxLayout(composition_box)

        export_l = QHBoxLayout()
        self.export_csv_btn = QPushButton("💾 Exportar CSV")
        self.export_csv_btn.clicked.connect(self.export_csv)
        export_l.addWidget(self.export_csv_btn)
        self.export_excel_btn = QPushButton("📊 Exportar Excel")
        self.export_excel_btn.clicked.connect(self.export_excel)
        export_l.addWidget(self.export_excel_btn)
        export_l.addStretch()
        comp_main.addLayout(export_l)

        comp_horizontal = QHBoxLayout()
        comp_main.addLayout(comp_horizontal)

        # Composição (esquerda)
        columns = ('Ativo', 'Peso Inicial (%)', 'Peso Atual (%)', 'Tipo')
        self.portfolio_tree = QTableWidget(0, len(columns))
        self.portfolio_tree.setHorizontalHeaderLabels(columns)
        self.portfolio_tree.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.portfolio_tree.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.portfolio_tree.setMinimumHeight(280)
        comp_horizontal.addWidget(self.portfolio_tree, 1)

        self.populate_portfolio_tree()

        # Gráfico (direita)
        chart_right = QGroupBox("📈 Evolução do Portfólio")
        QVBoxLayout(chart_right)
        comp_horizontal.addWidget(chart_right, 1)

        self.chart_layout.addWidget(composition_box)

        self.create_performance_chart_in_frame(chart_right)

    def get_portfolio_summary_desktop(self, weights, assets, returns_data):
        """Cria resumo do portfólio (IDÊNTICO À VERSÃO TKINTER)"""
        asset_final_values = []

        for i, asset in enumerate(assets):
            if abs(weights[i]) > 0.001:
                asset_returns = returns_data[asset].values
                asset_cumulative_return = np.sum(asset_returns)
                asset_final_value = weights[i] * (1 + asset_cumulative_return)
                asset_final_values.append(asset_final_value)
            else:
                asset_final_values.append(0)

        long_portfolio_value = sum([val for i, val in enumerate(asset_final_values) if weights[i] > 0])

        asset_current_weights = []
        for i, asset in enumerate(assets):
            if abs(weights[i]) > 0.001:
                current_weight = asset_final_values[i] / long_portfolio_value if long_portfolio_value > 0 else 0
                asset_current_weights.append(current_weight)
            else:
                asset_current_weights.append(0)

        significant_weights = np.abs(weights) > 0.001

        portfolio_df = pd.DataFrame({
            'Ativo': np.array(assets)[significant_weights],
            'Peso Inicial (%)': weights[significant_weights] * 100,
            'Peso Atual (%)': np.array(asset_current_weights)[significant_weights] * 100,
            'Tipo': ['SHORT' if w < 0 else 'LONG' for w in weights[significant_weights]]
        }).sort_values('Peso Inicial (%)', key=abs, ascending=False)

        return portfolio_df

    def populate_portfolio_tree(self):
        self.portfolio_tree.setRowCount(0)

        portfolio_summary = self.get_portfolio_summary_desktop(
            self.result['weights'], self.result['assets'], self.optimizer.returns_data)

        total_initial_positive = 0
        total_initial_negative = 0
        total_current_positive = 0
        total_current_negative = 0

        rows = []
        for _, row in portfolio_summary.iterrows():
            peso_inicial = row['Peso Inicial (%)']
            peso_atual = row['Peso Atual (%)']
            tipo = row['Tipo']

            rows.append((row['Ativo'], f"{peso_inicial:.2f}%", f"{peso_atual:.2f}%", tipo))

            if peso_inicial > 0:
                total_initial_positive += peso_inicial
                total_current_positive += peso_atual
            else:
                total_initial_negative += peso_inicial
                total_current_negative += peso_atual

        rows.append(("─────────", "─────────", "─────────", "─────"))
        rows.append(("TOTAL LONG", f"{total_initial_positive:.1f}%", f"{total_current_positive:.1f}%", ""))

        if abs(total_initial_negative) > 0.001:
            rows.append(("TOTAL SHORT", f"{total_initial_negative:.1f}%", f"{total_current_negative:.1f}%", ""))

        if abs(total_current_positive - total_initial_positive) > 0.1:
            diff_long = total_current_positive - total_initial_positive
            rows.append(("", "", "", ""))
            rows.append(("VARIAÇÃO LONG", "", f"{diff_long:+.1f}%", "🔄" if abs(diff_long) > 1 else "✅"))

        for r, values in enumerate(rows):
            self.portfolio_tree.insertRow(r)
            for c, val in enumerate(values):
                item = QTableWidgetItem(str(val))
                item.setTextAlignment(Qt.AlignCenter)
                if values[3] == 'SHORT':
                    item.setForeground(QColor("#b00020"))
                self.portfolio_tree.setItem(r, c, item)

    def create_performance_chart_in_frame(self, parent_frame):
        """Criar gráfico de performance completo - EIXO X IGUAL AO STREAMLIT"""
        clear_widget(parent_frame)

        usar_grafico_completo = False
        fig = None

        if hasattr(self, 'df_analise') and self.df_analise is not None:
            try:
                selected_assets = self.get_selected_assets()

                if hasattr(self, 'use_short') and self.use_short.get() and hasattr(self, 'short_weights'):
                    short_assets = list(self.short_weights.keys())
                    assets_used_in_optimization = selected_assets + short_assets
                else:
                    assets_used_in_optimization = selected_assets

                optimizer_extended = PortfolioOptimizer(self.df_analise, assets_used_in_optimization)

                if self.has_risk_free and hasattr(self.optimizer, 'risk_free_rate_total'):
                    final_risk_free_rate = self.optimizer.risk_free_rate_total
                else:
                    final_risk_free_rate = self.risk_free_var.get() / 100

                metrics_extended = optimizer_extended.calculate_portfolio_metrics(
                    self.result['weights'], final_risk_free_rate)

                dates_extended = getattr(optimizer_extended, 'dates', None)
                n_dias_otim = len(self.optimizer.returns_data)

                fig = Figure(figsize=(8, 5))
                ax = fig.add_subplot(111)

                if dates_extended is not None:
                    x_data = dates_extended
                else:
                    x_data = list(range(len(metrics_extended['portfolio_cumulative'])))

                portfolio_cumulative_extended = metrics_extended['portfolio_cumulative'] * 100

                ax.plot(x_data, portfolio_cumulative_extended, 'b-', linewidth=2.5, label='Portfólio Otimizado')

                if hasattr(optimizer_extended, 'risk_free_cumulative') and optimizer_extended.risk_free_cumulative is not None:
                    risk_free_cumulative_extended = optimizer_extended.risk_free_cumulative * 100
                    ax.plot(x_data, risk_free_cumulative_extended, color='orange', linestyle='--',
                            linewidth=2, label='Taxa de Referência')

                    if metrics_extended.get('excess_cumulative') is not None:
                        excess_cumulative_extended = metrics_extended['excess_cumulative'] * 100
                        ax.plot(x_data, excess_cumulative_extended, 'g:', linewidth=2, label='Excesso de Retorno')

                if len(x_data) > n_dias_otim:
                    if dates_extended is not None:
                        fim_otim_point = x_data.iloc[n_dias_otim - 1] if hasattr(x_data, 'iloc') else x_data[n_dias_otim - 1]
                    else:
                        fim_otim_point = n_dias_otim - 1
                    ax.axvline(x=fim_otim_point, color='red', linestyle='--', linewidth=2, alpha=0.8,
                               label='Fim Otimização')

                ax.set_title('Evolução do Retorno Acumulado - Período Completo', fontsize=10)
                ax.set_ylabel('Retorno Acumulado (%)')
                ax.legend(loc='upper left', fontsize=8)
                ax.grid(True, alpha=0.3)

                if dates_extended is not None:
                    try:
                        total_pontos = len(dates_extended)
                        n_ticks = 9

                        if total_pontos > n_ticks:
                            indices = np.linspace(0, total_pontos - 1, n_ticks, dtype=int)
                        else:
                            indices = np.arange(total_pontos)

                        datas_selecionadas = [dates_extended.iloc[i] for i in indices]
                        ax.set_xticks(datas_selecionadas)

                        formato = mdates.DateFormatter('%d/%m/%Y')
                        ax.xaxis.set_major_formatter(formato)

                        for lbl in ax.xaxis.get_majorticklabels():
                            lbl.set_rotation(0)
                            lbl.set_ha('center')
                            lbl.set_fontsize(8)
                    except Exception:
                        ax.locator_params(axis='x', nbins=7)

                usar_grafico_completo = True

            except Exception:
                usar_grafico_completo = False

        if not usar_grafico_completo:
            fig = Figure(figsize=(8, 5))
            ax = fig.add_subplot(111)

            periods = range(1, len(self.result['metrics']['portfolio_cumulative']) + 1)
            portfolio_cumulative = self.result['metrics']['portfolio_cumulative'] * 100

            ax.plot(periods, portfolio_cumulative, 'b-', linewidth=2.5, label='Portfólio Otimizado')
            ax.set_title('Evolução do Retorno Acumulado - Período de Otimização', fontsize=10)
            ax.set_xlabel('Dias')
            ax.set_ylabel('Retorno Acumulado (%)')
            ax.legend()
            ax.grid(True, alpha=0.3)

        fig.tight_layout()
        canvas = FigureCanvas(fig)
        parent_frame.layout().addWidget(canvas)
        canvas.draw()

    def create_performance_chart(self):
        """Criar gráfico de performance completo (com toolbar)"""
        clear_widget(self.chart_frame)

        fig = Figure(figsize=(10, 6))
        ax = fig.add_subplot(111)

        periods = range(1, len(self.result['metrics']['portfolio_cumulative']) + 1)
        portfolio_cumulative = self.result['metrics']['portfolio_cumulative'] * 100

        ax.plot(periods, portfolio_cumulative, 'b-', linewidth=2.5, label='Portfólio Otimizado')

        if hasattr(self.optimizer, 'risk_free_cumulative') and self.optimizer.risk_free_cumulative is not None:
            risk_free_cumulative = self.optimizer.risk_free_cumulative * 100
            ax.plot(periods, risk_free_cumulative, color='orange', linestyle='--', linewidth=2,
                    label='Taxa de Referência')

            excess = portfolio_cumulative - risk_free_cumulative
            ax.plot(periods, excess, 'g:', linewidth=2, label='Excesso de Retorno')

        ax.set_title('Evolução do Retorno Acumulado', fontsize=14, fontweight='bold')
        ax.set_xlabel('Dias de Negociação')
        ax.set_ylabel('Retorno Acumulado (%)')
        ax.legend(loc='upper left')
        ax.grid(True, alpha=0.3)

        final_return = self.result['metrics']['gv_final']
        ax.annotate(f'Retorno Final: {final_return:.2%}',
                    xy=(len(periods), final_return * 100),
                    xytext=(-60, -30), textcoords='offset points',
                    bbox=dict(boxstyle='round,pad=0.5', fc='yellow', alpha=0.7),
                    arrowprops=dict(arrowstyle='->', connectionstyle='arc3,rad=0'))

        fig.tight_layout()
        canvas = FigureCanvas(fig)
        self.chart_layout.addWidget(canvas)
        toolbar = NavigationToolbar(canvas, self.chart_frame)
        self.chart_layout.addWidget(toolbar)
        canvas.draw()

    # -------------------------------------------------------------------------
    # ABA 8: AUTO-OTIMIZAÇÃO
    # -------------------------------------------------------------------------
    def setup_auto_optimization_tab(self):
        main_l = QVBoxLayout(self.tab_auto)

        # Cabeçalho
        header_box = QGroupBox("🤖 Auto-Otimização Inteligente")
        header_l = QHBoxLayout(header_box)
        lbl = QLabel("Walk-Forward Optimization com Ranking Dinâmico Anti-Overfitting")
        lbl.setFont(bold())
        header_l.addWidget(lbl)
        header_l.addStretch()
        self.status_auto_label = QLabel("Aguardando configuração...")
        self.status_auto_label.setStyleSheet("color: blue;")
        header_l.addWidget(self.status_auto_label)
        main_l.addWidget(header_box)

        content_l = QHBoxLayout()
        main_l.addLayout(content_l, 1)

        left = QVBoxLayout()
        right = QVBoxLayout()
        content_l.addLayout(left, 45)   # 45% esquerda: parâmetros
        content_l.addLayout(right, 55)  # 55% direita: estimativa + tabela de 13 colunas

        # Grid de parâmetros (linhas 0/1 expandem; conteúdo ancorado no topo-esquerda)
        params_grid = QGridLayout()
        params_grid.setColumnStretch(0, 1)
        params_grid.setColumnStretch(1, 1)
        params_grid.setRowStretch(0, 1)
        params_grid.setRowStretch(1, 1)
        left.addLayout(params_grid, 1)

        # 1. Janelas de Otimização
        otim_box = QGroupBox("📅 Janelas de Otimização")
        otim_outer = QVBoxLayout(otim_box)
        otim_grid = QGridLayout()
        otim_grid.setColumnStretch(3, 1)  # empurra checkboxes para a esquerda
        otim_outer.addLayout(otim_grid)
        otim_outer.addStretch()
        self.otim_windows = {}
        for i, (period, default) in enumerate([('3m', False), ('6m', True), ('1a', True),
                                               ('2a', False), ('3a', False)]):
            cb = QCheckBox(period)
            cb.setChecked(default)
            r, c = divmod(i, 3)
            otim_grid.addWidget(cb, r, c)
            self.otim_windows[period] = BoolVar(cb)
        params_grid.addWidget(otim_box, 0, 0)

        # 2. Janelas de Validação / Step
        valid_box = QGroupBox("🔍 Janelas de Validação / Step")
        valid_l = QVBoxLayout(valid_box)
        hint = QLabel("Frequência de rebalanceamento = Período de avaliação")
        hint.setStyleSheet("color: gray;")
        valid_l.addWidget(hint)
        valid_grid = QGridLayout()
        valid_grid.setColumnStretch(2, 1)  # empurra checkboxes para a esquerda
        valid_l.addLayout(valid_grid)
        valid_l.addStretch()
        self.rebalance_periods = {}
        period_labels = {'1sem': '1 semana', '2sem': '2 semanas', '1mes': '1 mês',
                         '2mes': '2 meses', '3mes': '3 meses'}
        for i, (key, default) in enumerate([('1sem', False), ('2sem', False), ('1mes', True),
                                            ('2mes', True), ('3mes', False)]):
            cb = QCheckBox(period_labels[key])
            cb.setChecked(default)
            r, c = divmod(i, 2)
            valid_grid.addWidget(cb, r, c)
            self.rebalance_periods[key] = BoolVar(cb)
        params_grid.addWidget(valid_box, 0, 1)

        # 3. Objetivos
        obj_box = QGroupBox("🎯 Objetivos de Otimização")
        obj_l = QVBoxLayout(obj_box)
        self.objectives = {}
        obj_labels = {'sharpe': 'Maximizar Sharpe', 'sortino': 'Maximizar Sortino',
                      'volatility': 'Minimizar Risco',
                      'under_water': 'Minimizar Under Water',
                      'hc10': 'Maximizar Inc/[(1-R²)×Vol]',
                      'quality_linear': 'Qualidade da Linearidade',
                      'excess_hc10': 'Linearidade do Excesso',
                      'excess_sharpe': 'Sharpe do Excesso'}
        for key, default in [('sharpe', True), ('sortino', False),
                             ('volatility', False), ('under_water', False),
                             ('hc10', False), ('quality_linear', False),
                             ('excess_hc10', False), ('excess_sharpe', False)]:
            cb = QCheckBox(obj_labels[key])
            cb.setChecked(default)
            obj_l.addWidget(cb)
            self.objectives[key] = BoolVar(cb)
        obj_l.addStretch()
        params_grid.addWidget(obj_box, 1, 0)

        # 4. Posições Vendidas
        short_box = QGroupBox("🔻 Posições Vendidas")
        short_l = QVBoxLayout(short_box)
        chk_short = QCheckBox("Habilitar posições short")
        self.use_auto_shorts = BoolVar(chk_short)
        chk_short.toggled.connect(self.toggle_auto_shorts)
        short_l.addWidget(chk_short)

        self.auto_shorts_config = QWidget()
        asc_l = QVBoxLayout(self.auto_shorts_config)
        asc_l.setContentsMargins(0, 0, 0, 0)
        btn_auto_short = QPushButton("📋 Selecionar Ativos para Short")
        btn_auto_short.clicked.connect(self.open_auto_short_selection_window)
        asc_l.addWidget(btn_auto_short)
        self.auto_short_summary = QLabel("Nenhum ativo short configurado")
        self.auto_short_summary.setWordWrap(True)
        self.auto_short_summary.setStyleSheet("color: gray;")
        asc_l.addWidget(self.auto_short_summary)
        short_l.addWidget(self.auto_shorts_config)
        short_l.addStretch()
        self.auto_shorts_config.setEnabled(False)
        params_grid.addWidget(short_box, 1, 1)

        # 5. Configurações Globais
        config_box = QGroupBox("⚙️ Configurações Globais")
        config_l = QHBoxLayout(config_box)
        config_l.setAlignment(Qt.AlignTop)

        cfg_left = QVBoxLayout()
        lbl_rank = QLabel("🏆 Ranking de Ativos (por Score):")
        lbl_rank.setFont(bold())
        cfg_left.addWidget(lbl_rank)
        rank_l = QHBoxLayout()
        rank_l.addWidget(QLabel("Score Min:"))
        e_rank_min = QLineEdit()
        e_rank_min.setFixedWidth(50)
        self.rank_min_var = NumVar(e_rank_min, 0)
        rank_l.addWidget(e_rank_min)
        rank_l.addWidget(QLabel("Max:"))
        e_rank_max = QLineEdit()
        e_rank_max.setFixedWidth(50)
        self.rank_max_var = NumVar(e_rank_max, 100)
        rank_l.addWidget(e_rank_max)
        rank_l.addWidget(QLabel("(0-100)"))
        rank_l.addStretch()
        cfg_left.addLayout(rank_l)
        config_l.addLayout(cfg_left, 1)

        cfg_right = QVBoxLayout()
        lbl_w = QLabel("⚖️ Limites de Peso:")
        lbl_w.setFont(bold())
        cfg_right.addWidget(lbl_w)
        weight_l = QHBoxLayout()
        weight_l.addWidget(QLabel("Min:"))
        e_wmin = QLineEdit()
        e_wmin.setFixedWidth(50)
        self.weight_min_var = NumVar(e_wmin, 0)
        weight_l.addWidget(e_wmin)
        weight_l.addWidget(QLabel("% Max:"))
        e_wmax = QLineEdit()
        e_wmax.setFixedWidth(50)
        self.weight_max_var = NumVar(e_wmax, 30)
        weight_l.addWidget(e_wmax)
        weight_l.addWidget(QLabel("%"))
        weight_l.addStretch()
        cfg_right.addLayout(weight_l)
        config_l.addLayout(cfg_right, 1)

        cfg_meta = QVBoxLayout()
        lbl_m = QLabel("🎯 Meta de Retorno:")
        lbl_m.setFont(bold())
        cfg_meta.addWidget(lbl_m)
        chk_auto_meta = QCheckBox("Exigir meta")
        self.use_auto_meta = BoolVar(chk_auto_meta)
        cfg_meta.addWidget(chk_auto_meta)

        # Mesmos dois modos da aba Configuração
        self.auto_meta_mode_group = QButtonGroup(self)
        self.auto_meta_mode_var = RadioVar(default="relativa")

        meta_l2 = QHBoxLayout()
        rb_auto_rel = QRadioButton("")
        self.auto_meta_mode_group.addButton(rb_auto_rel)
        self.auto_meta_mode_var.add("relativa", rb_auto_rel)
        meta_l2.addWidget(rb_auto_rel)
        e_meta = QLineEdit()
        e_meta.setFixedWidth(50)
        self.auto_meta_var = NumVar(e_meta, 5)
        meta_l2.addWidget(e_meta)
        meta_l2.addWidget(QLabel("% acima da ref."))
        meta_l2.addStretch()
        cfg_meta.addLayout(meta_l2)

        meta_l3 = QHBoxLayout()
        rb_auto_abs = QRadioButton("")
        self.auto_meta_mode_group.addButton(rb_auto_abs)
        self.auto_meta_mode_var.add("absoluta", rb_auto_abs)
        meta_l3.addWidget(rb_auto_abs)
        e_meta_abs = QLineEdit()
        e_meta_abs.setFixedWidth(50)
        self.auto_meta_abs_var = NumVar(e_meta_abs, 15)
        meta_l3.addWidget(e_meta_abs)
        meta_l3.addWidget(QLabel("% ao ano"))
        meta_l3.addStretch()
        cfg_meta.addLayout(meta_l3)
        config_l.addLayout(cfg_meta, 1)

        params_grid.addWidget(config_box, 2, 0, 1, 2)

        # Coluna direita: controles e resultados
        estimate_box = QGroupBox("📊 Estimativa")
        est_l = QVBoxLayout(estimate_box)
        self.estimate_label = QLabel("Configure os parâmetros")
        self.estimate_label.setWordWrap(True)
        est_l.addWidget(self.estimate_label)
        btn_calc = QPushButton("🧮 Calcular")
        btn_calc.clicked.connect(self.calculate_estimate)
        est_l.addWidget(btn_calc)
        right.addWidget(estimate_box)

        self.auto_optimize_btn = QPushButton("🚀 INICIAR AUTO-OTIMIZAÇÃO")
        self.auto_optimize_btn.setStyleSheet(
            "QPushButton { background-color: #0078d4; color: white; font-weight: bold; padding: 8px; }")
        self.auto_optimize_btn.clicked.connect(self.start_auto_optimization)
        right.addWidget(self.auto_optimize_btn)

        results_box = QGroupBox("🏆 Top Configurações")
        res_l = QVBoxLayout(results_box)

        export_l = QHBoxLayout()
        b_csv = QPushButton("💾 Exportar CSV")
        b_csv.clicked.connect(self.export_auto_results_csv)
        export_l.addWidget(b_csv)
        b_xls = QPushButton("📊 Exportar Excel")
        b_xls.clicked.connect(self.export_auto_results_excel)
        export_l.addWidget(b_xls)
        export_l.addStretch()
        res_l.addLayout(export_l)

        # Cabeçalhos curtos de propósito: com 13 colunas, títulos longos são o que
        # estoura a largura (as células em si são curtas). O significado completo
        # de cada um fica no tooltip do cabeçalho.
        columns = ('#', 'Otimização', 'Rebalanc.', 'Objetivo', 'N_Ativos', 'Sharpe',
                   'Ret%', 'Meta%', 'Ref%', 'Vol%', 'VaR%', '>Ref%', '>0%')
        self.auto_results_columns = columns
        self.auto_results_tree = QTableWidget(0, len(columns))
        self.auto_results_tree.setHorizontalHeaderLabels(columns)

        _tips = {
            '#': 'Posição no ranking (ordenado por Sharpe out-of-sample)',
            'Otimização': 'Janela de otimização (in-sample) usada em cada step',
            'Rebalanc.': 'Frequência de rebalanceamento = período de avaliação de cada step',
            'Objetivo': 'Objetivo de otimização usado',
            'N_Ativos': 'Número médio de ativos na carteira por step',
            'Sharpe': 'Sharpe médio obtido FORA da amostra (validação)',
            'Ret%': 'Retorno anualizado obtido FORA da amostra (validação)',
            'Meta%': '% dos steps que cumpriram a Meta DENTRO da janela de otimização.\n'
                     '100% com Ret% abaixo do alvo = meta batida in-sample que não se\n'
                     'sustentou fora da amostra (restrição sem folga).\n'
                     'Abaixo de 100% = alvo inatingível em alguns steps, valendo o\n'
                     'fallback de maior retorno possível.\n'
                     '"—" = meta não utilizada.',
            'Ref%': 'Taxa de referência anualizada do período',
            'Vol%': 'Volatilidade anualizada',
            'VaR%': 'VaR 95% diário médio dos steps (risco de cauda)',
            '>Ref%': '% de steps cujo retorno superou a taxa de referência do período',
            '>0%': '% de steps com retorno absoluto positivo',
        }
        for _c, _name in enumerate(columns):
            if _name in _tips:
                self.auto_results_tree.horizontalHeaderItem(_c).setToolTip(_tips[_name])
        _auto_header = self.auto_results_tree.horizontalHeader()
        _auto_header.setSectionResizeMode(QHeaderView.ResizeToContents)  # colunas ajustam ao conteúdo
        self.auto_results_tree.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.auto_results_tree.setSortingEnabled(True)                   # ordenável ao clicar no cabeçalho
        res_l.addWidget(self.auto_results_tree, 1)

        right.addWidget(results_box, 1)

    def toggle_auto_shorts(self):
        self.auto_shorts_config.setEnabled(self.use_auto_shorts.get())

    def update_auto_short_summary(self):
        if not self.auto_short_weights:
            self.auto_short_summary.setText("Nenhum ativo short configurado")
            self.auto_short_summary.setStyleSheet("color: gray;")
            return
        linhas = [f"📉 {a}: {w*100:.1f}%" for a, w in self.auto_short_weights.items()]
        total = sum(self.auto_short_weights.values()) * 100
        self.auto_short_summary.setText("\n".join(linhas) + f"\nTotal Short: {total:.1f}%")
        self.auto_short_summary.setStyleSheet("")

    def open_auto_short_selection_window(self):
        """Selecionar vários ativos short (fixos) para a Auto-Otimização."""
        available_assets = list(self.assets_listbox.get(0, END))
        if not available_assets:
            messagebox.showerror("Erro", "Carregue os dados primeiro (aba Dados)!")
            return

        popup = QDialog(self)
        popup.setWindowTitle("Configuração de Ativos Short - Auto-Otimização")
        popup.resize(600, 700)
        popup.setModal(True)
        main_l = QVBoxLayout(popup)

        lbl = QLabel("Selecione os ativos short (peso fixo, aplicado a cada step):")
        lbl.setFont(bold(12))
        main_l.addWidget(lbl)

        search_l = QHBoxLayout()
        search_l.addWidget(QLabel("🔍 Buscar:"))
        search_entry = QLineEdit()
        search_l.addWidget(search_entry)
        main_l.addLayout(search_l)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll_content = QWidget()
        scroll_layout = QVBoxLayout(scroll_content)
        scroll.setWidget(scroll_content)
        main_l.addWidget(scroll, 1)

        asset_widgets = {}

        def toggle_asset_weight(asset):
            if asset in asset_widgets:
                w = asset_widgets[asset]
                w['weight_entry'].setEnabled(w['use_var'].get())

        def create_asset_widgets():
            _clear_layout(scroll_layout)
            asset_widgets.clear()
            term = search_entry.text().lower()
            filtered = [a for a in available_assets if term in a.lower()]
            for asset in filtered:
                row = QHBoxLayout()
                chk = QCheckBox(asset)
                chk.setChecked(asset in self.auto_short_weights)
                use_var = BoolVar(chk)
                row.addWidget(chk)
                row.addStretch()
                row.addWidget(QLabel("Peso (%):"))
                e_w = QLineEdit()
                e_w.setFixedWidth(80)
                initial = self.auto_short_weights.get(asset, -1.0) * 100
                weight_var = NumVar(e_w, initial)
                row.addWidget(e_w)
                e_w.setEnabled(chk.isChecked())
                chk.toggled.connect(lambda _c, a=asset: toggle_asset_weight(a))
                scroll_layout.addLayout(row)
                asset_widgets[asset] = {'use_var': use_var, 'weight_var': weight_var, 'weight_entry': e_w}
            scroll_layout.addStretch()

        search_entry.textChanged.connect(create_asset_widgets)
        create_asset_widgets()

        quick_box = QGroupBox("Ações Rápidas")
        quick_l = QHBoxLayout(quick_box)

        def select_all(select):
            for a, w in asset_widgets.items():
                w['use_var'].set(select)
                toggle_asset_weight(a)

        def apply_default_weight():
            dw = default_weight_var.get()
            for a, w in asset_widgets.items():
                if w['use_var'].get():
                    w['weight_var'].set(dw)

        b_all = QPushButton("Selecionar Todos")
        b_all.clicked.connect(lambda: select_all(True))
        quick_l.addWidget(b_all)
        b_clear = QPushButton("Limpar Todos")
        b_clear.clicked.connect(lambda: select_all(False))
        quick_l.addWidget(b_clear)
        quick_l.addWidget(QLabel("Peso padrão:"))
        e_dw = QLineEdit()
        e_dw.setFixedWidth(80)
        default_weight_var = NumVar(e_dw, -10.0)
        quick_l.addWidget(e_dw)
        b_apply_dw = QPushButton("Aplicar aos Selecionados")
        b_apply_dw.clicked.connect(apply_default_weight)
        quick_l.addWidget(b_apply_dw)
        main_l.addWidget(quick_box)

        def apply_configuration():
            new_weights = {}
            for asset, w in asset_widgets.items():
                if w['use_var'].get():
                    weight = w['weight_var'].get() / 100
                    if weight >= 0:
                        messagebox.showerror("Erro", f"Peso short deve ser negativo para {asset}!")
                        return
                    new_weights[asset] = weight
            self.auto_short_weights = new_weights
            self.update_auto_short_summary()
            popup.accept()
            messagebox.showinfo("Sucesso", f"Configurados {len(new_weights)} ativos short!")

        btn_l = QHBoxLayout()
        btn_l.addStretch()
        b_cancel = QPushButton("❌ Cancelar")
        b_cancel.clicked.connect(popup.reject)
        btn_l.addWidget(b_cancel)
        b_ok = QPushButton("✅ Aplicar Configuração")
        b_ok.clicked.connect(apply_configuration)
        btn_l.addWidget(b_ok)
        main_l.addLayout(btn_l)

        popup.exec_()

    # ========== MÉTODOS DE CONTROLE ==========

    def calculate_estimate(self):
        """Calcular estimativa de testes"""
        try:
            if not hasattr(self, 'df') or self.df is None:
                self.estimate_label.setText("⚠️ Carregue dados primeiro")
                return

            if 'Data' not in self.df.columns:
                self.estimate_label.setText("⚠️ Coluna Data não encontrada")
                return

            data_inicial = pd.to_datetime(self.df['Data'].iloc[0])
            data_final = pd.to_datetime(self.df['Data'].iloc[-1])
            total_calendar_days = (data_final - data_inicial).days + 1

            otim_count = sum(1 for var in self.otim_windows.values() if var.get())
            rebal_count = sum(1 for var in self.rebalance_periods.values() if var.get())
            obj_count = sum(1 for var in self.objectives.values() if var.get())

            if otim_count == 0 or rebal_count == 0 or obj_count == 0:
                self.estimate_label.setText("❌ Selecione ao menos 1 opção\nde cada categoria")
                return

            period_to_days = {
                '3m': 90, '6m': 180, '1a': 365, '2a': 730, '3a': 1095,
                '1sem': 7, '2sem': 14, '1mes': 30, '2mes': 60, '3mes': 90
            }

            selected_otim = [period for period, var in self.otim_windows.items() if var.get()]
            selected_rebal = [period for period, var in self.rebalance_periods.items() if var.get()]

            max_otim_days = max(period_to_days[period] for period in selected_otim)
            max_rebal_days = max(period_to_days[period] for period in selected_rebal)
            min_required = max_otim_days + max_rebal_days * 2

            if total_calendar_days < min_required:
                self.estimate_label.setText(
                    f"⚠️ Período insuficiente!\n"
                    f"Tem: {total_calendar_days} dias\n"
                    f"Precisa: {min_required} dias"
                )
                return

            total_steps_estimado = 0
            for otim_period in selected_otim:
                otim_days = period_to_days[otim_period]
                for rebal_period in selected_rebal:
                    rebal_days = period_to_days[rebal_period]
                    steps_possiveis = (total_calendar_days - otim_days) // rebal_days
                    if steps_possiveis > 1000000:
                        steps_usados = 1000000
                    else:
                        steps_usados = steps_possiveis
                    total_steps_estimado += steps_usados

            total_configs = otim_count * rebal_count * obj_count
            total_tests = total_steps_estimado * obj_count

            avg_steps = total_steps_estimado / (otim_count * rebal_count) if (otim_count * rebal_count) > 0 else 0

            self.estimate_label.setText(
                f"📊 {total_configs} configurações\n"
                f"⚡ ~{total_tests:,} testes\n"
                f"📅 {total_calendar_days} dias\n"
                f"📈 ~{avg_steps:.0f} steps/config"
            )

        except Exception as e:
            self.estimate_label.setText(f"❌ Erro: {str(e)}")

    def start_auto_optimization(self):
        if not self.validate_auto_config():
            return

        self.auto_optimize_btn.setEnabled(False)
        self.auto_optimize_btn.setText("🔄 Executando...")
        self.status_auto_label.setText("Iniciando auto-otimização...")
        self.status_auto_label.setStyleSheet("color: blue;")

        import threading
        thread = threading.Thread(target=self.run_auto_optimization_thread)
        thread.daemon = True
        thread.start()

    def run_auto_optimization_thread(self):
        """Thread principal de auto-otimização (UI atualizada via sinais Qt)"""
        try:
            configs = self.prepare_test_configurations()

            if len(configs) == 0:
                self.update_auto_status("Nenhuma configuração válida encontrada", 'red')
                return

            self.update_auto_status(f"Iniciando {len(configs)} configurações...", 'blue')

            results = []
            total_configs = len(configs)

            for i, config in enumerate(configs):
                config_desc = f"{config['otim_period']}+{config['rebal_period']}+{config['objective']}"
                status_msg = f"[{i+1}/{total_configs}] {config_desc}"
                self.update_auto_status(status_msg, 'blue')

                config_result = self.run_walk_forward_test(config)

                if config_result:
                    results.append(config_result)
                    success_msg = f"[{i+1}/{total_configs}] ✅ {config_desc} ({config_result['n_steps']} steps)"
                    self.update_auto_status(success_msg, 'green')
                else:
                    fail_msg = f"[{i+1}/{total_configs}] ❌ {config_desc} FALHOU"
                    self.update_auto_status(fail_msg, 'orange')

                import time
                time.sleep(0.1)

            if results:
                self.auto_signals.results_ready.emit(results)
                final_msg = f"🎉 Concluído! {len(results)}/{total_configs} configurações válidas"
                self.update_auto_status(final_msg, 'green')
            else:
                self.update_auto_status("❌ Nenhum resultado válido obtido", 'red')

        except Exception as e:
            self.update_auto_status(f"Erro: {str(e)}", 'red')
        finally:
            self.auto_signals.finished.emit()

    def update_auto_status(self, message, color='black'):
        """Atualizar status (thread-safe via sinal Qt)"""
        self.auto_signals.status.emit(message, color)

    def _set_auto_status(self, message, color):
        self.status_auto_label.setText(message)
        self.status_auto_label.setStyleSheet(f"color: {color};")

    def _auto_finished(self):
        self.auto_optimize_btn.setEnabled(True)
        self.auto_optimize_btn.setText("🚀 INICIAR AUTO-OTIMIZAÇÃO")

    def prepare_test_configurations(self):
        configs = []

        selected_otim = [period for period, var in self.otim_windows.items() if var.get()]
        selected_rebal = [period for period, var in self.rebalance_periods.items() if var.get()]
        selected_obj = [obj for obj, var in self.objectives.items() if var.get()]

        obj_mapping = {'sharpe': 'sharpe', 'sortino': 'sortino',
                       'volatility': 'volatility', 'under_water': 'under_water',
                       'hc10': 'hc10', 'quality_linear': 'quality_linear',
                       'excess_hc10': 'excess_hc10', 'excess_sharpe': 'excess_sharpe'}

        for otim_period in selected_otim:
            for rebal_period in selected_rebal:
                for obj_key in selected_obj:
                    config = {
                        'otim_period': otim_period,
                        'rebal_period': rebal_period,
                        'objective': obj_mapping[obj_key],
                        'rank_min': self.rank_min_var.get(),
                        'rank_max': self.rank_max_var.get(),
                        'weight_min': self.weight_min_var.get() / 100,
                        'weight_max': self.weight_max_var.get() / 100,
                        'use_shorts': self.use_auto_shorts.get() and len(self.auto_short_weights) > 0,
                        'short_weights': dict(self.auto_short_weights) if self.use_auto_shorts.get() else {},
                        # Meta relativa OU absoluta, conforme o modo escolhido
                        'target_return': ((self.auto_meta_var.get() / 100)
                                          if self.use_auto_meta.get()
                                          and self.auto_meta_mode_var.get() == 'relativa' else None),
                        'target_annual': ((self.auto_meta_abs_var.get() / 100)
                                          if self.use_auto_meta.get()
                                          and self.auto_meta_mode_var.get() == 'absoluta' else None),
                        'desc': f"{otim_period}_{rebal_period}_{obj_key}"
                    }
                    configs.append(config)

        return configs

    def run_walk_forward_test(self, config):
        """Executar teste walk-forward (IDÊNTICO À VERSÃO TKINTER)"""
        try:
            period_to_days = {
                '3m': 90, '6m': 180, '1a': 365, '2a': 730, '3a': 1095,
                '1sem': 7, '2sem': 14, '1mes': 30, '2mes': 60, '3mes': 90
            }

            otim_days = period_to_days[config['otim_period']]
            rebal_days = period_to_days[config['rebal_period']]

            if hasattr(self, 'dados_brutos') and self.dados_brutos is not None:
                df_trabalho = self.dados_brutos.copy()
            else:
                print("AVISO: Usando dados já processados. Resultados podem não coincidir com modo manual.")
                df_trabalho = self.df.copy()

            if 'Data' not in df_trabalho.columns:
                print("Coluna 'Data' não encontrada")
                return None

            data_inicial = pd.to_datetime(df_trabalho['Data'].iloc[0])
            data_final = pd.to_datetime(df_trabalho['Data'].iloc[-1])
            total_calendar_days = (data_final - data_inicial).days + 1

            print(f"=== DEBUG WALK-FORWARD (BASE ZERO ÚNICA + ACUMULAÇÃO) ===")
            print(f"Config: {config['desc']}")
            print(f"Data inicial: {data_inicial.strftime('%d/%m/%Y')}")
            print(f"Data final: {data_final.strftime('%d/%m/%Y')}")
            print(f"Total dias corridos: {total_calendar_days}")
            print(f"Janela otimização: {otim_days} dias corridos")
            print(f"Período rebalanceamento: {rebal_days} dias corridos")

            min_required_days = otim_days + rebal_days * 2
            if total_calendar_days < min_required_days:
                print(f"Período insuficiente: {total_calendar_days} < {min_required_days} dias corridos")
                return None

            max_possible_steps = (total_calendar_days - otim_days) // rebal_days
            max_steps = max_possible_steps

            if max_steps > 1000000:
                print(f"⚠️ AVISO: {max_steps} steps é muito! Limitando a 1000000.")
                max_steps = 1000000

            print(f"Steps possíveis: {max_possible_steps}")
            print(f"Max steps que será usado: {max_steps}")

            if max_steps < 1:
                print(f"Steps insuficientes: {max_steps}")
                return None

            ret_acumulado_config = 1.0
            taxa_ref_acumulada_config = 1.0
            primeira_data_validacao = None
            ultima_data_validacao = None
            step_metrics_individuais = []
            volatilidades = []

            for step_num in range(1, max_steps + 1):
                try:
                    inicio_otim_days = (step_num - 1) * rebal_days
                    fim_otim_days = inicio_otim_days + otim_days
                    fim_valid_days = fim_otim_days + rebal_days

                    data_inicio_otim = data_inicial + pd.Timedelta(days=inicio_otim_days)
                    data_fim_otim = data_inicial + pd.Timedelta(days=fim_otim_days - 1)
                    data_fim_valid = data_inicial + pd.Timedelta(days=fim_valid_days - 1)

                    print(f"\nStep {step_num}:")
                    print(f"  Otimização: {data_inicio_otim.strftime('%d/%m/%Y')} a {data_fim_otim.strftime('%d/%m/%Y')}")
                    print(f"  Validação: {(data_fim_otim + pd.Timedelta(days=1)).strftime('%d/%m/%Y')} a {data_fim_valid.strftime('%d/%m/%Y')}")

                    if data_fim_valid > data_final:
                        print(f"PAROU: data_fim_valid > data_final")
                        break

                    df_periodo_completo_bruto = df_trabalho[
                        (pd.to_datetime(df_trabalho['Data']) >= data_inicio_otim) &
                        (pd.to_datetime(df_trabalho['Data']) <= data_fim_valid)
                    ].copy().reset_index(drop=True)

                    print(f"  Período completo: {len(df_periodo_completo_bruto)} registros")
                    print(f"  Aplicando transformação base zero ÚNICA...")

                    df_completo_base0, cols_removidas = self.transformar_base_zero(
                        df_periodo_completo_bruto.set_index('Data')
                    )

                    if df_completo_base0 is None:
                        print(f"  Falha na transformação base zero")
                        continue

                    df_completo_base0 = df_completo_base0.reset_index()
                    df_completo_base0.rename(columns={'index': 'Data'}, inplace=True)

                    df_otim_base0 = df_completo_base0[
                        (pd.to_datetime(df_completo_base0['Data']) >= data_inicio_otim) &
                        (pd.to_datetime(df_completo_base0['Data']) <= data_fim_otim)
                    ].copy().reset_index(drop=True)

                    df_valid_base0 = df_completo_base0[
                        (pd.to_datetime(df_completo_base0['Data']) > data_fim_otim) &
                        (pd.to_datetime(df_completo_base0['Data']) <= data_fim_valid)
                    ].copy().reset_index(drop=True)

                    print(f"  Após separação: otim={len(df_otim_base0)}, valid={len(df_valid_base0)}")

                    if len(df_otim_base0) < 10 or len(df_valid_base0) < 2:
                        print(f"  Dados insuficientes após base zero")
                        continue

                    step_result = self.execute_single_step(
                        df_otim_base0, df_valid_base0, df_completo_base0, config, step_num
                    )

                    if step_result:
                        retorno_step = step_result['total_return']
                        ret_acumulado_config *= (1 + retorno_step)

                        if 'risk_free_period' in step_result:
                            taxa_ref_step = step_result['risk_free_period']
                            taxa_ref_acumulada_config *= (1 + taxa_ref_step)

                        if 'volatility' in step_result:
                            volatilidades.append(step_result['volatility'])

                        if primeira_data_validacao is None:
                            primeira_data_validacao = pd.to_datetime(df_valid_base0['Data'].iloc[0])
                            print(f"  📅 Primeira data validação guardada: {primeira_data_validacao.strftime('%d/%m/%Y')}")

                        ultima_data_validacao = pd.to_datetime(df_valid_base0['Data'].iloc[-1])

                        step_metrics_individuais.append(step_result)

                        print(f"  ✅ Step {step_num} concluído e acumulado")
                        print(f"     Retorno do step: {retorno_step:.2%}")
                        print(f"     Retorno acumulado até agora: {(ret_acumulado_config-1):.2%}")
                    else:
                        print(f"  ⚠️ Step {step_num} FALHOU")

                except Exception as e:
                    print(f"  ERRO no step {step_num}: {str(e)}")
                    import traceback
                    traceback.print_exc()

            print(f"\n=== LOOP CONCLUÍDO ===")
            print(f"Steps com sucesso: {len(step_metrics_individuais)}")

            if step_metrics_individuais and len(step_metrics_individuais) >= 1:
                retorno_total_config = ret_acumulado_config - 1
                taxa_ref_total_config = taxa_ref_acumulada_config - 1

                if primeira_data_validacao and ultima_data_validacao:
                    dias_totais = (ultima_data_validacao - primeira_data_validacao).days + 1

                    if dias_totais > 0:
                        retorno_anualizado = (1 + retorno_total_config) ** (365 / dias_totais) - 1
                        taxa_ref_anualizada = (1 + taxa_ref_total_config) ** (365 / dias_totais) - 1
                    else:
                        retorno_anualizado = 0
                        taxa_ref_anualizada = 0
                else:
                    retorno_anualizado = retorno_total_config
                    taxa_ref_anualizada = taxa_ref_total_config
                    dias_totais = 0

                vol_media = sum(volatilidades) / len(volatilidades) if volatilidades else 0

                if vol_media > 0:
                    sharpe_final = (retorno_anualizado - taxa_ref_anualizada) / vol_media
                else:
                    sharpe_final = 0

                print(f"📊 CÁLCULO SHARPE FINAL:")
                print(f"   Retorno anualizado: {retorno_anualizado:.2%}")
                print(f"   Taxa ref anualizada: {taxa_ref_anualizada:.2%}")
                print(f"   Volatilidade média: {vol_media:.2%}")
                print(f"   Sharpe = {sharpe_final:.3f}")

                return self.calculate_average_metrics_v2(
                    step_metrics_individuais, config, retorno_anualizado,
                    taxa_ref_anualizada, vol_media, sharpe_final,
                    retorno_total_config, dias_totais
                )
            else:
                print(f"Steps insuficientes para análise: {len(step_metrics_individuais)}")
                return None

        except Exception as e:
            print(f"Erro no walk-forward test: {str(e)}")
            import traceback
            traceback.print_exc()
            return None

    def execute_single_step(self, df_otim, df_valid, df_completo, config, step_num):
        """Step walk-forward (IDÊNTICO À VERSÃO TKINTER)"""
        try:
            print(f"\n🎯 STEP {step_num} - VALIDAÇÃO DE MÉTRICAS")

            ranking_result = calculate_asset_ranking(
                df_otim, self.risk_free_column_name,
                peso_inc=0.33, peso_desv=0.33, peso_cor=0.33
            )

            if not ranking_result:
                return None

            df_ranking = ranking_result['ranking']
            score_min = config['rank_min'] / 100
            score_max = config['rank_max'] / 100

            filtered_ranking = df_ranking[
                (df_ranking['Índice'] >= score_min) &
                (df_ranking['Índice'] <= score_max)
            ]

            if len(filtered_ranking) < 2:
                return None

            selected_assets = filtered_ranking['Ativo'].tolist()
            print(f"✅ {len(selected_assets)} ativos selecionados")

            original_df = self.df
            original_optimizer = self.optimizer
            original_result = getattr(self, 'result', None)

            try:
                self.df = df_otim.copy()

                # Shorts fixos configurados (podem ser vários), disponíveis neste período
                short_weights_cfg = config.get('short_weights', {})
                short_assets = [a for a in short_weights_cfg if a in df_otim.columns]
                # Um ativo não pode ser long (via ranking) e short ao mesmo tempo
                if short_assets:
                    selected_assets = [a for a in selected_assets if a not in short_assets]
                use_shorts_step = bool(config['use_shorts'] and short_assets and len(selected_assets) >= 1)

                if use_shorts_step:
                    all_assets = selected_assets + short_assets
                else:
                    all_assets = selected_assets

                self.optimizer = PortfolioOptimizer(self.df, all_assets)

                if hasattr(self.optimizer, 'risk_free_rate_total'):
                    risk_free_rate = self.optimizer.risk_free_rate_total
                    print(f"🎯 Taxa para otimização: {risk_free_rate:.4f} ({risk_free_rate:.2%})")
                else:
                    risk_free_rate = 0.0

                objective_map = {
                    'sharpe': 'sharpe', 'sortino': 'sortino',
                    'volatility': 'volatility', 'under_water': 'under_water',
                    'hc10': 'hc10', 'quality_linear': 'quality_linear',
                    'excess_hc10': 'excess_hc10', 'excess_sharpe': 'excess_sharpe'
                }

                target_return = config.get('target_return')
                target_annual = config.get('target_annual')

                if use_shorts_step:
                    print(f"🔄 OTIMIZAÇÃO COM SHORTS ({len(short_assets)} ativos)")
                    self.result = self.optimizer.optimize_portfolio_with_shorts(
                        selected_assets=selected_assets,
                        short_assets=short_assets,
                        short_weights={a: short_weights_cfg[a] for a in short_assets},
                        objective_type=objective_map[config['objective']],
                        target_return=target_return,
                        target_annual=target_annual,
                        max_weight=config['weight_max'],
                        min_weight=config['weight_min'],
                        risk_free_rate=risk_free_rate,
                        individual_constraints=None
                    )
                else:
                    print(f"📊 OTIMIZAÇÃO NORMAL (SEM SHORTS)")
                    self.result = self.optimizer.optimize_portfolio(
                        objective_type=objective_map[config['objective']],
                        target_return=target_return,
                        target_annual=target_annual,
                        max_weight=config['weight_max'],
                        min_weight=config['weight_min'],
                        risk_free_rate=risk_free_rate,
                        individual_constraints=None
                    )

                if not self.result['success']:
                    print(f"❌ Otimização falhou: {self.result.get('message', 'Erro desconhecido')}")
                    return None

                if self.result.get('degraded'):
                    # Degradação graciosa: o step segue, mas fica registrado no log
                    print(f"⚠️ ATENÇÃO: {self.result['degraded_message']}")

                # A meta é exigida DENTRO da janela de otimização (in-sample).
                # Guardamos aqui se ela foi cumprida neste step, para depois medir
                # quantos steps de fato bateram o alvo — o que separa "não atingiu
                # porque era inatingível" de "atingiu mas não se sustentou fora
                # da amostra". None = meta não estava em uso.
                meta_ok_step = self.result.get('meta_atingida') if self.result.get('meta_used') else None
                if meta_ok_step is False:
                    print(f"⚠️ Meta NÃO atingida neste step (alvo in-sample não alcançável)")

                optimized_weights = self.result['weights']
                optimized_assets = self.result['assets']

                print(f"\n🔍 COMPOSIÇÃO DO PORTFÓLIO:")
                for i, asset in enumerate(optimized_assets):
                    weight = optimized_weights[i]
                    if abs(weight) > 0.001:
                        tipo = "SHORT" if weight < 0 else "LONG"
                        print(f"   {asset}: {weight:.3f} ({weight*100:.1f}%) - {tipo}")

                print(f"✅ Otimização concluída")
                print("\n🔍 Calculando OUT-OF-SAMPLE (BASE ZERO ÚNICA)")

                available_assets = [asset for asset in optimized_assets if asset in df_completo.columns]
                if len(available_assets) != len(optimized_assets):
                    print("⚠️ Nem todos os ativos disponíveis no período completo")
                    return None

                optimizer_completo = PortfolioOptimizer(df_completo, available_assets)

                portfolio_returns_completo = np.dot(optimizer_completo.returns_data.values, optimized_weights)
                cumulative_completo = np.cumsum(portfolio_returns_completo)

                n_dias_otim = len(df_otim)
                n_dias_total = len(df_completo)
                n_dias_valid = n_dias_total - n_dias_otim

                print(f"📊 Divisão do período completo:")
                print(f"   Total: {n_dias_total} registros")
                print(f"   Otimização: {n_dias_otim} registros")
                print(f"   Validação: {n_dias_valid} registros")

                if n_dias_valid <= 0:
                    print("⚠️ Período de validação muito curto")
                    return None

                primeira_data_valid = pd.to_datetime(df_valid['Data'].iloc[0])
                ultima_data_valid = pd.to_datetime(df_valid['Data'].iloc[-1])
                n_dias_corridos_valid = (ultima_data_valid - primeira_data_valid).days + 1

                print(f"\n📅 Período validação: {primeira_data_valid.strftime('%d/%m/%Y')} a {ultima_data_valid.strftime('%d/%m/%Y')}")
                print(f"   Dias corridos: {n_dias_corridos_valid}")

                portfolio_acum_fim_otim = cumulative_completo[n_dias_otim - 1]
                portfolio_acum_fim_valid = cumulative_completo[-1]

                print(f"\n📊 RETORNO DO PORTFÓLIO (BASE ZERO ÚNICA):")
                print(f"   Acumulado fim otimização: {portfolio_acum_fim_otim:.6f}")
                print(f"   Acumulado fim validação: {portfolio_acum_fim_valid:.6f}")

                retorno_periodo_valid = (1 + portfolio_acum_fim_valid) / (1 + portfolio_acum_fim_otim) - 1
                print(f"   Retorno período validação: {retorno_periodo_valid:.4f} ({retorno_periodo_valid:.2%})")

                if n_dias_corridos_valid > 0:
                    retorno_anual_valid = (1 + retorno_periodo_valid) ** (365 / n_dias_corridos_valid) - 1
                else:
                    retorno_anual_valid = 0
                print(f"   Retorno anualizado: {retorno_anual_valid:.4f} ({retorno_anual_valid:.2%})")

                if hasattr(optimizer_completo, 'risk_free_cumulative') and optimizer_completo.risk_free_cumulative is not None:
                    try:
                        taxa_acum_fim_otim = optimizer_completo.risk_free_cumulative.iloc[n_dias_otim - 1]
                        taxa_acum_fim_valid = optimizer_completo.risk_free_cumulative.iloc[-1]

                        print(f"\n📊 TAXA LIVRE DE RISCO (BASE ZERO ÚNICA):")
                        print(f"   Acumulada fim otimização: {taxa_acum_fim_otim:.6f}")
                        print(f"   Acumulada fim validação: {taxa_acum_fim_valid:.6f}")

                        taxa_periodo_valid = (1 + taxa_acum_fim_valid) / (1 + taxa_acum_fim_otim) - 1
                        print(f"   Taxa período validação: {taxa_periodo_valid:.4f} ({taxa_periodo_valid:.2%})")

                        if n_dias_corridos_valid > 0:
                            taxa_anual_valid = (1 + taxa_periodo_valid) ** (365 / n_dias_corridos_valid) - 1
                        else:
                            taxa_anual_valid = 0
                        print(f"   Taxa anualizada: {taxa_anual_valid:.4f} ({taxa_anual_valid:.2%})")
                    except Exception as e:
                        print(f"   Erro no cálculo da taxa: {e}")
                        taxa_periodo_valid = 0
                        taxa_anual_valid = 0
                else:
                    taxa_periodo_valid = 0
                    taxa_anual_valid = 0
                    print("   Taxa livre não disponível")

                portfolio_cumulative_validacao = cumulative_completo[n_dias_otim:]

                if len(portfolio_cumulative_validacao) > 1:
                    portfolio_cumulative_validacao_completo = np.concatenate(
                        [[portfolio_acum_fim_otim], portfolio_cumulative_validacao]
                    )
                    variac_result_pu = (1 + portfolio_cumulative_validacao_completo[1:]) / (1 + portfolio_cumulative_validacao_completo[:-1])
                    portfolio_returns_pct_valid = variac_result_pu - 1
                    vol_valid = np.std(portfolio_returns_pct_valid, ddof=0) * np.sqrt(252)
                else:
                    vol_valid = 0
                    portfolio_returns_pct_valid = np.array([])

                print(f"\n📊 VOLATILIDADE:")
                print(f"   Volatilidade anualizada: {vol_valid:.4f} ({vol_valid:.2%})")

                excesso_anual_valid = retorno_anual_valid - taxa_anual_valid
                print(f"\n📊 PERFORMANCE:")
                print(f"   Excesso anualizado: {excesso_anual_valid:.4f} ({excesso_anual_valid:.2%})")

                if vol_valid > 0:
                    sharpe_valid = excesso_anual_valid / vol_valid
                else:
                    sharpe_valid = 0
                print(f"   Sharpe Ratio: {sharpe_valid:.3f}")

                if len(portfolio_returns_pct_valid) > 0:
                    mean_daily_return = np.mean(portfolio_returns_pct_valid)
                    std_daily_return = np.std(portfolio_returns_pct_valid, ddof=0)
                    var_95_daily_valid = mean_daily_return - 1.65 * std_daily_return
                else:
                    var_95_daily_valid = 0
                print(f"   VaR 95% diário: {var_95_daily_valid:.4f} ({var_95_daily_valid:.2%})")

                if len(portfolio_returns_pct_valid) > 0:
                    positive_returns = portfolio_returns_pct_valid[portfolio_returns_pct_valid > 0]
                    total_returns = len(portfolio_returns_pct_valid)
                    positive_return_pct = len(positive_returns) / total_returns if total_returns > 0 else 0
                else:
                    positive_return_pct = 0

                print(f"\n✅ RESUMO OUT-OF-SAMPLE:")
                print(f"   Retorno Total: {retorno_periodo_valid:.2%}")
                print(f"   Sharpe: {sharpe_valid:.3f}")

                return {
                    'n_assets': len(available_assets),
                    'sharpe': sharpe_valid,
                    'annual_return': retorno_anual_valid,
                    'volatility': vol_valid,
                    'positive_return_pct': positive_return_pct,
                    'total_return': retorno_periodo_valid,
                    'risk_free_annual': taxa_anual_valid,
                    'risk_free_period': taxa_periodo_valid,
                    'excess_annual': excesso_anual_valid,
                    'var_95': var_95_daily_valid,
                    'n_days': n_dias_corridos_valid,
                    'meta_atingida': meta_ok_step
                }

            finally:
                self.df = original_df
                self.optimizer = original_optimizer
                if original_result is not None:
                    self.result = original_result
                else:
                    if hasattr(self, 'result'):
                        delattr(self, 'result')

        except Exception as e:
            print(f"❌ Erro: {str(e)}")
            import traceback
            traceback.print_exc()
            return None

    def calculate_average_metrics_v2(self, step_metrics, config, retorno_anualizado,
                                     taxa_ref_anualizada, vol_media, sharpe_final,
                                     retorno_total, dias_totais):
        avg_metrics = {}

        values = [step['n_assets'] for step in step_metrics
                  if 'n_assets' in step and step['n_assets'] is not None]
        avg_metrics['n_assets'] = sum(values) / len(values) if values else 0

        # VaR 95% diário médio dos steps (risco de cauda)
        var_values = [step['var_95'] for step in step_metrics
                      if step.get('var_95') is not None]
        avg_metrics['var_95'] = sum(var_values) / len(var_values) if var_values else 0

        # % de PERÍODOS (steps) com resultado positivo — não por dia, por rebalanceamento:
        #  - Pos Abs : retorno do período > 0
        #  - Pos>Ref : retorno do período > taxa de referência do período (excesso do período > 0)
        n = len(step_metrics)
        if n > 0:
            avg_metrics['positive_return_pct'] = sum(
                1 for s in step_metrics if s.get('total_return', 0) > 0) / n
            avg_metrics['positive_vs_ref_pct'] = sum(
                1 for s in step_metrics if s.get('total_return', 0) > s.get('risk_free_period', 0)) / n
        else:
            avg_metrics['positive_return_pct'] = 0
            avg_metrics['positive_vs_ref_pct'] = 0

        # % de steps que cumpriram a META dentro da janela de otimização.
        # Diferente das colunas acima (que medem o resultado FORA da amostra),
        # esta mostra se o alvo era alcançável onde o otimizador podia agir:
        #  - 100% com Ret% abaixo do alvo => a meta foi batida in-sample mas
        #    não se sustentou out-of-sample (restrição sem folga)
        #  - abaixo de 100% => em alguns steps o alvo era inatingível e valeu o
        #    fallback de "maior retorno possível"
        # None quando a meta não estava em uso (coluna exibe "—").
        meta_flags = [s['meta_atingida'] for s in step_metrics
                      if s.get('meta_atingida') is not None]
        avg_metrics['meta_ok_pct'] = (
            sum(1 for f in meta_flags if f) / len(meta_flags) if meta_flags else None)

        avg_metrics['annual_return'] = retorno_anualizado
        avg_metrics['risk_free_annual'] = taxa_ref_anualizada
        avg_metrics['volatility'] = vol_media
        avg_metrics['sharpe'] = sharpe_final

        print(f"\n📊 MÉTRICAS FINAIS DA CONFIGURAÇÃO:")
        print(f"   Retorno Anual: {retorno_anualizado:.2%}")
        print(f"   Sharpe Final: {sharpe_final:.3f}")
        print(f"   Steps com sucesso: {len(step_metrics)}")

        result = {
            'config': config,
            'metrics': avg_metrics,
            'n_steps': len(step_metrics),
            'total_return': retorno_total,
            'total_days': dias_totais
        }

        return result

    def process_and_display_results(self, results):
        # Desabilita ordenação durante a inserção para evitar reordenação a cada linha
        self.auto_results_tree.setSortingEnabled(False)
        self.auto_results_tree.setRowCount(0)

        results.sort(key=lambda x: x['metrics']['sharpe'], reverse=True)

        obj_names = {
            'sharpe': 'Sharpe', 'sortino': 'Sortino', 'volatility': 'MinRisco', 'under_water': 'MinUW',
            'hc10': 'Inc/R²Vol', 'quality_linear': 'Qualidade', 'excess_hc10': 'Lin.Excesso',
            'excess_sharpe': 'Sharpe.Exc'
        }

        for i, result in enumerate(results):
            config = result['config']
            metrics = result['metrics']

            # (texto exibido, chave numérica para ordenação — None = ordena como texto)
            cells = (
                (str(i + 1), i + 1),
                (config['otim_period'], None),
                (config['rebal_period'], None),
                (obj_names.get(config['objective'], config['objective']), None),
                (str(int(metrics['n_assets'])), metrics['n_assets']),
                (f"{metrics['sharpe']:.3f}", metrics['sharpe']),
                (f"{metrics['annual_return']:.1%}", metrics['annual_return']),
                # MetaOK%: "—" quando a meta não foi usada (ordena por último)
                ((f"{metrics['meta_ok_pct']:.0%}", metrics['meta_ok_pct'])
                 if metrics.get('meta_ok_pct') is not None else ("—", -1.0)),
                (f"{metrics.get('risk_free_annual', 0):.1%}", metrics.get('risk_free_annual', 0)),
                (f"{metrics['volatility']:.1%}", metrics['volatility']),
                (f"{metrics.get('var_95', 0):.2%}", metrics.get('var_95', 0)),
                (f"{metrics.get('positive_vs_ref_pct', 0):.1%}", metrics.get('positive_vs_ref_pct', 0)),
                (f"{metrics['positive_return_pct']:.1%}", metrics['positive_return_pct']),
            )

            r = self.auto_results_tree.rowCount()
            self.auto_results_tree.insertRow(r)
            for c, (text, key) in enumerate(cells):
                item = _SortableItem(str(text))
                item.setTextAlignment(Qt.AlignCenter)
                if key is not None:
                    item.setData(Qt.UserRole, float(key))
                self.auto_results_tree.setItem(r, c, item)

        # Reabilita ordenação; começa ordenado por Rank (coluna 0, ascendente)
        self.auto_results_tree.setSortingEnabled(True)
        self.auto_results_tree.sortItems(0, Qt.AscendingOrder)

    # Especificação de exportação: (cabeçalho, tipo) por coluna, na ordem da tabela.
    # tipo: 'text' | 'int' | 'float' | 'pctN' (fração; N = casas decimais ao exibir)
    AUTO_EXPORT_SPEC = (
        ('Rank', 'int'),
        ('Otimização', 'text'),
        ('Rebalanceamento', 'text'),
        ('Objetivo', 'text'),
        ('N_Ativos', 'int'),
        ('Sharpe', 'float'),
        ('Retorno(%)', 'pct1'),
        ('Meta_OK(%)', 'pct0'),
        ('Taxa_Ref(%)', 'pct1'),
        ('Volatilidade(%)', 'pct1'),
        ('VaR95%(diário)', 'pct2'),
        ('Pos>Ref(%)', 'pct1'),
        ('Pos_Abs(%)', 'pct1'),
    )

    def _auto_results_rows(self, pct_as_number=False):
        """
        Linhas da tabela como valores TIPADOS (números de verdade, não texto),
        respeitando a ordenação atual da tabela.

        Os números vêm de Qt.UserRole (guardado na criação das células), não do
        texto exibido — assim a exportação não herda o "%" nem o separador
        decimal da interface.

        pct_as_number=False -> percentuais como fração (0.218); use com formato
                               de célula percentual (Excel).
        pct_as_number=True  -> percentuais já multiplicados (21.8); use em texto
                               puro (CSV), onde não há formatação.
        Células sem valor (Meta_OK "—") viram None (célula vazia).
        """
        spec = self.AUTO_EXPORT_SPEC
        rows = []
        for r in range(self.auto_results_tree.rowCount()):
            row = []
            for c in range(self.auto_results_tree.columnCount()):
                item = self.auto_results_tree.item(r, c)
                if item is None:
                    row.append(None)
                    continue
                kind = spec[c][1] if c < len(spec) else 'text'
                if kind == 'text':
                    row.append(item.text())
                    continue
                # "—" marca ausência de valor (meta não utilizada)
                if item.text().strip() in ('—', '-', ''):
                    row.append(None)
                    continue
                key = item.data(Qt.UserRole)
                if key is None:
                    row.append(item.text())
                elif kind == 'int':
                    row.append(int(round(float(key))))
                elif kind == 'float':
                    row.append(round(float(key), 6))
                else:
                    # Arredondamento necessário: multiplicar a fração por 100
                    # expõe o ruído binário (0.145*100 = 14.499999999999998).
                    # 4 casas na escala percentual preservam toda a precisão útil.
                    row.append(round(float(key) * 100, 4) if pct_as_number
                               else round(float(key), 6))
            rows.append(row)
        return rows

    def export_auto_results_csv(self):
        if self.auto_results_tree.rowCount() == 0:
            messagebox.showerror("Erro", "Não há resultados para exportar!")
            return

        try:
            filename = filedialog.asksaveasfilename(
                title="Salvar Resultados Auto-Otimização",
                defaultextension=".csv",
                filetypes=[("CSV files", "*.csv"), ("All files", "*.*")]
            )

            if filename:
                columns = [h for h, _ in self.AUTO_EXPORT_SPEC]
                # Percentuais já multiplicados: no CSV não há formato de célula,
                # então "21,8" sob o cabeçalho "Retorno(%)" é o que faz sentido.
                data = self._auto_results_rows(pct_as_number=True)
                df = pd.DataFrame(data, columns=columns)
                # Padrão brasileiro: separador ";" e decimal "," — assim o Excel
                # pt-BR abre o arquivo já com os números reconhecidos como número.
                df.to_csv(filename, index=False, encoding='utf-8-sig',
                          sep=';', decimal=',')
                messagebox.showinfo("Sucesso", f"Resultados exportados para:\n{filename}")

        except Exception as e:
            messagebox.showerror("Erro", f"Erro ao exportar CSV:\n{str(e)}")

    def export_auto_results_excel(self):
        if self.auto_results_tree.rowCount() == 0:
            messagebox.showerror("Erro", "Não há resultados para exportar!")
            return

        try:
            filename = filedialog.asksaveasfilename(
                title="Salvar Resultados Auto-Otimização",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")]
            )

            if filename:
                spec = self.AUTO_EXPORT_SPEC
                columns = [h for h, _ in spec]
                # Percentuais como FRAÇÃO: combinados com o formato de célula
                # percentual abaixo, viram percentual nativo do Excel. O separador
                # decimal passa a ser o do sistema (vírgula no Brasil).
                data = self._auto_results_rows(pct_as_number=False)
                df = pd.DataFrame(data, columns=columns)

                # Formato de número por tipo de coluna
                fmt_by_kind = {
                    'int': '0',
                    'float': '0.000',
                    'pct0': '0%',
                    'pct1': '0.0%',
                    'pct2': '0.00%',
                }

                with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                    df.to_excel(writer, sheet_name='Resultados Auto-Otimização', index=False)

                    worksheet = writer.sheets['Resultados Auto-Otimização']

                    # Cabeçalho em negrito e centralizado
                    from openpyxl.styles import Font, Alignment
                    for cell in worksheet[1]:
                        cell.font = Font(bold=True)
                        cell.alignment = Alignment(horizontal='center')

                    # Aplica o formato numérico e o alinhamento em cada coluna
                    for idx, (_, kind) in enumerate(spec, start=1):
                        number_format = fmt_by_kind.get(kind)
                        letter = worksheet.cell(row=1, column=idx).column_letter
                        for row_i in range(2, worksheet.max_row + 1):
                            cell = worksheet.cell(row=row_i, column=idx)
                            if number_format:
                                cell.number_format = number_format
                            cell.alignment = Alignment(horizontal='center')

                    # Largura das colunas pelo conteúdo já formatado
                    for column in worksheet.columns:
                        max_length = 0
                        column_letter = column[0].column_letter
                        for cell in column:
                            try:
                                if len(str(cell.value)) > max_length:
                                    max_length = len(str(cell.value))
                            except Exception:
                                pass
                        adjusted_width = min(max_length + 3, 30)
                        worksheet.column_dimensions[column_letter].width = adjusted_width

                    # Congela o cabeçalho e liga o autofiltro
                    worksheet.freeze_panes = 'A2'
                    worksheet.auto_filter.ref = worksheet.dimensions

                messagebox.showinfo("Sucesso", f"Resultados exportados para:\n{filename}")

        except Exception as e:
            messagebox.showerror("Erro", f"Erro ao exportar Excel:\n{str(e)}")

    def validate_auto_config(self):
        if not hasattr(self, 'df') or self.df is None:
            messagebox.showerror("Erro", "Carregue e processe dados primeiro!")
            return False

        if not any(var.get() for var in self.otim_windows.values()):
            messagebox.showerror("Erro", "Selecione pelo menos uma janela de otimização!")
            return False

        if not any(var.get() for var in self.rebalance_periods.values()):
            messagebox.showerror("Erro", "Selecione pelo menos um período de rebalanceamento!")
            return False

        if not any(var.get() for var in self.objectives.values()):
            messagebox.showerror("Erro", "Selecione pelo menos um objetivo!")
            return False

        return True


def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    window = PortfolioOptimizerGUI()
    window.show()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()
